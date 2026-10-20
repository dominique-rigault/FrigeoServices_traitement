"""Décisions du métier ligne par ligne (étape 6, morceau 6e1).

Ce module ne contient rien de propre à un client. Ce que les règles automatiques
du nettoyage ne savent pas trancher (une anomalie de cellule non corrigée, un écart
« à arbitrer » des règles de gestion, un doublon qui n'est pas exact, une clé en
double) est soumis au métier dans un classeur Excel d'arbitrage, puis ses décisions
sont relues, contrôlées et enregistrées dans un fichier hors du dépôt.

Trois sortes d'éléments (voir `elements_a_arbitrer`) :

- « cellule » : une anomalie de format, ou un écart « à arbitrer » d'une règle de
  gestion qui porte sur une seule colonne d'une ligne; décisions « corriger » (avec
  la valeur retenue) ou « garder » (la valeur est juste);
- « ligne » : une ligne en doublon; décisions « retirer » ou « garder »;
- « écart » : un écart « à arbitrer » d'une règle de gestion qui porte sur
  plusieurs colonnes, ou sur un total de période; décision « garder » (examiné).

Chaque décision est repérée par l'identifiant stable de l'anomalie (`id_anomalie`)
et garde une empreinte (hachage) de ce qui a été constaté : la valeur de la
cellule, la ligne entière hors lignage, ou les valeurs de l'écart. La valeur
elle-même n'est recopiée nulle part; seule la valeur retenue par le métier est
écrite. Une décision dont l'empreinte ne correspond plus aux données est caduque :
l'élément est de nouveau proposé au métier.

Fichiers, tous hors du dépôt :

- data/decisions/arbitrage_AAAA-MM.xlsx : le classeur d'arbitrage
  (`exporter_arbitrage`, `importer_arbitrage`); il porte des valeurs du client, y
  compris celles des colonnes sensibles;
- data/config/decisions.yaml : les décisions enregistrées et le journal des
  imports (`charger_decisions`, `ecrire_decisions`), avec une copie datée dans
  data/config/sauvegardes/ avant chaque écriture.

L'application des décisions aux tables nettoyées relève du morceau 6e2.
"""

from __future__ import annotations

import hashlib
import os
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import NamedTuple

import pandas as pd
import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.utils import get_column_letter

from . import racine_projet
from .anomalies import (
    COLONNES_ANOMALIES,
    ETAPE_DOUBLONS,
    ETAPE_FORMAT,
    ETAPE_GESTION,
    _index_lignage,
    _repere,
)
from .controles import ISSUE_A_ARBITRER, ISSUE_ANOMALIE, _colonnes_de_donnees, _valeur_constatee
from .dictionnaire import (
    _sauvegarder,
    chemin_dictionnaire,
    horodater,
    lire_instant,
)
from .doublons import SEPARATEUR, _texte
from .nettoyage import _convertir
from .profilage import _interpretations
from .revue import _ecrire_texte, _est_vide, _menu
from .typage import NATURES_LISIBLES, ecrire_date

SORTE_CELLULE = "cellule"
SORTE_LIGNE = "ligne"
SORTE_ECART = "écart"
SORTES = (SORTE_CELLULE, SORTE_LIGNE, SORTE_ECART)

DECISION_CORRIGER = "corriger"
DECISION_GARDER = "garder"
DECISION_RETIRER = "retirer"
DECISIONS = (DECISION_CORRIGER, DECISION_GARDER, DECISION_RETIRER)
DECISIONS_PAR_SORTE = {
    SORTE_CELLULE: (DECISION_CORRIGER, DECISION_GARDER),
    SORTE_LIGNE: (DECISION_RETIRER, DECISION_GARDER),
    SORTE_ECART: (DECISION_GARDER,),
}

ETAT_SANS_DECISION = "sans décision"
ETAT_VALABLE = "décision valable"
ETAT_CADUQUE = "décision caduque"
ETAT_INTROUVABLE = "ligne source introuvable"
ETATS = (ETAT_SANS_DECISION, ETAT_CADUQUE, ETAT_VALABLE, ETAT_INTROUVABLE)
# Seuls ces éléments sont proposés au métier.
ETATS_A_ARBITRER = (ETAT_SANS_DECISION, ETAT_CADUQUE)

LONGUEUR_EMPREINTE = 16
VALEUR_MASQUEE = "masqué"

COLONNES_ELEMENTS = (
    "id_anomalie",
    "sorte",
    "etat_decision",
    "etape",
    "table",
    "fichier_source",
    "feuille_source",
    "num_ligne_source",
    "cle",
    "champ",
    "valeur",
    "issue",
    "regle",
    "constat",
    "detail",
    "proposition",
    "empreinte",
)

FEUILLE_MODE_EMPLOI = "mode_emploi"
FEUILLE_ARBITRAGE = "arbitrage"
COLONNES_ARBITRAGE = (
    "table",
    "fichier_source",
    "feuille_source",
    "num_ligne_source",
    "cle",
    "champ",
    "valeur",
    "issue",
    "regle",
    "constat",
    "detail",
    "proposition",
    "sorte",
    "decision",
    "valeur_retenue",
    "commentaire",
    "id_anomalie",
    "empreinte",
)
MODIFIABLES_ARBITRAGE = ("decision", "valeur_retenue", "commentaire")
COLONNES_TECHNIQUES = ("id_anomalie", "empreinte")
LARGEURS_ARBITRAGE = {
    "table": 20,
    "fichier_source": 30,
    "feuille_source": 14,
    "num_ligne_source": 10,
    "cle": 22,
    "champ": 24,
    "valeur": 30,
    "issue": 11,
    "regle": 30,
    "constat": 34,
    "detail": 30,
    "proposition": 18,
    "sorte": 9,
    "decision": 11,
    "valeur_retenue": 24,
    "commentaire": 40,
    "id_anomalie": 14,
    "empreinte": 18,
}
REPERE_IMPORTS = "Imports de décisions déjà faits (repère technique)"

STRUCTURE_DECISIONS = 1
CHAMPS_DECISION = (
    "sorte",
    "table",
    "fichier_source",
    "feuille_source",
    "num_ligne_source",
    "champ",
    "constat",
    "empreinte",
    "decision",
    "valeur_retenue",
    "nature",
    "commentaire",
    "decide_le",
    "classeur",
)
CHAMPS_IMPORT = ("date", "classeur", "periode_fin", "decisions", "remplacees", *DECISIONS)
COLONNES_DECISIONS_LUES = (
    "id_anomalie",
    "sorte",
    "table",
    "fichier_source",
    "feuille_source",
    "num_ligne_source",
    "champ",
    "constat",
    "empreinte",
    "decision",
    "valeur_retenue",
    "nature",
    "commentaire",
    "ligne_excel",
)


# --------------------------------------------------------------------------------
# Emplacements et fichier des décisions


def chemin_arbitrage(periode_fin: str) -> Path:
    """Emplacement par défaut du classeur : data/decisions/, hors du dépôt."""
    return racine_projet() / "data" / "decisions" / f"arbitrage_{periode_fin}.xlsx"


def chemin_decisions() -> Path:
    """Emplacement par défaut des décisions : data/config/decisions.yaml.

    Les décisions portent des valeurs retenues par le métier : elles sont rangées
    avec les données, hors du dépôt, à côté du dictionnaire.
    """
    return chemin_dictionnaire().parent / "decisions.yaml"


def _vide() -> dict:
    return {"meta": {"structure": STRUCTURE_DECISIONS, "imports": []}, "decisions": {}}


def _valeur_admise(valeur) -> bool:
    """Valeur retenue que le fichier YAML sait écrire et relire à l'identique."""
    if isinstance(valeur, bool):
        return False
    return isinstance(valeur, (str, int, float, datetime))


def problemes_decisions(contenu) -> list[str]:
    """Problèmes de structure du contenu d'un fichier de décisions, tous à la fois."""
    if not isinstance(contenu, dict) or set(contenu) != {"meta", "decisions"}:
        return ["contenu attendu : les deux clés meta et decisions"]
    problemes = []
    meta = contenu["meta"]
    if not isinstance(meta, dict) or meta.get("structure") != STRUCTURE_DECISIONS:
        problemes.append(f"meta : structure {STRUCTURE_DECISIONS} attendue")
        meta = {}
    imports = meta.get("imports", [])
    if not isinstance(imports, list):
        problemes.append("meta : imports doit être une liste")
        imports = []
    precedent = None
    for rang, entree in enumerate(imports, start=1):
        ou = f"meta, import {rang}"
        if not isinstance(entree, dict) or set(entree) != set(CHAMPS_IMPORT):
            problemes.append(f"{ou} : champs attendus {', '.join(CHAMPS_IMPORT)}")
            continue
        instant = lire_instant(entree["date"])
        if instant is None:
            problemes.append(f"{ou} : date 'AAAA-MM-JJ HH:MM:SS' attendue")
        elif precedent is not None and instant <= precedent:
            problemes.append(f"{ou} : date antérieure ou égale à celle de l'import précédent")
        precedent = instant or precedent
    decisions = contenu["decisions"]
    if not isinstance(decisions, dict):
        return problemes + ["decisions : un dictionnaire par identifiant est attendu"]
    for identifiant, decision in decisions.items():
        ou = f"décision {identifiant}"
        if not isinstance(decision, dict) or set(decision) != set(CHAMPS_DECISION):
            problemes.append(f"{ou} : champs attendus {', '.join(CHAMPS_DECISION)}")
            continue
        sorte, choix = decision["sorte"], decision["decision"]
        if sorte not in SORTES:
            problemes.append(f"{ou} : sorte {sorte!r} inconnue")
        elif choix not in DECISIONS_PAR_SORTE[sorte]:
            problemes.append(f"{ou} : décision {choix!r} impossible pour la sorte {sorte}")
        if not isinstance(decision["empreinte"], str) or not decision["empreinte"]:
            problemes.append(f"{ou} : empreinte absente")
        if lire_instant(decision["decide_le"]) is None:
            problemes.append(f"{ou} : decide_le 'AAAA-MM-JJ HH:MM:SS' attendu")
        retenue = decision["valeur_retenue"]
        if choix == DECISION_CORRIGER and not _valeur_admise(retenue):
            problemes.append(f"{ou} : valeur retenue absente ou d'un type non admis")
        if choix != DECISION_CORRIGER and retenue is not None:
            problemes.append(f"{ou} : valeur retenue donnée sans la décision corriger")
    return problemes


def charger_decisions(chemin: str | Path | None = None) -> dict:
    """Décisions enregistrées : {"meta": {"structure", "imports"}, "decisions": {...}}.

    Fichier absent : aucune décision et aucun import. Un fichier mal formé est
    refusé, avec tous ses problèmes listés en une fois.
    """
    chemin = Path(chemin) if chemin is not None else chemin_decisions()
    if not chemin.exists():
        return _vide()
    with open(chemin, encoding="utf-8") as flux:
        contenu = yaml.safe_load(flux)
    problemes = problemes_decisions(contenu)
    if problemes:
        raise ValueError(f"Fichier de décisions invalide ({chemin})\n  " + "\n  ".join(problemes))
    return contenu


# --------------------------------------------------------------------------------
# Éléments à arbitrer


def _forme(valeur) -> str:
    """Texte d'une valeur pour l'empreinte, qui distingue un texte d'un nombre ou d'une date."""
    if valeur is None or (not isinstance(valeur, str) and pd.isna(valeur)):
        return "vide"
    if isinstance(valeur, str):
        return "texte:" + valeur
    if isinstance(valeur, (datetime, date, pd.Timestamp)):
        return "date:" + pd.Timestamp(valeur).isoformat()
    if isinstance(valeur, bool):
        return "booléen:" + str(valeur)
    try:
        return "nombre:" + repr(float(valeur))
    except (TypeError, ValueError):
        return "autre:" + str(valeur)


def _empreinte(parties) -> str:
    texte = "\x1f".join(parties)
    return hashlib.sha1(texte.encode("utf-8")).hexdigest()[:LONGUEUR_EMPREINTE]


def _affichee(valeur) -> str | None:
    """Texte d'une valeur pour le classeur (une date native s'écrit AAAA-MM-JJ)."""
    if isinstance(valeur, (datetime, date, pd.Timestamp)) and not pd.isna(valeur):
        return ecrire_date(valeur, "date native")
    return _valeur_constatee(valeur)


def _colonnes_locales(champ, donnees: pd.DataFrame) -> list[str] | None:
    """Colonnes de la table nommées par le champ, None si l'une n'en est pas une."""
    noms = [nom.strip() for nom in _texte(champ).split(",")]
    if not noms or any(nom not in _colonnes_de_donnees(donnees) for nom in noms):
        return None
    return noms


def _sorte(ligne, donnees: pd.DataFrame | None) -> str | None:
    """Sorte d'une anomalie, None si elle n'est pas soumise au métier."""
    if ligne.etape == ETAPE_FORMAT:
        return SORTE_CELLULE if ligne.issue == ISSUE_ANOMALIE else None
    if ligne.etape == ETAPE_DOUBLONS:
        return SORTE_LIGNE if ligne.issue in (ISSUE_ANOMALIE, ISSUE_A_ARBITRER) else None
    if ligne.etape == ETAPE_GESTION and ligne.issue == ISSUE_A_ARBITRER:
        sans_ligne = pd.isna(ligne.num_ligne_source)
        colonnes = None if donnees is None else _colonnes_locales(ligne.champ, donnees)
        if not sans_ligne and colonnes is not None and len(colonnes) == 1:
            return SORTE_CELLULE
        return SORTE_ECART
    return None


def elements_a_arbitrer(
    tables: dict[str, pd.DataFrame],
    anomalies: pd.DataFrame,
    decisions: dict | None = None,
) -> pd.DataFrame:
    """Éléments soumis au métier, avec l'état de leur décision.

    `tables` sont les tables sur lesquelles `anomalies` a été calculée (les tables
    nettoyées), `anomalies` la table rendue par `anomalies.reunir_anomalies` (ou
    `qualifier`), `decisions` le contenu rendu par `charger_decisions` (fichier
    par défaut).

    Éléments retenus, dans l'ordre de la table des anomalies : les anomalies de
    l'étape format (cellule); les lignes en doublon, anomalies ou à arbitrer
    (ligne); les écarts « à arbitrer » de l'étape gestion, cellule s'ils portent
    sur une seule colonne d'une ligne repérée, écart sinon. Les autres anomalies de
    gestion et les informations n'en font pas partie.

    La valeur affichée est lue dans la table par le lignage, y compris pour une
    colonne sensible : ce tableau porte des valeurs du client et ne doit pas être
    affiché. `empreinte` est le hachage de la valeur de la cellule (cellule), de
    la ligne hors lignage (ligne), ou des valeurs de l'écart (écart : valeurs des
    colonnes citées, lues dans la table; si l'une est d'une table liée, celles de
    la table et la valeur de la table des anomalies; sans lignage, la valeur et la
    clé de la table des anomalies).

    `etat_decision` : « sans décision », « décision caduque » (une décision existe
    pour cet identifiant, mais l'empreinte ou la sorte a changé), « décision
    valable », ou « ligne source introuvable » (lignage absent ou en double dans la
    table : l'élément ne peut pas être arbitré). Seuls les deux premiers états sont
    exportés dans le classeur.
    """
    manquantes = [c for c in COLONNES_ANOMALIES if c not in anomalies.columns]
    if manquantes:
        raise ValueError(
            f"Table des anomalies inattendue : colonnes absentes {', '.join(manquantes)}"
        )
    if decisions is None:
        decisions = charger_decisions()
    enregistrees = decisions["decisions"]
    index = {}
    for table, donnees in tables.items():
        donnees = donnees.reset_index(drop=True)
        index[table] = (_index_lignage(donnees), donnees)

    lignes = []
    for ligne in anomalies.itertuples(index=False):
        reperes, donnees = index.get(ligne.table, ({}, None))
        sorte = _sorte(ligne, donnees)
        if sorte is None:
            continue
        position = None
        if not pd.isna(ligne.num_ligne_source):
            position = reperes.get(
                _repere(ligne.fichier_source, ligne.feuille_source, ligne.num_ligne_source)
            )
        valeur, empreinte = ligne.valeur, None
        if position is not None:
            source = donnees.iloc[position]
            colonnes = _colonnes_locales(ligne.champ, donnees)
            if sorte == SORTE_LIGNE:
                empreinte = _empreinte(
                    f"{c}={_forme(source[c])}" for c in _colonnes_de_donnees(donnees)
                )
            elif colonnes is not None:
                empreinte = _empreinte(_forme(source[c]) for c in colonnes)
            else:
                # Écart qui cite une colonne d'une table liée : les colonnes de la
                # table, et la valeur de la table des anomalies.
                locales = [
                    nom.strip()
                    for nom in _texte(ligne.champ).split(",")
                    if nom.strip() in _colonnes_de_donnees(donnees)
                ]
                empreinte = _empreinte(
                    [*(_forme(source[c]) for c in locales), _forme(ligne.valeur)]
                )
            if colonnes is not None:
                textes = [_affichee(source[c]) for c in colonnes]
                valeur = (
                    None
                    if all(t is None for t in textes)
                    else SEPARATEUR.join(t or "" for t in textes)
                )
        elif sorte == SORTE_ECART and pd.isna(ligne.num_ligne_source):
            empreinte = _empreinte([_forme(ligne.valeur), _forme(ligne.cle)])
        if empreinte is None:
            etat = ETAT_INTROUVABLE
        else:
            existante = enregistrees.get(ligne.id_anomalie)
            if existante is None:
                etat = ETAT_SANS_DECISION
            elif existante["empreinte"] == empreinte and existante["sorte"] == sorte:
                etat = ETAT_VALABLE
            else:
                etat = ETAT_CADUQUE
        lignes.append(
            {
                "id_anomalie": ligne.id_anomalie,
                "sorte": sorte,
                "etat_decision": etat,
                "etape": ligne.etape,
                "table": ligne.table,
                "fichier_source": ligne.fichier_source,
                "feuille_source": ligne.feuille_source,
                "num_ligne_source": ligne.num_ligne_source,
                "cle": ligne.cle,
                "champ": ligne.champ,
                "valeur": valeur,
                "issue": ligne.issue,
                "regle": ligne.regle,
                "constat": ligne.constat,
                "detail": ligne.detail,
                "proposition": ligne.proposition,
                "empreinte": empreinte,
            }
        )
    elements = pd.DataFrame(lignes, columns=list(COLONNES_ELEMENTS), dtype=object)
    elements["num_ligne_source"] = pd.array(
        [None if pd.isna(n) else int(n) for n in elements["num_ligne_source"]], dtype="Int64"
    )
    return elements


COLONNE_DECISION_METIER = "decision_metier"


def ajouter_decisions(
    anomalies: pd.DataFrame, elements: pd.DataFrame, decisions: dict | None = None
) -> pd.DataFrame:
    """Table des anomalies avec, en dernière colonne, la décision du métier.

    `elements` est le tableau rendu par `elements_a_arbitrer` sur la même table des
    anomalies, `decisions` le contenu rendu par `charger_decisions` (fichier par
    défaut). `decision_metier` vaut, pour une anomalie dont l'élément a une décision
    enregistrée, la décision et sa date (« garder le 2026-10-20 »), suivies de
    « , caduque » si la donnée a changé depuis (l'élément sera proposé de nouveau
    au métier); elle est vide sinon. Une anomalie gardée reste une anomalie. La
    table reçue n'est pas modifiée.
    """
    if decisions is None:
        decisions = charger_decisions()
    enregistrees = decisions["decisions"]
    etats = dict(zip(elements["id_anomalie"], elements["etat_decision"]))
    textes = []
    for identifiant in anomalies["id_anomalie"]:
        etat = etats.get(identifiant)
        decision = enregistrees.get(identifiant)
        if decision is None or etat not in (ETAT_VALABLE, ETAT_CADUQUE):
            textes.append(None)
            continue
        texte = f"{decision['decision']} le {str(decision['decide_le'])[:10]}"
        textes.append(texte + (", caduque" if etat == ETAT_CADUQUE else ""))
    resultat = anomalies.copy()
    resultat[COLONNE_DECISION_METIER] = pd.Series(textes, index=anomalies.index, dtype=object)
    return resultat


# --------------------------------------------------------------------------------
# Classeur d'arbitrage : export


def _mode_emploi(periode_fin: str, elements: pd.DataFrame) -> list[str]:
    par_sorte = Counter(elements["sorte"])
    caduques = int((elements["etat_decision"] == ETAT_CADUQUE).sum())
    return [
        "Arbitrage des anomalies par le métier",
        f"Période de fin : {periode_fin}. Classeur généré le {date.today():%d/%m/%Y}.",
        f"Contenu : {len(elements)} lignes dans la feuille « {FEUILLE_ARBITRAGE} », dont "
        f"{par_sorte[SORTE_CELLULE]} cellules, {par_sorte[SORTE_LIGNE]} lignes et "
        f"{par_sorte[SORTE_ECART]} écarts; {caduques} lignes avaient une décision devenue "
        "caduque (la donnée a changé depuis).",
        REPERE_IMPORTS,
        "",
        "Ce qui est attendu",
        "Chaque ligne est un élément que les règles automatiques ne savent pas trancher. "
        "Pour chaque ligne, choisir une décision dans le menu déroulant de la colonne "
        "decision et, si besoin, écrire un commentaire. Une ligne laissée sans décision "
        "sera proposée de nouveau au prochain classeur.",
        "La colonne sorte dit quelles décisions sont possibles :",
        "cellule : corriger (la valeur est fausse, la valeur juste est écrite dans la "
        "colonne valeur_retenue) ou garder (la valeur est juste, ce qui est un cas "
        "normal; elle ne sera plus proposée tant qu'elle ne change pas).",
        "ligne : retirer (la ligne est retirée des données nettoyées) ou garder (la "
        "ligne est à conserver telle quelle).",
        "écart : garder (l'écart a été examiné; il ne sera plus proposé tant que les "
        "valeurs ne changent pas).",
        "La colonne valeur_retenue est préremplie avec la correction proposée par "
        "l'outil quand il en propose une : la vérifier avant de choisir corriger. Elle "
        "n'est lue qu'avec la décision corriger.",
        "Écrire la valeur retenue comme dans le fichier source. Une date peut s'écrire "
        "JJ/MM/AAAA, JJ.MM.AAAA, JJ-MM-AAAA, AAAA-MM-JJ ou AAAAMMJJ, un nombre avec une "
        "virgule ou un point décimal; une année sur deux chiffres est refusée. La "
        "valeur est ensuite écrite dans la forme attendue pour sa colonne.",
        "Une décision ne s'applique aux données nettoyées que si la valeur constatée "
        "n'a pas changé depuis : si la source est corrigée ou réexportée entre-temps, "
        "l'élément sera proposé de nouveau.",
        "",
        "Confidentialité",
        "Ce classeur porte des valeurs du client, y compris celles des colonnes "
        "sensibles. Il reste dans le dossier data/decisions et n'est ni transmis ni "
        "recopié ailleurs.",
        "",
        "Règles à respecter pour que le classeur puisse être réimporté",
        "Ne modifier que les colonnes decision, valeur_retenue et commentaire (les "
        "autres cellules sont verrouillées).",
        "Ne pas renommer, déplacer ni supprimer de feuilles, de colonnes ou de lignes, "
        "et ne pas ajouter de lignes.",
        "Choisir les décisions dans le menu déroulant : une décision qui ne convient "
        "pas à la sorte de la ligne est refusée à l'import.",
        "Un classeur ne s'importe qu'une fois : après son import, ou celui d'un autre "
        "classeur, il est à supprimer puis à réexporter.",
        "Avant l'import : enregistrer au format .xlsx, sans changer le nom du fichier, "
        "puis fermer le classeur dans Excel.",
        "Le filtre de la ligne d'en-tête peut être utilisé librement.",
    ]


def exporter_arbitrage(
    elements: pd.DataFrame,
    periode_fin: str,
    chemin: str | Path | None = None,
    *,
    nb_imports: int,
    ecraser: bool = False,
) -> Path:
    """Écrit le classeur d'arbitrage et rend son chemin.

    `elements` est le tableau rendu par `elements_a_arbitrer`; seuls les éléments
    sans décision ou à la décision caduque sont écrits. `nb_imports` est le nombre
    d'imports déjà enregistrés (`len(charger_decisions()["meta"]["imports"])`) : il
    est inscrit dans le classeur, et l'import refusera le classeur si un autre
    import a eu lieu depuis. Sans `ecraser=True`, refuse de remplacer un classeur
    existant : un arbitrage y est peut-être en cours.
    """
    if isinstance(nb_imports, bool) or not isinstance(nb_imports, int) or nb_imports < 0:
        raise ValueError(f"nb_imports {nb_imports!r} : un entier positif ou nul est attendu")
    manquantes = [c for c in COLONNES_ELEMENTS if c not in elements.columns]
    if manquantes:
        raise ValueError(f"Éléments inattendus : colonnes absentes {', '.join(manquantes)}")
    chemin = Path(chemin) if chemin is not None else chemin_arbitrage(periode_fin)
    if chemin.exists() and not ecraser:
        raise FileExistsError(f"{chemin} existe déjà : export refusé (ecraser=True pour forcer)")
    a_arbitrer = elements[elements["etat_decision"].isin(ETATS_A_ARBITRER)]

    classeur = Workbook()
    notice = classeur.active
    notice.title = FEUILLE_MODE_EMPLOI
    notice.column_dimensions["A"].width = 120
    for numero, texte in enumerate(_mode_emploi(periode_fin, a_arbitrer), start=1):
        cellule = notice.cell(row=numero, column=1)
        _ecrire_texte(cellule, texte)
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
        if texte == REPERE_IMPORTS:
            notice.cell(row=numero, column=2, value=nb_imports)
        if numero == 1 or texte.startswith(("Ce qui", "Confidentialité", "Règles à")):
            cellule.font = Font(bold=True)
    notice.protection.sheet = True

    feuille = classeur.create_sheet(FEUILLE_ARBITRAGE)
    for indice, nom in enumerate(COLONNES_ARBITRAGE, start=1):
        cellule = feuille.cell(row=1, column=indice, value=nom)
        cellule.font = Font(bold=True)
        lettre = cellule.column_letter
        feuille.column_dimensions[lettre].width = LARGEURS_ARBITRAGE[nom]
        if nom in COLONNES_TECHNIQUES:
            feuille.column_dimensions[lettre].hidden = True
    menus = {
        sorte: _menu(feuille, choix) for sorte, choix in DECISIONS_PAR_SORTE.items()
    }
    colonne_decision = COLONNES_ARBITRAGE.index("decision") + 1
    for numero, element in enumerate(a_arbitrer.to_dict("records"), start=2):
        proposition = element["proposition"]
        prerempli = (
            proposition
            if element["sorte"] == SORTE_CELLULE
            and not _est_vide(proposition)
            and proposition != VALEUR_MASQUEE
            else None
        )
        valeurs = {
            **element,
            "decision": None,
            "valeur_retenue": prerempli,
            "commentaire": None,
        }
        for indice, nom in enumerate(COLONNES_ARBITRAGE, start=1):
            cellule = feuille.cell(row=numero, column=indice)
            valeur = valeurs.get(nom)
            if nom == "num_ligne_source":
                cellule.value = None if pd.isna(valeur) else int(valeur)
            else:
                _ecrire_texte(cellule, None if _est_vide(valeur) else _texte(valeur))
            if nom in MODIFIABLES_ARBITRAGE:
                cellule.protection = Protection(locked=False)
        menus[element["sorte"]].add(feuille.cell(row=numero, column=colonne_decision))
    fin = max(len(a_arbitrer) + 1, 2)
    derniere = get_column_letter(len(COLONNES_ARBITRAGE))
    feuille.freeze_panes = "B2"
    feuille.auto_filter.ref = f"A1:{derniere}{fin}"
    feuille.protection.sheet = True
    feuille.protection.autoFilter = False
    feuille.protection.formatColumns = False

    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin


# --------------------------------------------------------------------------------
# Classeur d'arbitrage : import


class Arbitrage(NamedTuple):
    """Décisions lues dans un classeur d'arbitrage, contrôlées mais pas encore enregistrées."""

    decisions: pd.DataFrame
    avertissements: list[str]
    classeur: Path
    repere: int


def _lire_feuille(classeur, nom: str, colonnes: tuple[str, ...]) -> list[tuple[int, dict]]:
    """Lignes non vides d'une feuille, avec leur numéro de ligne Excel."""
    if nom not in classeur.sheetnames:
        raise ValueError(
            f"feuille « {nom} » absente (feuille renommée : classeur à réexporter)"
        )
    lignes = classeur[nom].iter_rows(values_only=True)
    entete = tuple(next(lignes, ()))
    while entete and entete[-1] is None:
        entete = entete[:-1]
    if entete != colonnes:
        raise ValueError(
            f"feuille « {nom} » : en-tête modifié, attendu {', '.join(colonnes)} "
            "(classeur exporté par une version antérieure : à réexporter)"
        )
    resultat = []
    for numero, valeurs in enumerate(lignes, start=2):
        valeurs = tuple(valeurs[: len(colonnes)]) + (None,) * (len(colonnes) - len(valeurs))
        if all(_est_vide(valeur) for valeur in valeurs):
            continue
        resultat.append((numero, dict(zip(colonnes, valeurs))))
    return resultat


def _lire_repere(classeur) -> int:
    """Nombre d'imports déjà enregistrés quand le classeur a été exporté."""
    if FEUILLE_MODE_EMPLOI not in classeur.sheetnames:
        raise ValueError(f"feuille « {FEUILLE_MODE_EMPLOI} » absente")
    for ligne in classeur[FEUILLE_MODE_EMPLOI].iter_rows(max_col=2, values_only=True):
        if ligne and ligne[0] == REPERE_IMPORTS:
            nombre = ligne[1] if len(ligne) > 1 else None
            if isinstance(nombre, int) and not isinstance(nombre, bool) and nombre >= 0:
                return nombre
            break
    raise ValueError(
        f"feuille « {FEUILLE_MODE_EMPLOI} » : repère des imports absent ou modifié, "
        "classeur à réexporter"
    )


def repere_arbitrage(chemin: str | Path) -> int:
    """Nombre d'imports déjà enregistrés quand le classeur a été exporté.

    Comparé à `len(charger_decisions()["meta"]["imports"])`, il dit si le classeur
    est encore importable (égal) ou déjà importé, ou périmé (inférieur).
    """
    classeur = load_workbook(Path(chemin), read_only=True, data_only=True)
    try:
        return _lire_repere(classeur)
    finally:
        classeur.close()


def _choix(valeur) -> str:
    """Texte d'une décision, normalisé (NFC, espaces en bord ôtés)."""
    if valeur is None:
        return ""
    return unicodedata.normalize("NFC", str(valeur)).strip()


def _numero(valeur) -> int | None:
    if valeur is None or isinstance(valeur, bool):
        return None
    try:
        nombre = float(valeur)
    except (TypeError, ValueError):
        return None
    return int(nombre) if nombre.is_integer() else None


def lire_valeur_retenue(valeur, nature: str):
    """Valeur retenue écrite dans la représentation de la nature de sa colonne.

    Rend (valeur, None), ou (None, raison du refus). Les espaces en bord sont
    ôtés. Une colonne de texte (ou à la nature non lisible) reçoit le texte tel
    quel. Pour une nature de date, d'heure ou de nombre, une valeur déjà conforme
    est gardée, une autre est lue dans les formes élargies de la même famille,
    comme au nettoyage (`nettoyage._convertir`, sans date de coupure : une année
    sur deux chiffres est refusée), puis écrite dans la représentation de la
    nature.
    """
    if isinstance(valeur, str):
        valeur = valeur.strip()
    if _est_vide(valeur):
        return None, "valeur retenue vide"
    if nature not in NATURES_LISIBLES:
        return _texte(valeur), None
    if nature not in ("date native", "nombre natif") and isinstance(valeur, str):
        if len(_interpretations(pd.Series([valeur], dtype=object))[nature]):
            return valeur, None
    resultat = _convertir(valeur, nature, None)
    if isinstance(resultat, str):
        return None, resultat
    convertie = resultat[0]
    if isinstance(convertie, pd.Timestamp):
        convertie = convertie.to_pydatetime()
    return convertie, None


def importer_arbitrage(
    chemin: str | Path, dictionnaire: dict, *, nb_imports: int
) -> Arbitrage:
    """Lit un classeur d'arbitrage et contrôle chaque décision, sans rien enregistrer.

    `dictionnaire` est le dictionnaire enregistré (nature des colonnes),
    `nb_imports` le nombre d'imports déjà enregistrés
    (`len(charger_decisions()["meta"]["imports"])`). Un classeur exporté avant un
    autre import est refusé (périmé). Seules les lignes qui portent une décision
    sont rendues; une ligne sans décision n'en est pas une.

    Tous les problèmes bloquants sont listés en une fois dans une ValueError, avec
    le numéro de ligne Excel : décision hors menu, décision qui ne convient pas à
    la sorte, « corriger » sans valeur retenue, valeur retenue illisible dans la
    nature de sa colonne, identifiant absent ou en double, colonne inconnue du
    dictionnaire, deux corrections différentes d'une même cellule. Les
    avertissements (valeur retenue ou commentaire ignorés) ne bloquent pas.
    """
    chemin = Path(chemin)
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    try:
        try:
            lignes = _lire_feuille(classeur, FEUILLE_ARBITRAGE, COLONNES_ARBITRAGE)
            repere = _lire_repere(classeur)
        except ValueError as erreur:
            raise ValueError(f"Classeur d'arbitrage invalide ({chemin})\n  {erreur}") from None
    finally:
        classeur.close()
    if repere != nb_imports:
        raise ValueError(
            f"Classeur d'arbitrage périmé ({chemin}) : exporté après {repere} imports, "
            f"le fichier des décisions en compte {nb_imports}. Ce classeur a déjà été "
            "importé, ou un autre l'a été depuis son export : à supprimer puis à réexporter."
        )

    problemes, avertissements, decisions = [], [], []
    vus: dict[str, int] = {}
    corrections: dict[tuple, tuple[int, object]] = {}
    for numero, ligne in lignes:
        ou = f"ligne {numero}"
        identifiant = _choix(ligne["id_anomalie"])
        sorte = _choix(ligne["sorte"])
        choix = _choix(ligne["decision"])
        if not identifiant:
            problemes.append(f"{ou} : identifiant absent (colonne id_anomalie masquée)")
            continue
        if identifiant in vus:
            problemes.append(
                f"{ou} : identifiant {identifiant} déjà présent ligne {vus[identifiant]}"
            )
            continue
        vus[identifiant] = numero
        if sorte not in SORTES:
            problemes.append(f"{ou} : sorte {sorte!r} inconnue")
            continue
        if not choix:
            if not _est_vide(ligne["commentaire"]):
                avertissements.append(f"{ou} : commentaire sans décision, ignoré")
            continue
        if choix not in DECISIONS:
            problemes.append(f"{ou} : décision {choix!r} hors du menu ({', '.join(DECISIONS)})")
            continue
        if choix not in DECISIONS_PAR_SORTE[sorte]:
            problemes.append(
                f"{ou} : décision {choix} impossible pour un élément de sorte {sorte} "
                f"(possibles : {', '.join(DECISIONS_PAR_SORTE[sorte])})"
            )
            continue
        table, champ = _texte(ligne["table"]), _texte(ligne["champ"])
        nature, retenue = None, None
        if choix == DECISION_CORRIGER:
            if table not in dictionnaire or champ not in dictionnaire[table]:
                problemes.append(f"{ou} : colonne {table}.{champ} inconnue du dictionnaire")
                continue
            nature = dictionnaire[table][champ]["nature"]["regle"]
            if _est_vide(ligne["valeur_retenue"]):
                problemes.append(f"{ou} : décision corriger sans valeur retenue")
                continue
            retenue, raison = lire_valeur_retenue(ligne["valeur_retenue"], nature)
            if raison is not None:
                problemes.append(
                    f"{ou} : valeur retenue illisible dans la nature « {nature} » de "
                    f"{table}.{champ} ({raison})"
                )
                continue
            cellule = (
                table,
                _texte(ligne["fichier_source"]),
                _texte(ligne["feuille_source"]),
                _numero(ligne["num_ligne_source"]),
                champ,
            )
            if cellule in corrections and corrections[cellule][1] != retenue:
                problemes.append(
                    f"{ou} : la même cellule reçoit une autre valeur ligne "
                    f"{corrections[cellule][0]}"
                )
                continue
            corrections[cellule] = (numero, retenue)
        elif not _est_vide(ligne["valeur_retenue"]) and _texte(ligne["valeur_retenue"]) != _texte(
            ligne["proposition"]
        ):
            avertissements.append(f"{ou} : valeur retenue ignorée avec la décision {choix}")
        commentaire = ligne["commentaire"]
        decisions.append(
            {
                "id_anomalie": identifiant,
                "sorte": sorte,
                "table": table,
                "fichier_source": _texte(ligne["fichier_source"]) or None,
                "feuille_source": _texte(ligne["feuille_source"]) or None,
                "num_ligne_source": _numero(ligne["num_ligne_source"]),
                "champ": champ or None,
                "constat": None if _est_vide(ligne["constat"]) else _texte(ligne["constat"]),
                "empreinte": _texte(ligne["empreinte"]),
                "decision": choix,
                "valeur_retenue": retenue,
                "nature": nature,
                "commentaire": None if _est_vide(commentaire) else str(commentaire).strip(),
                "ligne_excel": numero,
            }
        )
        if not decisions[-1]["empreinte"]:
            problemes.append(f"{ou} : empreinte absente (colonne empreinte masquée)")
    if problemes:
        raise ValueError(
            f"Classeur d'arbitrage invalide ({chemin}), {len(problemes)} problèmes\n  "
            + "\n  ".join(problemes)
        )
    return Arbitrage(
        pd.DataFrame(decisions, columns=list(COLONNES_DECISIONS_LUES), dtype=object),
        avertissements,
        chemin,
        repere,
    )


# --------------------------------------------------------------------------------
# Enregistrement


def ecrire_decisions(
    arbitrage: Arbitrage,
    periode_fin: str,
    chemin: str | Path | None = None,
) -> tuple[Path, dict | None]:
    """Enregistre les décisions d'un classeur importé et rend le chemin et l'entrée du journal.

    Une décision remplace celle du même identifiant (décision caduque). Le
    journal des imports reçoit une entrée (date, classeur, période, nombre de
    décisions, de décisions remplacées, et par décision). Rien n'est écrit si le
    classeur ne porte aucune décision (entrée None) : le classeur reste alors
    importable.

    Refusé si un import a été enregistré depuis la lecture du classeur, ou si la
    date de l'import n'est pas postérieure à celle du précédent. Le fichier en
    place est d'abord copié à l'identique dans data/config/sauvegardes/, puis le
    nouveau contenu est écrit dans un fichier provisoire, relu et comparé, avant
    de remplacer le fichier en place.
    """
    chemin = Path(chemin) if chemin is not None else chemin_decisions()
    contenu = charger_decisions(chemin)
    imports = contenu["meta"]["imports"]
    if len(imports) != arbitrage.repere:
        raise ValueError(
            f"Classeur {arbitrage.classeur.name} périmé : {len(imports)} imports enregistrés, "
            f"{arbitrage.repere} à son export. Rien n'est écrit."
        )
    if arbitrage.decisions.empty:
        return chemin, None
    instant = horodater()
    if imports and lire_instant(instant) <= lire_instant(imports[-1]["date"]):
        raise ValueError(
            f"{chemin} : l'import ({instant}) doit être postérieur au précédent "
            f"({imports[-1]['date']}); relancer l'écriture"
        )
    decisions = dict(contenu["decisions"])
    remplacees = 0
    for ligne in arbitrage.decisions.to_dict("records"):
        identifiant = ligne["id_anomalie"]
        remplacees += identifiant in decisions
        decisions[identifiant] = {
            **{c: ligne[c] for c in CHAMPS_DECISION if c in ligne},
            "decide_le": instant,
            "classeur": arbitrage.classeur.name,
        }
    nombres = Counter(arbitrage.decisions["decision"])
    entree = {
        "date": instant,
        "classeur": arbitrage.classeur.name,
        "periode_fin": periode_fin,
        "decisions": len(arbitrage.decisions),
        "remplacees": int(remplacees),
        **{choix: int(nombres[choix]) for choix in DECISIONS},
    }
    nouveau = {
        "meta": {"structure": STRUCTURE_DECISIONS, "imports": [*imports, entree]},
        "decisions": decisions,
    }
    problemes = problemes_decisions(nouveau)
    if problemes:
        raise ValueError("Décisions invalides, rien n'est écrit\n  " + "\n  ".join(problemes))

    chemin.parent.mkdir(parents=True, exist_ok=True)
    if chemin.exists():
        _sauvegarder(chemin, periode_fin)
    provisoire = chemin.with_name(chemin.name + ".tmp")
    try:
        with open(provisoire, "w", encoding="utf-8", newline="\n") as flux:
            yaml.safe_dump(
                nouveau, flux, allow_unicode=True, sort_keys=False,
                default_flow_style=False, width=1000,
            )
        if charger_decisions(provisoire) != nouveau:
            raise RuntimeError(f"Relecture de {provisoire} différente du contenu écrit")
        os.replace(provisoire, chemin)
    finally:
        provisoire.unlink(missing_ok=True)
    return chemin, entree

