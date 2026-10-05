"""Chargement des fichiers bruts sans modification, avec lignage.

Les formats (encodage, séparateur, feuille, ligne d'en-tête) sont lus dans la
configuration des sources, portée par le périmètre. Ce module ne contient aucun
nom propre à un client.

Chaque ligne chargée conserve
- fichier_source, le nom du fichier d'origine
- feuille_source, le nom de la feuille pour un fichier Excel
- num_ligne_source, le numéro de ligne physique, la première ligne du fichier étant la ligne 1
- periode_export, la période du fichier
- date_chargement, l'horodatage du chargement

Les valeurs des fichiers texte sont conservées telles quelles, sous forme de texte.
Les valeurs des fichiers Excel conservent le type lu dans la cellule.
"""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pandas as pd

from frigeo.perimetre import Perimetre, tables_du_flux


def liste_tables(config: dict) -> list[str]:
    """Noms des tables déclarées dans la configuration, dans l'ordre des flux."""
    return [t for flux in config["flux"].values() for t in tables_du_flux(flux)]


def _localiser(config: dict, table: str) -> tuple[str, dict]:
    """Flux d'origine et options d'une table."""
    for nom, flux in config["flux"].items():
        tables = tables_du_flux(flux)
        if table in tables:
            return nom, tables[table]
    raise RuntimeError(f"Table {table} absente de la configuration des sources")


def lire_texte(chemin: Path, encodage: str, sep: str, ligne_entete: int = 1):
    """Renvoie l'en-tête et la liste des (numéro de ligne physique, valeurs)."""
    try:
        with open(chemin, encoding=encodage, newline="") as fichier:
            lecteur = csv.reader(fichier, delimiter=sep)
            entete = None
            for valeurs in lecteur:
                if lecteur.line_num >= ligne_entete:
                    entete = valeurs
                    break
            if entete is None:
                raise RuntimeError(f"Fichier vide ou en-tête introuvable {chemin.name}")
            precedent = lecteur.line_num
            lignes = []
            for valeurs in lecteur:
                lignes.append((precedent + 1, valeurs))
                precedent = lecteur.line_num
    except UnicodeDecodeError as erreur:
        raise RuntimeError(
            f"Encodage inattendu pour {chemin.name}, attendu {encodage}, octet {erreur.start}"
        ) from erreur
    return entete, lignes


def lire_excel(chemin: Path, feuille: str, ligne_entete: int = 1):
    """Renvoie l'en-tête et la liste des (numéro de ligne, valeurs) d'une feuille."""
    classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    try:
        if feuille not in classeur.sheetnames:
            raise RuntimeError(f"Feuille {feuille} absente de {chemin.name}")
        rangees = list(classeur[feuille].iter_rows(values_only=True))
    finally:
        classeur.close()
    if len(rangees) < ligne_entete:
        raise RuntimeError(f"Feuille {feuille} vide dans {chemin.name}")
    entete = [None if c is None else str(c) for c in rangees[ligne_entete - 1]]
    lignes = [
        (numero, list(valeurs))
        for numero, valeurs in enumerate(rangees[ligne_entete:], start=ligne_entete + 1)
    ]
    return entete, lignes


def _assembler(entete, lignes, fichier, feuille=None, sep=None) -> pd.DataFrame:
    n = len(entete)
    donnees = [v[:n] + [None] * (n - len(v)) for _, v in lignes]
    table = pd.DataFrame(donnees, columns=entete, dtype=object)
    if sep is not None:
        table["nb_champs_source"] = [len(v) for _, v in lignes]
        table["champs_excedentaires"] = [
            sep.join(v[n:]) if len(v) > n else None for _, v in lignes
        ]
    table.insert(0, "periode_export", fichier["periode"])
    table.insert(0, "num_ligne_source", [numero for numero, _ in lignes])
    table.insert(0, "feuille_source", feuille)
    table.insert(0, "fichier_source", fichier["fichier"])
    return table


def charger_table(perimetre: Perimetre, table: str, date_chargement=None) -> pd.DataFrame:
    """Charge et assemble tous les exemplaires d'une table du périmètre."""
    date_chargement = date_chargement or pd.Timestamp.now().floor("s")
    flux, options = _localiser(perimetre.config, table)
    spec = perimetre.config["flux"][flux]
    fichiers = [f for f in perimetre.fichiers if f["flux"] == flux]
    if not fichiers:
        raise RuntimeError(f"Aucun fichier retenu dans le périmètre pour la table {table}")
    morceaux = []
    for f in fichiers:
        if spec["format"] == "texte":
            entete, lignes = lire_texte(
                f["chemin"], spec["encodage"], spec["separateur"], spec["ligne_entete"]
            )
            morceaux.append(_assembler(entete, lignes, f, sep=spec["separateur"]))
        else:
            feuille = options["feuille"]
            entete, lignes = lire_excel(f["chemin"], feuille, spec["ligne_entete"])
            morceaux.append(_assembler(entete, lignes, f, feuille=feuille))
    resultat = pd.concat(morceaux, ignore_index=True)
    resultat["date_chargement"] = date_chargement
    return resultat


def charger_tout(perimetre: Perimetre) -> dict[str, pd.DataFrame]:
    """Charge toutes les tables avec le même horodatage."""
    horodatage = pd.Timestamp.now().floor("s")
    return {
        t: charger_table(perimetre, t, horodatage)
        for t in liste_tables(perimetre.config)
    }


def bilan_chargement(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Nombre de lignes chargées par table, fichier et feuille."""
    lignes = []
    for table, donnees in tables.items():
        comptes = donnees.groupby(
            ["fichier_source", "feuille_source"], dropna=False
        ).size()
        for (fichier, feuille), nombre in comptes.items():
            lignes.append(
                {
                    "table": table,
                    "fichier_source": fichier,
                    "feuille_source": feuille if isinstance(feuille, str) else None,
                    "lignes_chargees": int(nombre),
                }
            )
    return pd.DataFrame(lignes)