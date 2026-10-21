"""Inventaire des fichiers du périmètre, avant et après chargement (étape 1).

Les formats (encodage, séparateur, ligne d'en-tête, feuilles) sont lus dans
config/sources.yaml, porté par le périmètre. Ce module ne contient aucun nom propre
à un client, et ses résultats ne contiennent aucune valeur des données : seulement
des noms de fichiers, de feuilles, de tables et de colonnes, et des décomptes.

- `sonder_fichiers` lit chaque fichier retenu sans passer par le chargement :
  marque d'ordre des octets, fins de ligne, décodage, nombre de champs par
  enregistrement, nombre de lignes par feuille, en-têtes;
- `entetes_par_source` regroupe les fichiers d'un même flux (et d'une même
  feuille) par en-tête, pour repérer un export dont les colonnes ont changé;
- `comparer_lignes` rapproche les lignes chargées des lignes attendues d'après le
  sondage, pour repérer des lignes perdues;
- `encodage_suspect` décompte, par table et par colonne, les cellules que la règle
  `encodage` de l'étape 3 signale (voir `controles.constats_encodage`), avec le
  mécanisme probable déduit du format déclaré de la source.
"""

from __future__ import annotations

import csv
import io
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd

from .controles import CONSTAT_DECODAGE, CONSTAT_REMPLACEMENT, constats_encodage
from .nettoyage import reparer_encodage
from .perimetre import Perimetre, tables_du_flux
from .profilage import COLONNES_LIGNAGE

# Nombre de longueurs d'enregistrement (nombre de champs) gardées par fichier texte.
LONGUEURS_AFFICHEES = 4

MECANISME_DOUBLE = "texte déjà mal décodé dans la source (double encodage)"
MECANISME_EXCEL = "texte mal décodé avant son enregistrement dans le classeur"
MECANISME_PERDU = "caractère déjà perdu dans la source"

COLONNES_TEXTE = [
    "fichier",
    "flux",
    "encodage",
    "separateur",
    "ligne_entete",
    "bom_utf8",
    "crlf",
    "lf_seul",
    "decodage",
    "colonnes_entete",
    "enregistrements",
    "longueurs",
]
COLONNES_EXCEL = [
    "fichier",
    "flux",
    "feuille",
    "declaree",
    "ligne_entete",
    "lignes",
    "colonnes",
    "lignes_donnees",
]
COLONNES_ENCODAGE = [
    "table",
    "colonne",
    "constat",
    "mecanisme",
    "cellules",
    "valeurs_distinctes",
    "reparables",
    "reparees_presentes_ailleurs",
]


@dataclass
class Inventaire:
    """Résultat de `sonder_fichiers` : un tableau par format et les en-têtes lus.

    `entetes` a une ligne par fichier (et par feuille pour Excel) : flux, feuille
    (None pour un fichier texte), fichier, en-tête (tuple des noms de colonnes, None
    si le fichier n'a pas pu être décodé).
    """

    texte: pd.DataFrame
    excel: pd.DataFrame
    entetes: pd.DataFrame


def sonder_texte(chemin: Path, encodage: str, sep: str, ligne_entete: int = 1):
    """Structure d'un fichier texte et son en-tête, lus comme au chargement.

    Rend (mesures, en-tête). Les mesures comptent les fins de ligne sur les octets,
    puis, si le décodage réussit, les colonnes de l'en-tête, les enregistrements qui
    le suivent (un enregistrement peut tenir sur plusieurs lignes physiques) et les
    longueurs d'enregistrement les plus fréquentes, {nombre de champs: nombre
    d'enregistrements}. L'en-tête vaut None si le décodage échoue.
    """
    brut = Path(chemin).read_bytes()
    mesures = {
        "bom_utf8": brut.startswith(b"\xef\xbb\xbf"),
        "crlf": brut.count(b"\r\n"),
        "lf_seul": brut.count(b"\n") - brut.count(b"\r\n"),
        "decodage": "ok",
        "colonnes_entete": None,
        "enregistrements": None,
        "longueurs": None,
    }
    try:
        texte = brut.decode(encodage)
    except UnicodeDecodeError as erreur:
        mesures["decodage"] = f"échec à l'octet {erreur.start}"
        return mesures, None
    lecteur = csv.reader(io.StringIO(texte, newline=""), delimiter=sep)
    entete = None
    for valeurs in lecteur:
        if lecteur.line_num >= ligne_entete:
            entete = valeurs
            break
    if entete is None:
        mesures["decodage"] = "en-tête introuvable"
        return mesures, None
    longueurs = Counter(len(valeurs) for valeurs in lecteur)
    mesures.update(
        colonnes_entete=len(entete),
        enregistrements=sum(longueurs.values()),
        longueurs=dict(longueurs.most_common(LONGUEURS_AFFICHEES)),
    )
    return mesures, tuple(entete)


def _lignes_feuille(feuille) -> int:
    """Nombre de lignes d'une feuille d'après ses dimensions, ou en la parcourant
    quand le classeur ne les donne pas."""
    if feuille.max_row is not None:
        return feuille.max_row
    return sum(1 for _ in feuille.iter_rows(values_only=True))


def sonder_excel(chemin: Path, feuilles_declarees, ligne_entete: int = 1) -> list[dict]:
    """Structure de chaque feuille d'un classeur, déclarée ou non.

    Rend une liste de dictionnaires (feuille, declaree, lignes, colonnes,
    lignes_donnees, entete). L'en-tête est lu à la ligne d'en-tête du flux.
    """
    resultat = []
    classeur = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    try:
        for feuille in classeur.worksheets:
            rangee = next(
                feuille.iter_rows(
                    min_row=ligne_entete, max_row=ligne_entete, values_only=True
                ),
                (),
            )
            lignes = _lignes_feuille(feuille)
            resultat.append(
                {
                    "feuille": feuille.title,
                    "declaree": feuille.title in feuilles_declarees,
                    "lignes": lignes,
                    "colonnes": feuille.max_column,
                    "lignes_donnees": max(lignes - ligne_entete, 0),
                    "entete": tuple(None if c is None else str(c) for c in rangee),
                }
            )
    finally:
        classeur.close()
    return resultat


def sonder_fichiers(perimetre: Perimetre) -> Inventaire:
    """Sonde tous les fichiers du périmètre selon leur format déclaré."""
    lignes_texte, lignes_excel, entetes = [], [], []
    for fichier in perimetre.fichiers:
        spec = perimetre.config["flux"][fichier["flux"]]
        ligne_entete = spec["ligne_entete"]
        if spec["format"] == "texte":
            mesures, entete = sonder_texte(
                fichier["chemin"], spec["encodage"], spec["separateur"], ligne_entete
            )
            lignes_texte.append(
                {
                    "fichier": fichier["chemin_relatif"],
                    "flux": fichier["flux"],
                    "encodage": spec["encodage"],
                    "separateur": spec["separateur"],
                    "ligne_entete": ligne_entete,
                    **mesures,
                }
            )
            entetes.append(
                {
                    "flux": fichier["flux"],
                    "feuille": None,
                    "fichier": fichier["chemin_relatif"],
                    "entete": entete,
                }
            )
        else:
            declarees = {
                options["feuille"] for options in tables_du_flux(spec).values()
            }
            for feuille in sonder_excel(fichier["chemin"], declarees, ligne_entete):
                entete = feuille.pop("entete")
                lignes_excel.append(
                    {
                        "fichier": fichier["chemin_relatif"],
                        "flux": fichier["flux"],
                        "ligne_entete": ligne_entete,
                        **feuille,
                    }
                )
                if feuille["declaree"]:
                    entetes.append(
                        {
                            "flux": fichier["flux"],
                            "feuille": feuille["feuille"],
                            "fichier": fichier["chemin_relatif"],
                            "entete": entete,
                        }
                    )
    return Inventaire(
        texte=pd.DataFrame(lignes_texte, columns=COLONNES_TEXTE),
        excel=pd.DataFrame(lignes_excel, columns=COLONNES_EXCEL),
        entetes=pd.DataFrame(entetes, columns=["flux", "feuille", "fichier", "entete"]),
    )


def entetes_par_source(inventaire: Inventaire) -> pd.DataFrame:
    """En-têtes distincts par flux et par feuille déclarée.

    Une ligne par en-tête distinct : flux, feuille, nombre de fichiers, fichiers,
    colonnes (noms séparés par « | », None si le fichier n'a pas pu être lu) et
    `divergent`, vrai quand le flux ou la feuille a plusieurs en-têtes distincts.
    """
    colonnes = ["flux", "feuille", "fichiers", "nb_fichiers", "colonnes", "divergent"]
    if inventaire.entetes.empty:
        return pd.DataFrame(columns=colonnes)
    lignes = []
    groupes = inventaire.entetes.groupby(["flux", "feuille"], dropna=False, sort=False)
    for (flux, feuille), groupe in groupes:
        par_entete: dict = {}
        for fichier, entete in zip(groupe["fichier"], groupe["entete"]):
            par_entete.setdefault(entete, []).append(fichier)
        for entete, fichiers in par_entete.items():
            lignes.append(
                {
                    "flux": flux,
                    "feuille": feuille if isinstance(feuille, str) else None,
                    "fichiers": ", ".join(Path(f).name for f in fichiers),
                    "nb_fichiers": len(fichiers),
                    "colonnes": None
                    if entete is None
                    else "|".join("" if c is None else str(c) for c in entete),
                    "divergent": len(par_entete) > 1,
                }
            )
    return pd.DataFrame(lignes, columns=colonnes)


def comparer_lignes(bilan: pd.DataFrame, inventaire: Inventaire) -> pd.DataFrame:
    """Lignes chargées et lignes attendues d'après le sondage, par fichier et feuille.

    `bilan` est le résultat de `chargement.bilan_chargement`. L'attendu est le
    nombre d'enregistrements après l'en-tête pour un fichier texte, le nombre de
    lignes de la feuille moins la ligne d'en-tête et celles qui la précèdent pour
    Excel. `ecart` vaut lignes chargées moins attendu.
    """
    attendu_texte = {
        Path(f).name: n
        for f, n in zip(
            inventaire.texte["fichier"], inventaire.texte["enregistrements"]
        )
    }
    attendu_excel = {
        (Path(f).name, s): n
        for f, s, n in zip(
            inventaire.excel["fichier"],
            inventaire.excel["feuille"],
            inventaire.excel["lignes_donnees"],
        )
    }
    comparaison = bilan.copy()
    comparaison["attendu"] = [
        attendu_excel.get((f, s)) if isinstance(s, str) else attendu_texte.get(f)
        for f, s in zip(comparaison["fichier_source"], comparaison["feuille_source"])
    ]
    comparaison["attendu"] = pd.to_numeric(comparaison["attendu"]).astype("Int64")
    comparaison["ecart"] = comparaison["lignes_chargees"] - comparaison["attendu"]
    return comparaison


def _mecanisme(spec: dict, constat: str) -> str:
    """Mécanisme probable d'une cellule suspecte, d'après le format déclaré.

    Le chargement décode strictement un fichier texte : une séquence suspecte dans
    un fichier lu en UTF-8 y figurait déjà sous cette forme (texte UTF-8 relu dans un
    encodage d'un octet puis réenregistré en UTF-8); dans un fichier lu dans un
    encodage d'un octet, elle vient d'octets UTF-8 présents dans ce fichier. Un
    classeur Excel enregistre du texte : la séquence y était avant l'enregistrement.
    """
    if constat == CONSTAT_REMPLACEMENT:
        return MECANISME_PERDU
    if spec["format"] != "texte":
        return MECANISME_EXCEL
    encodage = spec["encodage"].lower().replace("_", "-")
    if encodage in ("utf-8", "utf8", "utf-8-sig"):
        return MECANISME_DOUBLE
    return f"octets UTF-8 dans un fichier déclaré {spec['encodage']}"


def _flux_par_table(config: dict) -> dict[str, dict]:
    return {
        table: spec
        for spec in config["flux"].values()
        for table in tables_du_flux(spec)
    }


def encodage_suspect(tables: dict[str, pd.DataFrame], config: dict) -> pd.DataFrame:
    """Décompte des cellules à l'encodage suspect, sans aucune valeur.

    Une ligne par table, colonne et constat de la règle `encodage` : mécanisme
    probable, nombre de cellules, de valeurs distinctes, de cellules que la
    réparation du nettoyage (`nettoyage.reparer_encodage`) sait réparer, et de
    cellules dont la forme réparée existe telle quelle ailleurs dans la colonne.
    Une cellule qui porte les deux constats est comptée sous chacun.
    """
    flux = _flux_par_table(config)
    lignes = []
    for table, donnees in tables.items():
        for colonne in donnees.columns:
            if colonne in COLONNES_LIGNAGE:
                continue
            serie = donnees[colonne]
            textes = serie[[isinstance(v, str) for v in serie]]
            if textes.empty:
                continue
            constats = {v: constats_encodage(v) for v in textes.unique()}
            suspects = {v for v, c in constats.items() if c}
            if not suspects:
                continue
            saines = set(textes.unique()) - suspects
            effectifs = textes[textes.isin(suspects)].value_counts()
            reparations = {v: reparer_encodage(v).texte for v in suspects}
            for constat in (CONSTAT_DECODAGE, CONSTAT_REMPLACEMENT):
                valeurs = [v for v in suspects if constat in constats[v]]
                if not valeurs:
                    continue
                reparables = [
                    v
                    for v in valeurs
                    if reparations[v] is not None and reparations[v] != v
                ]
                lignes.append(
                    {
                        "table": table,
                        "colonne": colonne,
                        "constat": constat,
                        "mecanisme": _mecanisme(flux[table], constat)
                        if table in flux
                        else None,
                        "cellules": int(effectifs[valeurs].sum()),
                        "valeurs_distinctes": len(valeurs),
                        "reparables": int(effectifs[reparables].sum()),
                        "reparees_presentes_ailleurs": int(
                            effectifs[
                                [v for v in reparables if reparations[v] in saines]
                            ].sum()
                        ),
                    }
                )
    return pd.DataFrame(lignes, columns=COLONNES_ENCODAGE)


def texte_export_encodage(suspects: pd.DataFrame, periode_fin: str) -> str:
    """Texte à transmettre à l'assistant pour la partie du rapport sur les cellules
    à l'encodage suspect : décomptes par table, colonne, constat et mécanisme.

    Refuse un tableau qui porterait d'autres colonnes que celles de
    `encodage_suspect` (garantie qu'aucune valeur n'est transmise).
    """
    autres = [c for c in suspects.columns if c not in COLONNES_ENCODAGE]
    if autres:
        raise ValueError(
            "Export refusé : colonnes inattendues " + ", ".join(map(str, autres))
        )
    entete = f"Cellules à l'encodage suspect, périmètre jusqu'à {periode_fin}"
    if suspects.empty:
        return f"{entete}\nAucune cellule suspecte."
    total = int(suspects["cellules"].sum())
    lignes = [
        entete,
        f"{len(suspects)} ligne(s), {total} cellule(s) comptées par constat",
        (
            "table | colonne | constat | mécanisme | cellules | valeurs distinctes "
            "| réparables | réparées présentes ailleurs"
        ),
    ]
    for ligne in suspects.itertuples(index=False):
        lignes.append(" | ".join(str(getattr(ligne, c)) for c in COLONNES_ENCODAGE))
    return "\n".join(lignes)
