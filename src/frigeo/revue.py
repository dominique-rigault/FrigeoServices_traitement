"""Classeur Excel de revue du dictionnaire par le métier (étape 2d).

Le dictionnaire généré n'est ni ouvert dans Excel ni modifié à la main : il est
exporté dans un classeur où le métier choisit un statut par règle et par valeur,
puis le classeur est relu, contrôlé et appliqué au dictionnaire. Ce module ne
contient rien de propre à un client et n'écrit que le classeur.
"""

from __future__ import annotations

import unicodedata
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path, PureWindowsPath
from typing import NamedTuple

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.datavalidation import DataValidation

from . import racine_projet
from .dictionnaire import (
    A_DECIDER,
    AUCUNE_LISTE,
    LISTE_FERMEE,
    ORIGINE_A_ARBITRER,
    ORIGINE_AJOUT,
    ORIGINE_LISTE,
    ORIGINE_OBSERVEE,
    REGLE_ECARTEE,
    REGLES_RETENUES,
    STATUT_INITIAL,
    STATUTS,
    STATUTS_CIBLE,
    STATUTS_VALEUR_OBSERVEE,
    est_masquee,
    horodater,
    problemes_dictionnaire,
)

FEUILLE_MODE_EMPLOI = "mode_emploi"
FEUILLE_REGLES = "regles"
FEUILLE_VALEURS = "valeurs"
# Feuille technique, masquée : une colonne par table, pour les menus des lignes d'ajout.
FEUILLE_LISTES = "listes"

# Règles revues dans le classeur. La règle `format` du dictionnaire n'y figure pas
# encore : elle reste telle que la génération et la fusion la donnent.
REGLES_REVUES = ("obligatoire", "nature", "valeurs")

COLONNES_REGLES = (
    "table",
    "colonne",
    "regle",
    "proposition",
    "regle_en_vigueur",
    "motif",
    "nb_a_arbitrer",
    "statut",
    "regle_retenue",
    "commentaire",
)
COLONNES_VALEURS = (
    "table",
    "colonne",
    "valeur",
    "effectif",
    "origine",
    "remarque",
    "statut",
    "remplacement",
    "commentaire",
)
# Colonnes que le métier peut modifier; les autres cellules sont verrouillées.
MODIFIABLES_REGLES = ("statut", "regle_retenue", "commentaire")
MODIFIABLES_VALEURS = ("statut", "remplacement", "commentaire")
# Colonnes à saisir sur une ligne ajoutée en bas de la feuille des valeurs.
SAISIE_AJOUT = ("table", "colonne", "valeur", "statut", "commentaire")
# Statuts admis sur une valeur ajoutée déjà enregistrée : la garder ou la retirer.
STATUTS_VALEUR_AJOUTEE = ("documenté", STATUT_INITIAL)

LIGNES_AJOUT = 200
# Affiché à la place de l'effectif d'une valeur d'une colonne sensible.
EFFECTIF_MASQUE = "masqué"
# Ligne de la feuille du mode d'emploi qui porte, en colonne B, le nombre de revues
# déjà appliquées au dictionnaire au moment de l'export.
REPERE_REVUES = "Revues déjà appliquées au dictionnaire (repère technique)"
COLONNES_RAPPORT_REVUE = (
    "table",
    "colonne",
    "regle",
    "valeur",
    "evenement",
    "avant",
    "apres",
)
LARGEURS = {
    "table": 22,
    "colonne": 26,
    "regle": 12,
    "proposition": 26,
    "regle_en_vigueur": 22,
    "motif": 55,
    "nb_a_arbitrer": 14,
    "statut": 13,
    "regle_retenue": 24,
    "commentaire": 50,
    "valeur": 34,
    "effectif": 10,
    "origine": 16,
    "remarque": 24,
    "remplacement": 34,
}


def chemin_revue(periode_fin: str) -> Path:
    """Emplacement par défaut du classeur : data/revue/, hors du dépôt."""
    return racine_projet() / "data" / "revue" / f"dictionnaire_revue_{periode_fin}.xlsx"


def lignes_regles(dictionnaire: dict) -> list[dict]:
    """Lignes de la feuille des règles, dans l'ordre du dictionnaire.

    Trois lignes par colonne : `obligatoire`, `nature` et `valeurs` (les règles de
    `REGLES_REVUES`, sans la règle `format`). La ligne
    `valeurs` porte la question « cette colonne est-elle une liste fermée ? » : sa
    proposition s'affiche « liste de N valeurs », ou « aucune » quand la génération
    ne propose pas de liste (le métier peut alors en déclarer une). Chaque ligne
    affiche la décision enregistrée : statut, règle en vigueur, règle retenue (au
    statut « documenté ») et commentaire.
    """
    lignes = []
    for table, colonnes in dictionnaire.items():
        for colonne, regles in colonnes.items():
            for nom in REGLES_REVUES:
                regle = regles[nom]
                proposition = regle["proposition"]
                if nom == "valeurs" and proposition == LISTE_FERMEE:
                    proposees = sum(
                        element["origine"] == ORIGINE_LISTE for element in regle["liste"]
                    )
                    proposition = f"liste de {proposees} valeurs"
                ecarts = sum(
                    valeur
                    for valeur in regle["a_arbitrer"].values()
                    if isinstance(valeur, int) and not isinstance(valeur, bool)
                )
                lignes.append(
                    {
                        "table": table,
                        "colonne": colonne,
                        "regle": nom,
                        "proposition": proposition,
                        "regle_en_vigueur": regle["regle"],
                        "motif": regle["motif"],
                        "nb_a_arbitrer": ecarts,
                        "statut": regle["statut"],
                        "regle_retenue": (
                            regle["regle"] if regle["statut"] == "documenté" else None
                        ),
                        "commentaire": regle["commentaire"] or None,
                    }
                )
    return lignes


def _remarque(valeur: str) -> str:
    """Signale ce qui ne se voit pas à l'écran dans une valeur."""
    remarques = []
    if valeur != valeur.strip(" \t"):
        remarques.append("espaces en bord")
    if " " in valeur or " " in valeur:
        remarques.append("espace insécable")
    return ", ".join(remarques)


def lignes_valeurs(dictionnaire: dict) -> list[dict]:
    """Lignes de la feuille des valeurs : une ligne par valeur des listes du dictionnaire.

    Chaque ligne porte l'effectif, l'origine et la décision enregistrés dans le
    dictionnaire (statut, remplacement, commentaire), valeurs ajoutées et valeurs
    observées des colonnes candidates comprises. Un effectif masqué (colonne
    sensible) s'affiche « masqué ». Une colonne dont la liste est vide ne donne
    aucune ligne.
    """
    return [
        {
            "table": table,
            "colonne": colonne,
            "valeur": element["valeur"],
            "effectif": (
                EFFECTIF_MASQUE if element["effectif"] is None else element["effectif"]
            ),
            "origine": element["origine"],
            "remarque": _remarque(element["valeur"]),
            "statut": element["statut"],
            "remplacement": element["remplacement"],
            "commentaire": element["commentaire"] or None,
        }
        for table, colonnes in dictionnaire.items()
        for colonne, regles in colonnes.items()
        for element in regles["valeurs"]["liste"]
    ]


def _sans_objet(ligne: dict) -> bool:
    """Ligne `valeurs` sans liste proposée ni décision : il n'y a rien à décider."""
    return (
        ligne["regle"] == "valeurs"
        and ligne["statut"] == STATUT_INITIAL
        and ligne["proposition"] == AUCUNE_LISTE
    )


def _compter_reste(regles: list[dict], valeurs: list[dict]) -> dict:
    sans_objet = sum(_sans_objet(ligne) for ligne in regles)
    attente = [
        ligne
        for ligne in regles
        if ligne["statut"] == STATUT_INITIAL and not _sans_objet(ligne)
    ]
    ecartees = {
        (ligne["table"], ligne["colonne"])
        for ligne in regles
        if ligne["regle"] == "valeurs" and ligne["statut"] == "invalide"
    }
    non_declarees = {
        (ligne["table"], ligne["colonne"]) for ligne in regles if _sans_objet(ligne)
    }
    observees = [ligne for ligne in valeurs if ligne["statut"] == STATUT_INITIAL]
    restantes = [
        ligne
        for ligne in observees
        if (ligne["table"], ligne["colonne"]) not in ecartees | non_declarees
    ]
    candidates = sum(
        (ligne["table"], ligne["colonne"]) in non_declarees for ligne in observees
    )
    return {
        "regles": len(regles) - sans_objet,
        "regles_sans_liste": sans_objet,
        "regles_a_traiter": len(attente),
        "regles_a_decider": sum(ligne["proposition"] == A_DECIDER for ligne in attente),
        "regles_avec_ecarts": sum(ligne["nb_a_arbitrer"] > 0 for ligne in attente),
        "valeurs": len(valeurs),
        "valeurs_a_traiter": len(restantes),
        "valeurs_a_arbitrer": sum(
            ligne["origine"] == ORIGINE_A_ARBITRER for ligne in restantes
        ),
        "valeurs_listes_ecartees": len(observees) - len(restantes) - candidates,
        "valeurs_non_declarees": candidates,
        "complet": not attente and not restantes,
    }


def reste_a_traiter(dictionnaire: dict) -> dict:
    """Décompte des lignes du classeur que le métier n'a pas encore revues.

    Une ligne reste à traiter tant que son statut est « observé ». Les nombres sont
    ceux qu'afficherait le mode d'emploi d'un classeur exporté de ce dictionnaire :
    `regles` et `valeurs` (nombres de lignes à revoir dans les deux feuilles),
    `regles_a_traiter`, dont `regles_a_decider` (proposition « à décider ») et
    `regles_avec_ecarts` (écarts à arbitrer), `valeurs_a_traiter`, dont
    `valeurs_a_arbitrer` (valeurs rares). Les valeurs observées d'une colonne dont la liste n'est pas
    déclarée sont comptées à part, dans `valeurs_non_declarees` : elles ne sont à
    revoir que si le métier déclare la liste. Les lignes « valeurs » à « observé » dont
    la proposition est « aucune » sont comptées à part, dans `regles_sans_liste` :
    aucune liste n'y est proposée, il n'y a donc rien à décider par défaut. Les valeurs non revues d'une liste écartée sont comptées à part, dans
    `valeurs_listes_ecartees` : elles n'ont pas à être revues. `complet` est vrai
    quand il ne reste ni règle ni valeur à traiter.
    """
    return _compter_reste(lignes_regles(dictionnaire), lignes_valeurs(dictionnaire))


def _reste_a_traiter(regles: list[dict], valeurs: list[dict]) -> list[str]:
    """Rubrique du mode d'emploi qui dit, à l'export, ce qui reste à revoir."""
    reste = _compter_reste(regles, valeurs)
    return [
        "Ce qui reste à traiter",
        "Une ligne reste à traiter tant que son statut est « observé » : le filtre de "
        "la colonne statut les isole. Les nombres ci-dessous sont ceux du jour de "
        "l'export.",
        f"Feuille « {FEUILLE_REGLES} » : {reste['regles_a_traiter']} lignes à "
        f"« observé » sur {reste['regles']}. Parmi elles, {reste['regles_a_decider']} "
        f"ont une proposition « {A_DECIDER} » (les données ne permettent pas de "
        "proposer une règle : la donner par le statut documenté) et "
        f"{reste['regles_avec_ecarts']} ont des écarts à arbitrer (colonne "
        "nb_a_arbitrer supérieure à 0). Ce sont les lignes à regarder en premier.",
        f"S'y ajoutent {reste['regles_sans_liste']} lignes « valeurs » à « observé » "
        f"dont la proposition est « {AUCUNE_LISTE} » : aucune liste n'y est proposée, "
        "elles n'ont pas à être revues, sauf pour déclarer une liste (voir les cas "
        "particuliers).",
        f"Feuille « {FEUILLE_VALEURS} » : {reste['valeurs_a_traiter']} lignes à "
        f"« observé » sur {reste['valeurs']}, dont {reste['valeurs_a_arbitrer']} "
        f"d'origine « {ORIGINE_A_ARBITRER} » (valeurs rares, souvent des variantes ou "
        f"des erreurs), à regarder en premier. S'y ajoutent "
        f"{reste['valeurs_listes_ecartees']} valeurs de listes écartées, qui n'ont pas "
        f"à être revues, et {reste['valeurs_non_declarees']} valeurs d'origine "
        f"« {ORIGINE_OBSERVEE} », à revoir seulement si vous déclarez la liste de leur "
        "colonne (voir les cas particuliers).",
        "La revue est complète quand plus aucune ligne n'est à « observé », hors "
        "lignes sans liste proposée, valeurs des listes écartées et valeurs des listes "
        "non déclarées. Une revue partielle est possible : les lignes "
        "laissées à « observé » se retrouvent telles quelles dans le classeur suivant.",
    ]


def _mode_emploi(periode_fin: str, regles: list[dict], valeurs: list[dict]) -> list[str]:
    return [
        "Revue du dictionnaire des données",
        f"Période de fin : {periode_fin}. Classeur généré le {date.today():%d/%m/%Y}.",
        f"Contenu : {len(regles)} lignes dans la feuille « {FEUILLE_REGLES} » (trois "
        f"par colonne des données), {len(valeurs)} lignes dans la feuille "
        f"« {FEUILLE_VALEURS} ».",
        REPERE_REVUES,
        "",
        *_reste_a_traiter(regles, valeurs),
        "",
        "Ce qui est attendu",
        "Pour chaque ligne, choisir un statut dans le menu déroulant et, si besoin, "
        "écrire un commentaire. Une ligne laissée à « observé » n'est pas revue : "
        "elle pourra l'être plus tard.",
        "Le classeur affiche les décisions déjà enregistrées et fait foi pour chacune "
        "de ses lignes : remettre une ligne à « observé » annule la décision "
        "précédente.",
        "",
        f"Feuille « {FEUILLE_REGLES} » : les règles proposées pour chaque colonne",
        "valide : la proposition est la règle. Impossible si la proposition est "
        "« à décider ».",
        "invalide : la proposition est écartée et rien ne la remplace. Dire pourquoi "
        "en commentaire.",
        "documenté : vous donnez vous-même la règle, dans la colonne regle_retenue "
        "(à remplir avec ce statut, à laisser vide avec les autres).",
        "La colonne regle_en_vigueur rappelle la règle appliquée aujourd'hui. Une "
        "règle déjà revue la garde, même si la proposition a changé depuis : pour en "
        "changer, choisir documenté et donner la règle dans regle_retenue.",
        "Sur une ligne « valeurs », la question est : cette colonne n'accepte-t-elle "
        "qu'une liste fermée de valeurs ? Le statut invalide écarte toute la liste, "
        f"sans avoir à revoir ses valeurs dans la feuille « {FEUILLE_VALEURS} ». Quand "
        f"la proposition est « {AUCUNE_LISTE} », l'outil ne propose pas de liste : la "
        "ligne peut rester à « observé ».",
        "",
        f"Feuille « {FEUILLE_VALEURS} » : les valeurs des listes proposées",
        "valide : la valeur est légitime.",
        "invalide : la valeur est une erreur, elle sera signalée en anomalie.",
        "remplacement (facultatif) : pour une valeur invalide, la valeur correcte "
        "qui la remplacera. Elle se recopie à l'identique depuis une valeur de la "
        "même colonne, au statut valide ou documenté.",
        "La colonne remarque signale ce qui ne se voit pas à l'écran (espaces en "
        "bord, espace insécable).",
        "Pour ajouter une valeur légitime absente de la liste : sur une ligne vide "
        "en bas de la feuille, choisir la table puis la colonne dans les menus "
        "déroulants (le menu de la colonne dépend de la table choisie), saisir la "
        "valeur et choisir le statut documenté.",
        "",
        "Cas particuliers",
        "Annuler une décision (règle ou valeur) : remettre son statut à « observé ». "
        "La décision et sa date de revue sont effacées.",
        "Retirer une valeur ajoutée (origine « ajoutée ») : remettre son statut à "
        "« observé ». La valeur disparaît de la liste si elle n'est pas observée "
        "dans les données; sinon elle redevient une valeur à arbitrer.",
        f"Écarter toute une liste : statut invalide sur la ligne « valeurs » de la "
        f"feuille « {FEUILLE_REGLES} », avec un commentaire. Les valeurs et leurs "
        "statuts sont conservés, sans effet tant que la liste est écartée.",
        "Statut invalide sur une règle : plus aucun contrôle de ce type ne portera "
        "sur la colonne.",
        "Changer une règle déjà validée quand la proposition a changé : choisir "
        "documenté et donner la règle dans regle_retenue. Laisser valide conserve "
        "l'ancienne règle, celle de la colonne regle_en_vigueur.",
        "Déclarer correcte une valeur absente des données : l'ajouter en bas de la "
        f"feuille « {FEUILLE_VALEURS} », au statut documenté.",
        "Déclarer une liste que l'outil n'a pas proposée (proposition "
        f"« {AUCUNE_LISTE} ») : sur la ligne « valeurs » de la feuille "
        f"« {FEUILLE_REGLES} », choisir documenté et « {LISTE_FERMEE} » dans "
        f"regle_retenue. Puis, dans la feuille « {FEUILLE_VALEURS} », revoir les "
        f"valeurs d'origine « {ORIGINE_OBSERVEE} » de la colonne quand il y en a "
        "(valide, ou invalide avec son remplacement), et ajouter en bas de la feuille, "
        "au statut documenté, les valeurs correctes absentes des données. Remettre la "
        "ligne à « observé » annule la déclaration; les valeurs revues et ajoutées "
        "restent tant qu'elles ne sont pas remises à « observé » une à une.",
        f"Colonne sensible (motif « colonne sensible », effectif « {EFFECTIF_MASQUE} ») : "
        "l'outil n'y propose aucune valeur et n'y affiche aucun effectif. Sa liste "
        "se déclare en ajoutant les valeurs à la main : n'y saisir que des libellés "
        "de nomenclature, jamais une information propre à une personne.",
        "Remplacer une valeur erronée : statut invalide, puis la valeur correcte dans "
        "la colonne remplacement. Une valeur qui sert de remplacement ne peut être "
        "ni invalidée ni retirée tant que d'autres lignes la visent.",
        "Modifier seulement un commentaire : il est enregistré sans changer la date "
        "de la décision.",
        "Reprendre une revue : un classeur ne s'applique qu'une fois. Après son "
        "application, ou celle d'un autre classeur, il est à réexporter.",
        "Avant l'import : enregistrer puis fermer le classeur dans Excel.",
        "",
        "Règles à respecter pour que le classeur puisse être réimporté",
        "Ne modifier que les colonnes statut, regle_retenue, remplacement et "
        "commentaire (les autres cellules sont verrouillées).",
        "Ne pas renommer, déplacer ni supprimer de feuilles, de colonnes ou de lignes "
        f"(la feuille masquée « {FEUILLE_LISTES} » alimente les menus des lignes "
        "ajoutées).",
        "Ne pas ajouter de lignes, sauf en bas de la feuille des valeurs comme "
        "indiqué ci-dessus.",
        "Choisir les statuts dans le menu déroulant : une valeur saisie hors d'un "
        "menu est refusée.",
        "Enregistrer au format .xlsx, sans changer le nom du fichier.",
        "Le filtre de la ligne d'en-tête peut être utilisé librement.",
    ]


def _ecrire_texte(cellule, valeur) -> None:
    """Écrit une valeur comme texte, pour qu'Excel ne la transforme pas.

    Sans cela, « 076 » deviendrait 76, « 2026-03-01 » une date et « =1+1 » une formule.
    """
    cellule.value = valeur
    cellule.number_format = "@"
    if isinstance(valeur, str):
        cellule.data_type = "s"


def _menu(feuille, choix=None, formule: str | None = None) -> DataValidation:
    """Menu déroulant d'une liste de choix, ou d'une plage donnée par une formule.

    Excel refuse une valeur saisie hors du menu.
    """
    if formule is None:
        formule = '"' + ",".join(choix) + '"'
    menu = DataValidation(
        type="list", formula1=formule, allow_blank=True, showErrorMessage=True
    )
    menu.error = "Choisir une valeur dans le menu déroulant."
    menu.errorTitle = "Valeur non prévue"
    feuille.add_data_validation(menu)
    return menu


def _menus_d_ajout(classeur, feuille, dictionnaire: dict, premiere: int, derniere: int) -> None:
    """Menus `table` et `colonne` des lignes d'ajout de la feuille des valeurs.

    Les choix sont rangés dans une feuille technique masquée, une colonne par table :
    le nom de la table en première ligne, ses colonnes en dessous. Le menu `colonne`
    d'une ligne ne propose que les colonnes de la table choisie sur cette ligne.
    """
    listes = classeur.create_sheet(FEUILLE_LISTES)
    for indice, (table, colonnes) in enumerate(dictionnaire.items(), start=1):
        _ecrire_texte(listes.cell(row=1, column=indice), table)
        for numero, colonne in enumerate(colonnes, start=2):
            _ecrire_texte(listes.cell(row=numero, column=indice), colonne)
    listes.protection.sheet = True
    listes.sheet_state = "hidden"
    if not dictionnaire or derniere < premiere:
        return

    nom = quote_sheetname(FEUILLE_LISTES)
    tables = f"{nom}!$A$1:${get_column_letter(len(dictionnaire))}$1"
    col_table = get_column_letter(COLONNES_VALEURS.index("table") + 1)
    col_colonne = get_column_letter(COLONNES_VALEURS.index("colonne") + 1)
    _menu(feuille, formule=tables).add(f"{col_table}{premiere}:{col_table}{derniere}")
    # La référence à la table est relative à la ligne : elle suit chaque ligne d'ajout.
    rang = f"MATCH(${col_table}{premiere},{nom}!$1:$1,0)-1"
    colonnes = (
        f"OFFSET({nom}!$A$1,1,{rang},"
        f"COUNTA(OFFSET({nom}!$A:$A,0,{rang}))-1,1)"
    )
    _menu(feuille, formule=colonnes).add(f"{col_colonne}{premiere}:{col_colonne}{derniere}")


def _remplir(feuille, colonnes, lignes, modifiables, lignes_ajout=0) -> None:
    """Écrit l'en-tête et les lignes, puis protège la feuille hors colonnes modifiables."""
    for indice, nom in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=1, column=indice, value=nom)
        cellule.font = Font(bold=True)
        feuille.column_dimensions[cellule.column_letter].width = LARGEURS[nom]
    for numero, ligne in enumerate(lignes, start=2):
        for indice, nom in enumerate(colonnes, start=1):
            cellule = feuille.cell(row=numero, column=indice)
            valeur = ligne.get(nom)
            if isinstance(valeur, int):
                cellule.value = valeur
            else:
                _ecrire_texte(cellule, valeur)
            if nom in modifiables:
                cellule.protection = Protection(locked=False)
    fin = len(lignes) + 1
    for numero in range(fin + 1, fin + lignes_ajout + 1):
        for indice, nom in enumerate(colonnes, start=1):
            if nom in SAISIE_AJOUT:
                cellule = feuille.cell(row=numero, column=indice)
                cellule.number_format = "@"
                cellule.protection = Protection(locked=False)
    derniere = feuille.cell(row=1, column=len(colonnes)).column_letter
    # La ligne d'en-tête et les colonnes table et colonne restent visibles.
    feuille.freeze_panes = "C2"
    feuille.auto_filter.ref = f"A1:{derniere}{max(fin + lignes_ajout, 2)}"
    feuille.protection.sheet = True
    feuille.protection.autoFilter = False
    feuille.protection.formatColumns = False


def exporter_revue(
    dictionnaire: dict,
    periode_fin: str,
    chemin: str | Path | None = None,
    *,
    nb_revues: int,
    ecraser: bool = False,
) -> Path:
    """Écrit le classeur de revue du dictionnaire et rend son chemin.

    Le classeur affiche les décisions déjà enregistrées dans le dictionnaire.
    `nb_revues` est le nombre de revues déjà appliquées au dictionnaire
    (`len(charger_meta()["revues"])`) : il est inscrit dans le classeur, et l'import
    refusera le classeur si une revue a été appliquée depuis. Sans `ecraser=True`,
    refuse de remplacer un classeur existant : une revue y est peut-être en cours.
    """
    if isinstance(nb_revues, bool) or not isinstance(nb_revues, int) or nb_revues < 0:
        raise ValueError(f"nb_revues {nb_revues!r} : un entier positif ou nul est attendu")
    chemin = Path(chemin) if chemin is not None else chemin_revue(periode_fin)
    if chemin.exists() and not ecraser:
        raise FileExistsError(
            f"{chemin} existe déjà : export refusé (ecraser=True pour forcer)"
        )
    regles = lignes_regles(dictionnaire)
    valeurs = lignes_valeurs(dictionnaire)

    classeur = Workbook()
    notice = classeur.active
    notice.title = FEUILLE_MODE_EMPLOI
    notice.column_dimensions["A"].width = 120
    for numero, texte in enumerate(
        _mode_emploi(periode_fin, regles, valeurs), start=1
    ):
        cellule = notice.cell(row=numero, column=1)
        _ecrire_texte(cellule, texte)
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
        if texte == REPERE_REVUES:
            notice.cell(row=numero, column=2, value=nb_revues)
        if numero == 1 or texte.startswith(("Ce qui", "Feuille", "Cas particuliers", "Règles à")):
            cellule.font = Font(bold=True)
    notice.protection.sheet = True

    feuille = classeur.create_sheet(FEUILLE_REGLES)
    _remplir(feuille, COLONNES_REGLES, regles, MODIFIABLES_REGLES)
    statut = COLONNES_REGLES.index("statut") + 1
    retenue = COLONNES_REGLES.index("regle_retenue") + 1
    menu_statut = _menu(feuille, STATUTS)
    menus_retenue = {nom: _menu(feuille, choix) for nom, choix in REGLES_RETENUES.items()}
    for numero, ligne in enumerate(regles, start=2):
        menu_statut.add(feuille.cell(row=numero, column=statut))
        # Sur une ligne « valeurs », le menu n'a qu'un choix : déclarer une liste fermée.
        menus_retenue[ligne["regle"]].add(feuille.cell(row=numero, column=retenue))

    feuille = classeur.create_sheet(FEUILLE_VALEURS)
    _remplir(feuille, COLONNES_VALEURS, valeurs, MODIFIABLES_VALEURS, LIGNES_AJOUT)
    statut = COLONNES_VALEURS.index("statut") + 1
    menu_statut = _menu(feuille, STATUTS)
    for numero in range(2, len(valeurs) + LIGNES_AJOUT + 2):
        menu_statut.add(feuille.cell(row=numero, column=statut))
    _menus_d_ajout(
        classeur, feuille, dictionnaire, len(valeurs) + 2, len(valeurs) + LIGNES_AJOUT + 1
    )

    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin


class Revue(NamedTuple):
    """Décisions lues dans un classeur de revue, contrôlées mais pas encore appliquées."""

    regles: pd.DataFrame
    valeurs: pd.DataFrame
    avertissements: list[str]


def _choix(valeur) -> str:
    """Texte d'un statut ou d'une règle retenue, normalisé (NFC, espaces en bord ôtés)."""
    if valeur is None:
        return ""
    return unicodedata.normalize("NFC", str(valeur)).strip()


def _commentaire(valeur) -> str:
    return "" if valeur is None else str(valeur).strip()


def _est_vide(valeur) -> bool:
    return valeur is None or (isinstance(valeur, str) and valeur.strip() == "")


def _lire_feuille(classeur, nom: str, colonnes: tuple[str, ...]) -> list[tuple[int, dict]]:
    """Lignes non vides d'une feuille, avec leur numéro de ligne Excel.

    Lève une ValueError si la feuille manque ou si son en-tête n'est pas celui exporté.
    """
    if nom not in classeur.sheetnames:
        raise ValueError(f"feuille « {nom} » absente")
    lignes = classeur[nom].iter_rows(values_only=True)
    entete = tuple(next(lignes, ()))
    while entete and entete[-1] is None:
        entete = entete[:-1]
    if entete != colonnes:
        raise ValueError(
            f"feuille « {nom} » : en-tête modifié, attendu {', '.join(colonnes)} "
            "(classeur exporté par une version antérieure : à réexporter)"
        )
    resultat = []
    for numero, valeurs in enumerate(lignes, start=2):
        valeurs = tuple(valeurs[: len(colonnes)]) + (None,) * (len(colonnes) - len(valeurs))
        if all(_est_vide(valeur) for valeur in valeurs):
            continue
        resultat.append((numero, dict(zip(colonnes, valeurs))))
    return resultat


def _lire_repere(classeur) -> int:
    """Nombre de revues déjà appliquées au dictionnaire quand le classeur a été exporté."""
    if FEUILLE_MODE_EMPLOI not in classeur.sheetnames:
        raise ValueError(f"feuille « {FEUILLE_MODE_EMPLOI} » absente")
    for ligne in classeur[FEUILLE_MODE_EMPLOI].iter_rows(max_col=2, values_only=True):
        if ligne and ligne[0] == REPERE_REVUES:
            nombre = ligne[1] if len(ligne) > 1 else None
            if isinstance(nombre, int) and not isinstance(nombre, bool) and nombre >= 0:
                return nombre
            break
    raise ValueError(
        f"feuille « {FEUILLE_MODE_EMPLOI} » : repère des revues absent ou modifié, "
        "classeur à réexporter"
    )


def _controler_regles(lignes, dictionnaire) -> tuple[list[dict], list[str]]:
    attendues = {
        (l["table"], l["colonne"], l["regle"]): l for l in lignes_regles(dictionnaire)
    }
    decisions, problemes, vues = [], [], set()
    for numero, ligne in lignes:
        ou = f"{FEUILLE_REGLES}, ligne {numero}"
        cle = (ligne["table"], ligne["colonne"], ligne["regle"])
        if cle not in attendues:
            problemes.append(f"{ou} : règle inconnue {cle} (classeur périmé ou ligne ajoutée)")
            continue
        if cle in vues:
            problemes.append(f"{ou} : règle {cle} présente plusieurs fois")
            continue
        vues.add(cle)
        for champ in ("proposition", "regle_en_vigueur"):
            if ligne[champ] != attendues[cle][champ]:
                problemes.append(
                    f"{ou} : {champ} {ligne[champ]!r} différente du dictionnaire "
                    f"({attendues[cle][champ]!r}), classeur périmé, à réexporter"
                )
        en_place = dictionnaire[cle[0]][cle[1]][cle[2]]
        statut = _choix(ligne["statut"])
        retenue = _choix(ligne["regle_retenue"])
        commentaire = _commentaire(ligne["commentaire"])
        regle = ligne["regle"]
        # Une règle déjà validée le reste, même si la proposition a changé depuis.
        validation = statut == "valide" and en_place["statut"] != "valide"
        if statut == "":
            problemes.append(f"{ou} : statut vide (laisser « observé » si la règle n'est pas revue)")
        elif statut not in STATUTS:
            problemes.append(f"{ou} : statut inconnu {statut!r}")
        elif validation and en_place["proposition"] == A_DECIDER:
            problemes.append(
                f"{ou} : une proposition « {A_DECIDER} » ne peut pas être validée "
                "(choisir documenté et une règle retenue)"
            )
        elif validation and en_place["proposition"] == REGLE_ECARTEE:
            problemes.append(
                f"{ou} : aucune liste n'est proposée, il n'y a rien à valider "
                "(laisser le statut ou remettre « observé »)"
            )
        elif statut == "invalide" and commentaire == "":
            problemes.append(f"{ou} : commentaire obligatoire avec le statut invalide")
        declaration = (
            regle == "valeurs" and statut == "documenté" and en_place["statut"] != "documenté"
        )
        if declaration and en_place["proposition"] == LISTE_FERMEE:
            problemes.append(
                f"{ou} : une liste est déjà proposée, il n'y a rien à déclarer "
                "(choisir valide)"
            )
        elif statut == "documenté" and retenue == "":
            problemes.append(f"{ou} : regle_retenue obligatoire avec le statut documenté")
        elif statut == "documenté" and retenue not in REGLES_RETENUES[regle]:
            problemes.append(f"{ou} : regle_retenue {retenue!r} hors du menu de la règle {regle}")
        elif statut != "documenté" and retenue:
            problemes.append(f"{ou} : regle_retenue réservée au statut documenté")
        decisions.append(
            {
                "table": cle[0],
                "colonne": cle[1],
                "regle": regle,
                "statut": statut,
                "regle_retenue": retenue or None,
                "commentaire": commentaire,
                "ligne_excel": numero,
            }
        )
    for cle in attendues.keys() - vues:
        problemes.append(f"{FEUILLE_REGLES} : règle {cle} absente du classeur")
    return decisions, problemes


def _controler_remplacements(decisions: list[dict]) -> list[str]:
    """Chaque remplacement vise une valeur valide ou documentée de la même colonne.

    Le classeur porte toutes les valeurs de la colonne : le statut de la cible est
    celui du classeur, valeur ajoutée dans le même classeur comprise. Une cible
    invalidée, retirée ou remise à « observé » donne un seul problème, sur sa ligne,
    avec les lignes qui la visent.
    """
    cibles = {
        (d["table"], d["colonne"], d["valeur"]): d
        for d in decisions
        if isinstance(d["valeur"], str)
    }
    problemes, visees = [], {}
    for d in decisions:
        cible = d["remplacement"]
        if not isinstance(cible, str) or d["statut"] != "invalide" or cible == d["valeur"]:
            continue
        cle = (d["table"], d["colonne"], cible)
        if cle not in cibles:
            problemes.append(
                f"{FEUILLE_VALEURS}, ligne {d['ligne_excel']} : remplacement {cible!r} "
                f"absent des valeurs de {d['table']}.{d['colonne']} (le recopier à "
                "l'identique)"
            )
        else:
            visees.setdefault(cle, []).append(d["ligne_excel"])
    for cle, lignes in visees.items():
        cible = cibles[cle]
        if cible["statut"] not in STATUTS_CIBLE:
            problemes.append(
                f"{FEUILLE_VALEURS}, ligne {cible['ligne_excel']} : valeur {cle[2]!r} au "
                f"statut {cible['statut']!r} alors qu'elle sert de remplacement "
                f"(lignes {', '.join(str(n) for n in lignes)}), une valeur de "
                f"remplacement doit être au statut {' ou '.join(STATUTS_CIBLE)}"
            )
    return problemes


def _controler_valeurs(lignes, dictionnaire) -> tuple[list[dict], list[str]]:
    exportees = {
        (l["table"], l["colonne"], l["valeur"]): l["origine"]
        for l in lignes_valeurs(dictionnaire)
    }
    colonnes = {(table, colonne) for table, cols in dictionnaire.items() for colonne in cols}
    decisions, problemes, vues = [], [], set()
    for numero, ligne in lignes:
        ou = f"{FEUILLE_VALEURS}, ligne {numero}"
        table, colonne, valeur = ligne["table"], ligne["colonne"], ligne["valeur"]
        statut = _choix(ligne["statut"])
        remplacement = None if _est_vide(ligne["remplacement"]) else ligne["remplacement"]
        ajout = _est_vide(ligne["origine"])
        cle = (table, colonne, valeur)
        if ajout:
            origine = ORIGINE_AJOUT
            if (table, colonne) not in colonnes:
                problemes.append(f"{ou} : colonne {table}.{colonne} absente du dictionnaire")
            if _est_vide(valeur):
                problemes.append(f"{ou} : valeur vide")
            elif not isinstance(valeur, str):
                problemes.append(
                    f"{ou} : valeur {valeur!r} lue comme {type(valeur).__name__}, "
                    "la saisir comme texte"
                )
            elif cle in exportees or cle in vues:
                problemes.append(f"{ou} : valeur {valeur!r} déjà présente pour {table}.{colonne}")
            if statut != "documenté":
                problemes.append(f"{ou} : une valeur ajoutée doit porter le statut documenté")
            if remplacement is not None:
                problemes.append(f"{ou} : remplacement interdit sur une valeur ajoutée")
        else:
            origine = ligne["origine"]
            if cle not in exportees or exportees[cle] != origine:
                problemes.append(
                    f"{ou} : valeur {valeur!r} de {table}.{colonne} inconnue du "
                    "dictionnaire (classeur périmé ou ligne modifiée)"
                )
                continue
            if cle in vues:
                problemes.append(f"{ou} : valeur {valeur!r} présente plusieurs fois")
                continue
            if origine == ORIGINE_AJOUT:
                if statut not in STATUTS_VALEUR_AJOUTEE:
                    problemes.append(
                        f"{ou} : statut {statut!r} impossible sur une valeur ajoutée "
                        "(documenté pour la garder, observé pour la retirer)"
                    )
            elif statut not in STATUTS_VALEUR_OBSERVEE:
                problemes.append(
                    f"{ou} : statut {statut!r} impossible sur une valeur observée "
                    f"(choisir parmi {', '.join(STATUTS_VALEUR_OBSERVEE)})"
                )
            if remplacement is None:
                pass
            elif not isinstance(remplacement, str):
                problemes.append(
                    f"{ou} : remplacement {remplacement!r} lu comme "
                    f"{type(remplacement).__name__}, le saisir comme texte"
                )
            elif statut != "invalide":
                problemes.append(f"{ou} : remplacement réservé au statut invalide")
            elif remplacement == valeur:
                problemes.append(f"{ou} : une valeur ne peut pas être son propre remplacement")
        vues.add(cle)
        decisions.append(
            {
                "table": table,
                "colonne": colonne,
                "valeur": valeur,
                "origine": origine,
                "statut": statut,
                "remplacement": remplacement,
                "commentaire": _commentaire(ligne["commentaire"]),
                "ligne_excel": numero,
            }
        )
    for cle in exportees.keys() - vues:
        problemes.append(f"{FEUILLE_VALEURS} : valeur {cle} absente du classeur")
    problemes.extend(_controler_remplacements(decisions))
    return decisions, problemes


def _controler_listes(regles: list[dict], valeurs: list[dict], dictionnaire) -> list[str]:
    """Cohérence entre la ligne `valeurs` d'une colonne et les valeurs du classeur.

    Une liste déclarée (statut « documenté ») contient au moins une valeur au statut
    valide ou documenté. Sans liste proposée, il n'y a rien à écarter tant que la
    colonne n'a aucune valeur.
    """
    par_colonne: dict[tuple, list[str]] = {}
    for v in valeurs:
        par_colonne.setdefault((v["table"], v["colonne"]), []).append(v["statut"])
    problemes = []
    for r in regles:
        if r["regle"] != "valeurs":
            continue
        ou = f"{FEUILLE_REGLES}, ligne {r['ligne_excel']}"
        statuts = par_colonne.get((r["table"], r["colonne"]), [])
        en_place = dictionnaire[r["table"]][r["colonne"]]["valeurs"]
        if r["statut"] == "documenté":
            declaree = en_place["statut"] == "documenté"
            if not declaree and en_place["proposition"] == LISTE_FERMEE:
                continue  # déjà signalé : une liste proposée se valide
            if not any(statut in STATUTS_CIBLE for statut in statuts):
                problemes.append(
                    f"{ou} : liste déclarée sans aucune valeur au statut "
                    f"{' ou '.join(STATUTS_CIBLE)} pour {r['table']}.{r['colonne']} "
                    f"(ajouter ses valeurs en bas de la feuille « {FEUILLE_VALEURS} »)"
                )
        elif (
            r["statut"] == "invalide"
            and en_place["statut"] != "invalide"
            and en_place["proposition"] == AUCUNE_LISTE
            and not statuts
        ):
            problemes.append(
                f"{ou} : aucune liste n'est proposée et la colonne n'a aucune valeur, "
                "il n'y a rien à écarter (remettre « observé »)"
            )
    return problemes


def _avertissements(regles: list[dict], valeurs: list[dict]) -> list[str]:
    """Valeurs revues d'une liste écartée ou non déclarée (non bloquant)."""
    etats = {
        (r["table"], r["colonne"]): r["statut"] for r in regles if r["regle"] == "valeurs"
    }
    revues: dict[tuple[str, str], int] = {}
    for v in valeurs:
        cle = (v["table"], v["colonne"])
        ecartee = etats.get(cle) == "invalide"
        non_declaree = etats.get(cle) == STATUT_INITIAL and v["origine"] == ORIGINE_OBSERVEE
        if (ecartee or non_declaree) and v["statut"] in ("valide", "invalide"):
            revues[cle] = revues.get(cle, 0) + 1
    return [
        (
            f"{table}.{colonne} : liste écartée, mais {n} valeurs revues "
            "(statuts conservés, sans effet tant que la liste est écartée)"
            if etats[(table, colonne)] == "invalide"
            else f"{table}.{colonne} : {n} valeurs revues, mais liste non déclarée "
            "(statuts conservés, sans effet tant que la liste n'est pas déclarée)"
        )
        for (table, colonne), n in sorted(revues.items())
    ]


def importer_revue(chemin: str | Path, dictionnaire: dict, *, nb_revues: int) -> Revue:
    """Lit un classeur de revue et contrôle chaque décision, sans modifier le dictionnaire.

    `dictionnaire` est le dictionnaire courant et `nb_revues` le nombre de revues
    qui lui ont déjà été appliquées (`len(charger_meta()["revues"])`). Un classeur
    périmé est refusé : exporté avant la dernière revue appliquée (il annulerait
    les décisions prises depuis), ou avant une régénération qui a changé une
    proposition. Tous les problèmes bloquants sont listés en une fois dans une
    ValueError, avec la feuille et le numéro de ligne Excel. Les avertissements ne
    bloquent pas l'import.
    """
    chemin = Path(chemin)
    classeur = load_workbook(chemin, read_only=True, data_only=True)
    try:
        try:
            lignes_r = _lire_feuille(classeur, FEUILLE_REGLES, COLONNES_REGLES)
            lignes_v = _lire_feuille(classeur, FEUILLE_VALEURS, COLONNES_VALEURS)
            repere = _lire_repere(classeur)
        except ValueError as erreur:
            raise ValueError(f"Classeur de revue invalide ({chemin})\n  {erreur}") from None
    finally:
        classeur.close()
    if repere != nb_revues:
        raise ValueError(
            f"Classeur de revue périmé ({chemin}) : exporté après {repere} revues, le "
            f"dictionnaire en compte {nb_revues}. Ce classeur a déjà été appliqué, ou "
            "un autre l'a été depuis son export : à réexporter."
        )

    regles, problemes_r = _controler_regles(lignes_r, dictionnaire)
    valeurs, problemes_v = _controler_valeurs(lignes_v, dictionnaire)
    problemes = problemes_r + problemes_v + _controler_listes(regles, valeurs, dictionnaire)
    if problemes:
        raise ValueError(
            f"Classeur de revue invalide ({chemin}), {len(problemes)} problèmes\n  "
            + "\n  ".join(problemes)
        )
    return Revue(
        pd.DataFrame(regles, columns=[
            "table", "colonne", "regle", "statut", "regle_retenue", "commentaire", "ligne_excel",
        ]),
        pd.DataFrame(valeurs, columns=[
            "table", "colonne", "valeur", "origine", "statut", "remplacement", "commentaire",
            "ligne_excel",
        ]),
        _avertissements(regles, valeurs),
    )


def _etat_regle(regle: dict) -> str:
    return f"{regle['statut']}, règle {regle['regle']}"


def _etat_valeur(element: dict) -> str:
    if element["remplacement"] is None:
        return element["statut"]
    return f"{element['statut']}, remplacée par {element['remplacement']!r}"


def _texte_ou_rien(valeur) -> str | None:
    """Cellule de texte d'un tableau de décisions, None si elle est vide."""
    return valeur if isinstance(valeur, str) and valeur != "" else None


def appliquer_revue(
    dictionnaire: dict,
    revue: Revue,
    classeur: str | Path,
    instant: datetime | str | None = None,
) -> tuple[dict, pd.DataFrame, dict | None]:
    """Applique au dictionnaire les décisions d'un classeur importé.

    `revue` est le résultat de `importer_revue` pour ce même dictionnaire, `classeur`
    le chemin du classeur (seul son nom est gardé au journal) et `instant` la date et
    l'heure de l'application (maintenant par défaut, voir `horodater`). Rien n'est
    modifié ni écrit : la fonction rend le dictionnaire mis à jour, le rapport des
    changements (une ligne par règle ou valeur touchée) et l'entrée à ajouter au
    journal des revues par `ecrire_dictionnaire(..., revue=entree)`. Sans aucun
    changement, l'entrée est None et il n'y a rien à écrire.

    Le classeur fait foi pour chacune de ses lignes. Une décision change quand le
    statut, la règle retenue ou le remplacement change : la date `revu_le` prend
    alors `instant`, au format 'AAAA-MM-JJ HH:MM:SS'. Un commentaire modifié seul est
    enregistré sans toucher à la date.

    Règle : « valide » adopte la proposition du moment, « documenté » la règle
    retenue (sur une ligne `valeurs`, une liste fermée déclarée par le métier),
    « invalide » ne laisse aucune règle, et le retour à « observé » rend la
    proposition. Une décision inchangée garde sa règle, même si la proposition a
    changé depuis.

    Valeur : une ligne ajoutée entre dans la liste (origine « ajoutée », statut
    « documenté », effectif 0, ou masqué dans une colonne sensible). Une valeur
    remise à « observé » est retirée si son effectif est nul ou masqué; sinon elle redevient une valeur observée, à arbitrer si elle
    avait été ajoutée.
    """
    jour = horodater(instant)
    resultat = deepcopy(dictionnaire)
    evenements = []
    compte = {
        "regles_changees": 0,
        "valeurs_changees": 0,
        "valeurs_ajoutees": 0,
        "commentaires_changes": 0,
    }

    def commenter(element, commentaire, change, *cle):
        if commentaire == element["commentaire"]:
            return
        if not change:
            evenements.append((*cle, "commentaire modifié", element["commentaire"], commentaire))
            compte["commentaires_changes"] += 1
        element["commentaire"] = commentaire

    for d in revue.regles.itertuples(index=False):
        regle = resultat[d.table][d.colonne][d.regle]
        retenue = _texte_ou_rien(d.regle_retenue)
        avant = _etat_regle(regle)
        change = d.statut != regle["statut"] or (
            d.statut == "documenté" and retenue != regle["regle"]
        )
        if change:
            regle["statut"] = d.statut
            if d.statut == "documenté":
                regle["regle"] = retenue
            elif d.statut == "invalide":
                regle["regle"] = REGLE_ECARTEE
            else:
                regle["regle"] = regle["proposition"]
            regle["revu_le"] = None if d.statut == STATUT_INITIAL else jour
            evenements.append(
                (d.table, d.colonne, d.regle, "", "décision changée", avant, _etat_regle(regle))
            )
            compte["regles_changees"] += 1
        commenter(regle, d.commentaire, change, d.table, d.colonne, d.regle, "")

    for d in revue.valeurs.itertuples(index=False):
        regle_valeurs = resultat[d.table][d.colonne]["valeurs"]
        liste = regle_valeurs["liste"]
        cle = (d.table, d.colonne, "valeurs", d.valeur)
        element = next((e for e in liste if e["valeur"] == d.valeur), None)
        if element is None:
            liste.append(
                {
                    "valeur": d.valeur,
                    "origine": ORIGINE_AJOUT,
                    "effectif": None if est_masquee(regle_valeurs) else 0,
                    "statut": "documenté",
                    "commentaire": d.commentaire,
                    "revu_le": jour,
                    "remplacement": None,
                }
            )
            evenements.append((*cle, "valeur ajoutée", "", "documenté"))
            compte["valeurs_ajoutees"] += 1
            continue
        remplacement = _texte_ou_rien(d.remplacement)
        avant = _etat_valeur(element)
        change = d.statut != element["statut"] or remplacement != element["remplacement"]
        if change:
            compte["valeurs_changees"] += 1
            if d.statut == STATUT_INITIAL and not element["effectif"]:
                # Plus observée et plus revue : rien ne justifie de la garder.
                liste.remove(element)
                evenements.append((*cle, "valeur retirée", avant, ""))
                continue
            element["statut"] = d.statut
            element["remplacement"] = remplacement
            element["revu_le"] = None if d.statut == STATUT_INITIAL else jour
            if d.statut == STATUT_INITIAL and element["origine"] == ORIGINE_AJOUT:
                element["origine"] = ORIGINE_A_ARBITRER
            evenements.append((*cle, "décision changée", avant, _etat_valeur(element)))
        commenter(element, d.commentaire, change, *cle)

    problemes = problemes_dictionnaire(resultat)
    if problemes:
        raise ValueError(
            "Revue inapplicable, le dictionnaire obtenu serait incohérent\n  "
            + "\n  ".join(problemes)
        )
    rapport = pd.DataFrame(evenements, columns=list(COLONNES_RAPPORT_REVUE))
    if not evenements:
        return resultat, rapport, None
    entree = {"date": jour, "classeur": PureWindowsPath(str(classeur)).name, **compte}
    return resultat, rapport, entree
