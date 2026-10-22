"""Formulaire des règles de contrôle par le métier (phase 5.1 du mode opératoire).

Ce module ne contient rien de propre à un client. Les règles que le dictionnaire ne
peut pas déduire des données (nombre de décimales, bornes, clé de Luhn) et les
déclarations de doublons sont écrites dans config/controles.yaml. Plutôt que de faire
écrire ce fichier à la main, le métier les revoit dans un classeur Excel :

- le classeur montre, à côté de propositions tirées des données (valeurs observées),
  les règles en vigueur lues dans config/controles.yaml, si bien qu'il donne en clair
  l'état effectif des règles;
- le métier choisit un statut par ligne, comme dans la revue du dictionnaire, et
  saisit au besoin ses propres valeurs;
- le classeur relu réécrit config/controles.yaml (morceau suivant).

Le fichier YAML ne garde pas de statut : une règle présente a été donnée par le
métier, elle sort au statut « documenté » avec ses valeurs. Une proposition refusée
(statut « invalide ») est retenue sous la clé `ecartees`, pour ne pas revenir à
chaque cycle; les commentaires du métier le sont sous la clé `commentaires`. Ces deux
clés sont sans effet sur les contrôles.

Le classeur porte des valeurs observées (minimum, maximum) : il est rangé hors du
dépôt, dans data/formulaires/. Les minimums et maximums d'une colonne sensible y sont
masqués, sans proposition de bornes.
"""

from __future__ import annotations

import hashlib
from datetime import date
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, Protection
from openpyxl.utils import get_column_letter, quote_sheetname

from . import racine_projet
from .confidentialite import est_sensible
from .controles import (
    CLE_COMMENTAIRES,
    CLE_DOUBLONS,
    CLE_ECARTEES,
    CLE_REGLES_COMPLEMENTAIRES,
    EXCLUSIONS,
    _cle_luhn_valide,
    _nombre,
    _parametre_en_texte,
    _trop_de_decimales,
    chemin_regles_complementaires,
)
from .dictionnaire import STATUTS, _valeurs_renseignees
from .revue import _ecrire_texte, _menu
from .typage import NATURES_NOMBRE

FEUILLE_MODE_EMPLOI = "mode_emploi"
FEUILLE_DECIMALES = "décimales"
FEUILLE_BORNES = "bornes"
FEUILLE_LUHN = "clé_luhn"
FEUILLE_DOUBLONS = "doublons"
FEUILLE_TABLES = "tables"
FEUILLE_MOTIFS = "motifs"
# Feuille technique, masquée : la liste des tables, pour le menu des lignes d'ajout.
FEUILLE_LISTES = "listes"

# Types de règle revus dans le formulaire (clés de config/controles.yaml), avec leur
# feuille. Le motif n'y est qu'affiché : une expression régulière n'est pas l'affaire
# du métier.
TYPES_FORMULAIRE = ("decimales", "bornes", "cle_luhn")
FEUILLES_TYPE = {
    "decimales": FEUILLE_DECIMALES,
    "bornes": FEUILLE_BORNES,
    "cle_luhn": FEUILLE_LUHN,
}
# Les déclarations de doublons : nom dans config/controles.yaml et nom affiché.
GROUPES_DOUBLONS = {"cle": "clé", "unique": "unique", "presume_unique": "présumé unique"}
OUI = "oui"
NON = "non"
AUCUNE = "aucune"
MASQUE = "masqué"
ORIGINE_EN_VIGUEUR = "en vigueur"
ORIGINE_PROPOSEE = "proposée"
ORIGINE_ECARTEE = "écartée"
STATUT_OBSERVE, STATUT_VALIDE, STATUT_INVALIDE, STATUT_DOCUMENTE = STATUTS
# Part minimale de clés de Luhn valides pour proposer la règle.
SEUIL_LUHN = 0.95
# Natures des colonnes où une clé candidate peut être proposée : un montant ou une
# date ne sont pas des identifiants, même sans répétition.
NATURES_CLE = ("texte", "entier (texte)")
LIGNES_AJOUT = 50
REPERE_EMPREINTE = "Empreinte de config/controles.yaml à l'export (repère technique)"

COLONNES_DECIMALES = (
    "table",
    "colonne",
    "nature",
    "valeurs",
    "observé",
    "proposition",
    "en_vigueur",
    "statut",
    "décimales",
    "commentaire",
)
COLONNES_BORNES = (
    "table",
    "colonne",
    "nature",
    "valeurs",
    "observé",
    "proposition",
    "en_vigueur",
    "statut",
    "min",
    "min_exclu",
    "max",
    "max_exclu",
    "commentaire",
    "proposition_min",
    "proposition_max",
)
COLONNES_LUHN = (
    "table",
    "colonne",
    "nature",
    "valeurs",
    "observé",
    "proposition",
    "en_vigueur",
    "statut",
    "commentaire",
)
COLONNES_DOUBLONS = ("table", "type", "colonnes", "origine", "observé", "statut", "commentaire")
COLONNES_TABLES = ("table", "lignes", "clé_en_vigueur", "identiques_hors_clé")
COLONNES_MOTIFS = ("table", "colonne", "motif")
COLONNES_TYPE = {
    "decimales": COLONNES_DECIMALES,
    "bornes": COLONNES_BORNES,
    "cle_luhn": COLONNES_LUHN,
}
MODIFIABLES = {
    "decimales": ("statut", "décimales", "commentaire"),
    "bornes": ("statut", "min", "min_exclu", "max", "max_exclu", "commentaire"),
    "cle_luhn": ("statut", "commentaire"),
}
MODIFIABLES_DOUBLONS = ("type", "statut", "commentaire")
SAISIE_AJOUT_DOUBLONS = ("table", "type", "colonnes", "statut", "commentaire")
# Colonnes techniques, masquées : la proposition des bornes en nombres, relue à
# l'import.
COLONNES_MASQUEES = ("proposition_min", "proposition_max")
LARGEURS = {
    "table": 22,
    "colonne": 26,
    "nature": 22,
    "valeurs": 10,
    "observé": 34,
    "proposition": 16,
    "en_vigueur": 16,
    "statut": 13,
    "décimales": 11,
    "min": 11,
    "min_exclu": 11,
    "max": 11,
    "max_exclu": 11,
    "commentaire": 50,
    "proposition_min": 8,
    "proposition_max": 8,
    "type": 16,
    "colonnes": 40,
    "origine": 12,
    "lignes": 10,
    "clé_en_vigueur": 34,
    "identiques_hors_clé": 20,
    "motif": 50,
}


def chemin_formulaire(periode_fin: str) -> Path:
    """Emplacement par défaut du formulaire : data/formulaires/, hors du dépôt."""
    return racine_projet() / "data" / "formulaires" / f"controles_{periode_fin}.xlsx"


def empreinte_controles(chemin: str | Path | None = None) -> str:
    """Empreinte du contenu de config/controles.yaml, « absent » sans fichier.

    Les fins de ligne sont ramenées à LF : une conversion par Git sous Windows ne
    change pas l'empreinte.
    """
    chemin = Path(chemin) if chemin is not None else chemin_regles_complementaires()
    if not chemin.is_file():
        return "absent"
    contenu = chemin.read_bytes().replace(b"\r\n", b"\n")
    return hashlib.sha256(contenu).hexdigest()[:16]


# --- Observations ---------------------------------------------------------------


def _decimales(nombre) -> int:
    """Nombre de décimales significatives d'un nombre lu par `controles._nombre`."""
    if isinstance(nombre, Decimal):
        return max(0, -nombre.normalize().as_tuple().exponent)
    return next(d for d in range(16) if not _trop_de_decimales(nombre, d))


def _arrondi(nombre):
    """Nombre écrit dans config/controles.yaml : un entier, ou un décimal arrondi."""
    if isinstance(nombre, Decimal):
        nombre = float(nombre)
    if float(nombre).is_integer():
        return int(nombre)
    return round(float(nombre), 10)


def texte_nombre(nombre) -> str:
    """Nombre écrit à la française (virgule décimale), sans zéro inutile."""
    texte = format(Decimal(repr(_arrondi(nombre))).normalize(), "f")
    return texte.replace(".", ",")


def observer_nombres(serie: pd.Series) -> dict:
    """Valeurs numériques d'une colonne : nombre, minimum, maximum et décimales.

    Seules comptent les valeurs que lisent les contrôles (nombre natif, ou texte à
    virgule ou point décimal, sans séparateur de milliers) : les autres relèvent de
    la règle `nature`. Sans valeur lisible, minimum, maximum et décimales sont None.
    """
    nombres = [n for n in map(_nombre, _valeurs_renseignees(serie)) if n is not None]
    if not nombres:
        return {"valeurs": 0, "min": None, "max": None, "decimales": None}
    return {
        "valeurs": len(nombres),
        "min": min(nombres),
        "max": max(nombres),
        "decimales": max(_decimales(n) for n in nombres),
    }


def observer_luhn(serie: pd.Series) -> dict:
    """Valeurs faites de chiffres d'une colonne, et part de celles dont la clé de Luhn
    est valide (None sans valeur de chiffres)."""
    chiffres = [
        texte
        for texte in (str(v) for v in _valeurs_renseignees(serie))
        if texte.isascii() and texte.isdigit()
    ]
    if not chiffres:
        return {"valeurs": 0, "part_valide": None}
    valides = sum(_cle_luhn_valide(texte) for texte in chiffres)
    return {"valeurs": len(chiffres), "part_valide": valides / len(chiffres)}


def observer_groupe(donnees: pd.DataFrame, colonnes: list[str]) -> dict:
    """Lignes d'une table où le groupe de colonnes est entièrement renseigné, et lignes
    dont les valeurs du groupe se répètent (comparées telles quelles)."""
    renseignees = pd.Series(True, index=donnees.index)
    for colonne in colonnes:
        renseignees &= donnees.index.isin(_valeurs_renseignees(donnees[colonne]).index)
    sous_table = donnees.loc[renseignees, colonnes].astype(str)
    return {
        "lignes": len(donnees),
        "renseignees": len(sous_table),
        "repetees": int(sous_table.duplicated(keep=False).sum()),
    }


def _texte_groupe(observation: dict) -> str:
    texte = f"{observation['renseignees']} lignes renseignées sur {observation['lignes']}"
    if observation["repetees"]:
        return f"{texte}, {observation['repetees']} lignes répétées"
    return f"{texte}, aucune répétition"


# --- Contenu de config/controles.yaml -------------------------------------------


def _section(contenu: dict, cle: str) -> dict:
    valeur = contenu.get(cle)
    return valeur if isinstance(valeur, dict) else {}


def _ecartees(contenu: dict, nom: str) -> dict:
    """Propositions écartées pour un type de règle ou les doublons, par table."""
    return _section(_section(contenu, CLE_ECARTEES), nom)


def _commentaires(contenu: dict, nom: str) -> dict:
    """Commentaires du métier pour un type de règle ou les doublons, par table."""
    return _section(_section(contenu, CLE_COMMENTAIRES), nom)


def texte_groupe(colonnes: list[str]) -> str:
    """Colonnes d'un groupe de doublons, telles que le formulaire les écrit."""
    return ", ".join(colonnes)


def _groupes_en_vigueur(declaration: dict) -> list[tuple[str, list[str]]]:
    groupes = []
    if "cle" in declaration:
        groupes.append(("cle", list(declaration["cle"])))
    for nom in ("unique", "presume_unique"):
        groupes.extend((nom, list(groupe)) for groupe in declaration.get(nom, []))
    return groupes


# --- Lignes du formulaire -------------------------------------------------------


def _colonnes_de_nature(dictionnaire: dict, natures) -> list[tuple[str, str, str]]:
    """(table, colonne, nature en vigueur) des colonnes d'une des natures, dans
    l'ordre du dictionnaire."""
    return [
        (table, colonne, regles["nature"]["regle"])
        for table, colonnes in dictionnaire.items()
        for colonne, regles in colonnes.items()
        if regles["nature"]["regle"] in natures
    ]


def _statut_export(en_vigueur, ecartee: bool) -> str:
    if en_vigueur is not None:
        return STATUT_DOCUMENTE
    return STATUT_INVALIDE if ecartee else STATUT_OBSERVE


def lignes_decimales(tables: dict, dictionnaire: dict, contenu: dict) -> list[dict]:
    """Lignes de la feuille des décimales : une par colonne de nature numérique.

    La proposition est le plus grand nombre de décimales observé; « aucune » sans
    valeur lisible.
    """
    regles = _section(contenu, CLE_REGLES_COMPLEMENTAIRES)
    ecartees = _ecartees(contenu, "decimales")
    commentaires = _commentaires(contenu, "decimales")
    lignes = []
    for table, colonne, nature in _colonnes_de_nature(dictionnaire, NATURES_NOMBRE):
        observation = observer_nombres(tables[table][colonne])
        en_vigueur = regles.get(table, {}).get(colonne, {}).get("decimales")
        decimales = observation["decimales"]
        lignes.append(
            {
                "table": table,
                "colonne": colonne,
                "nature": nature,
                "valeurs": observation["valeurs"],
                "observé": decimales,
                "proposition": AUCUNE if decimales is None else decimales,
                "en_vigueur": en_vigueur,
                "statut": _statut_export(en_vigueur, colonne in ecartees.get(table, [])),
                "décimales": en_vigueur,
                "commentaire": commentaires.get(table, {}).get(colonne),
            }
        )
    return lignes


def lignes_bornes(
    tables: dict, dictionnaire: dict, contenu: dict, sensibles: dict
) -> list[dict]:
    """Lignes de la feuille des bornes : une par colonne de nature numérique.

    La proposition est l'intervalle observé, bornes incluses; « aucune » sans valeur
    lisible ou pour une colonne sensible, dont le minimum et le maximum observés
    sont masqués.
    """
    regles = _section(contenu, CLE_REGLES_COMPLEMENTAIRES)
    ecartees = _ecartees(contenu, "bornes")
    commentaires = _commentaires(contenu, "bornes")
    lignes = []
    for table, colonne, nature in _colonnes_de_nature(dictionnaire, NATURES_NOMBRE):
        observation = observer_nombres(tables[table][colonne])
        en_vigueur = regles.get(table, {}).get(colonne, {}).get("bornes")
        sensible = est_sensible(sensibles, table, colonne)
        if sensible and observation["valeurs"]:
            observe, proposition, mini, maxi = MASQUE, AUCUNE, None, None
        elif observation["valeurs"]:
            mini, maxi = _arrondi(observation["min"]), _arrondi(observation["max"])
            observe = f"de {texte_nombre(mini)} à {texte_nombre(maxi)}"
            proposition = f"{texte_nombre(mini)} à {texte_nombre(maxi)}"
        else:
            observe, proposition, mini, maxi = None, AUCUNE, None, None
        en_vigueur = en_vigueur if isinstance(en_vigueur, dict) else None
        ligne = {
            "table": table,
            "colonne": colonne,
            "nature": nature,
            "valeurs": observation["valeurs"],
            "observé": observe,
            "proposition": proposition,
            "en_vigueur": (
                None if en_vigueur is None else _parametre_en_texte("bornes", en_vigueur)
            ),
            "statut": _statut_export(en_vigueur, colonne in ecartees.get(table, [])),
            "min": None,
            "min_exclu": None,
            "max": None,
            "max_exclu": None,
            "commentaire": commentaires.get(table, {}).get(colonne),
            "proposition_min": mini,
            "proposition_max": maxi,
        }
        for borne, exclusion in EXCLUSIONS.items():
            if en_vigueur is not None and borne in en_vigueur:
                ligne[borne] = en_vigueur[borne]
                ligne[exclusion] = OUI if en_vigueur.get(exclusion) else NON
        lignes.append(ligne)
    return lignes


def lignes_luhn(tables: dict, dictionnaire: dict, contenu: dict) -> list[dict]:
    """Lignes de la feuille des clés de Luhn : une par colonne « entier (texte) ».

    La règle est proposée quand au moins `SEUIL_LUHN` des valeurs faites de chiffres
    ont une clé valide (un SIRET, par exemple).
    """
    regles = _section(contenu, CLE_REGLES_COMPLEMENTAIRES)
    ecartees = _ecartees(contenu, "cle_luhn")
    commentaires = _commentaires(contenu, "cle_luhn")
    lignes = []
    for table, colonne, nature in _colonnes_de_nature(dictionnaire, ("entier (texte)",)):
        observation = observer_luhn(tables[table][colonne])
        en_vigueur = OUI if regles.get(table, {}).get(colonne, {}).get("cle_luhn") else None
        part = observation["part_valide"]
        lignes.append(
            {
                "table": table,
                "colonne": colonne,
                "nature": nature,
                "valeurs": observation["valeurs"],
                "observé": (
                    None if part is None
                    else f"{texte_nombre(round(100 * part, 1))} % de clés valides"
                ),
                "proposition": OUI if part is not None and part >= SEUIL_LUHN else AUCUNE,
                "en_vigueur": en_vigueur,
                "statut": _statut_export(en_vigueur, colonne in ecartees.get(table, [])),
                "commentaire": commentaires.get(table, {}).get(colonne),
            }
        )
    return lignes


def _candidates(donnees: pd.DataFrame, dictionnaire_table: dict) -> list[str]:
    """Colonnes d'une table toujours renseignées et sans valeur répétée."""
    if len(donnees) < 2:
        return []
    candidates = []
    for colonne, regles in dictionnaire_table.items():
        if regles["nature"]["regle"] not in NATURES_CLE:
            continue
        observation = observer_groupe(donnees, [colonne])
        if observation["renseignees"] == len(donnees) and not observation["repetees"]:
            candidates.append(colonne)
    return candidates


def lignes_doublons(tables: dict, dictionnaire: dict, contenu: dict) -> list[dict]:
    """Lignes de la feuille des doublons : un groupe de colonnes par ligne.

    Les groupes en vigueur d'abord (statut « documenté »), puis les clés candidates
    (une colonne toujours renseignée, sans valeur répétée, de nature texte ou entier)
    qui ne sont pas déjà déclarées seules, au statut « observé », puis les groupes
    écartés (statut « invalide »). Une clé candidate est proposée comme clé si la
    table n'en a pas, comme unique sinon.
    """
    declarations = _section(contenu, CLE_DOUBLONS)
    ecartees = _ecartees(contenu, "doublons")
    commentaires = _commentaires(contenu, "doublons")
    lignes = []
    for table, donnees in tables.items():
        if table not in dictionnaire:
            continue
        en_vigueur = _groupes_en_vigueur(declarations.get(table, {}))
        refuses = [list(groupe) for groupe in ecartees.get(table, [])]
        connus = {frozenset(groupe) for _, groupe in en_vigueur} | {
            frozenset(groupe) for groupe in refuses
        }
        avec_cle = any(nom == "cle" for nom, _ in en_vigueur)
        proposes = [
            ("unique" if avec_cle else "cle", [colonne])
            for colonne in _candidates(donnees, dictionnaire[table])
            if frozenset([colonne]) not in connus
        ]
        for origine, statut, groupes in (
            (ORIGINE_EN_VIGUEUR, STATUT_DOCUMENTE, en_vigueur),
            (ORIGINE_PROPOSEE, STATUT_OBSERVE, proposes),
            (ORIGINE_ECARTEE, STATUT_INVALIDE, [("unique", groupe) for groupe in refuses]),
        ):
            for nom, groupe in groupes:
                presentes = [colonne for colonne in groupe if colonne in donnees.columns]
                observation = (
                    _texte_groupe(observer_groupe(donnees, groupe))
                    if len(presentes) == len(groupe)
                    else None
                )
                lignes.append(
                    {
                        "table": table,
                        "type": GROUPES_DOUBLONS[nom],
                        "colonnes": texte_groupe(groupe),
                        "origine": origine,
                        "observé": observation,
                        "statut": statut,
                        "commentaire": commentaires.get(table, {}).get(texte_groupe(groupe)),
                    }
                )
    return lignes


def lignes_tables(tables: dict, dictionnaire: dict, contenu: dict) -> list[dict]:
    """Lignes de la feuille des tables : faut-il chercher les lignes identiques hors
    de la clé (« non » quand deux enregistrements peuvent y être pareils) ?"""
    declarations = _section(contenu, CLE_DOUBLONS)
    return [
        {
            "table": table,
            "lignes": len(donnees),
            "clé_en_vigueur": texte_groupe(declarations.get(table, {}).get("cle", [])) or None,
            "identiques_hors_clé": (
                NON if declarations.get(table, {}).get("identique_hors_cle") is False else OUI
            ),
        }
        for table, donnees in tables.items()
        if table in dictionnaire
    ]


def lignes_motifs(contenu: dict) -> list[dict]:
    """Motifs en vigueur, affichés en lecture seule."""
    return [
        {"table": table, "colonne": colonne, "motif": regles["motif"]}
        for table, colonnes in _section(contenu, CLE_REGLES_COMPLEMENTAIRES).items()
        for colonne, regles in (colonnes or {}).items()
        if isinstance(regles, dict) and "motif" in regles
    ]


# --- Export ---------------------------------------------------------------------


def _a_traiter(lignes: list[dict]) -> int:
    return sum(ligne["statut"] == STATUT_OBSERVE for ligne in lignes)


def _mode_emploi(periode_fin: str, lignes: dict) -> list[str]:
    decimales, bornes, luhn = lignes["decimales"], lignes["bornes"], lignes["cle_luhn"]
    doublons = lignes["doublons"]
    return [
        "Formulaire des règles de contrôle",
        f"Période de fin : {periode_fin}. Formulaire généré le {date.today():%d/%m/%Y}.",
        "Ce formulaire donne les règles de contrôle que le dictionnaire des données ne "
        "peut pas déduire des fichiers : nombre de décimales, bornes, clé de contrôle "
        "de Luhn, et les colonnes qui ne doivent pas se répéter. Il montre les règles en "
        "vigueur (colonne en_vigueur, statut documenté) à côté de ce que l'outil observe "
        "dans les données (colonnes observé et proposition). Une fois revu, il réécrit "
        "le fichier des règles de contrôle.",
        "",
        "Ce qui reste à traiter",
        f"Feuille « {FEUILLE_DECIMALES} » : {_a_traiter(decimales)} lignes à « observé » "
        f"sur {len(decimales)}. Feuille « {FEUILLE_BORNES} » : {_a_traiter(bornes)} sur "
        f"{len(bornes)}. Feuille « {FEUILLE_LUHN} » : {_a_traiter(luhn)} sur {len(luhn)}. "
        f"Feuille « {FEUILLE_DOUBLONS} » : {_a_traiter(doublons)} sur {len(doublons)}.",
        "Une revue partielle est possible : une ligne laissée à « observé » est "
        "reprise telle quelle dans le formulaire suivant.",
        "",
        "Ce qui est attendu",
        "Pour chaque ligne, choisir un statut dans le menu déroulant et, si besoin, "
        "écrire un commentaire (il est conservé).",
        "observé : la ligne n'est pas revue; pas de règle.",
        "valide : la proposition devient la règle. Impossible si la proposition est "
        f"« {AUCUNE} ».",
        "documenté : vous donnez vous-même la règle, dans les colonnes à saisir de la "
        "ligne. Une règle en vigueur sort avec ce statut et ses valeurs : pour la "
        "modifier, changer les valeurs en gardant documenté.",
        "invalide : pas de règle sur cette colonne, et la proposition ne sera plus "
        "faite. Pour une règle en vigueur, invalide la retire.",
        "Remettre « observé » sur une règle en vigueur la retire aussi, mais la ligne "
        "sera de nouveau à traiter au formulaire suivant.",
        "Une valeur saisie sur une ligne qui n'est pas au statut documenté est ignorée "
        "(un avertissement le signale à l'import).",
        "",
        f"Feuille « {FEUILLE_DECIMALES} » : le nombre maximal de décimales d'un nombre",
        "Une ligne par colonne numérique, codes compris : pour un code ou un "
        "identifiant (SIRET, compte), choisir invalide. La proposition est le plus grand "
        "nombre de décimales observé. Colonne à saisir : décimales.",
        "",
        f"Feuille « {FEUILLE_BORNES} » : le minimum et le maximum admis",
        "La proposition est le plus petit et le plus grand nombre observés, bornes "
        "incluses. Attention : si les données contiennent une erreur (un taux de 80 au "
        "lieu de 0,8), elle fait partie de la proposition. Colonnes à saisir : min, max "
        "(l'un ou l'autre suffit), et pour chacun min_exclu ou max_exclu (oui si la "
        "borne elle-même est interdite : « plus de 0 » s'écrit min 0 et min_exclu oui). "
        f"Pour une colonne sensible, le minimum et le maximum observés sont « {MASQUE} » "
        "et rien n'est proposé : donner les bornes au statut documenté.",
        "",
        f"Feuille « {FEUILLE_LUHN} » : la clé de contrôle de Luhn",
        "Une ligne par colonne de chiffres. La règle est proposée quand au moins "
        f"{texte_nombre(100 * SEUIL_LUHN)} % des valeurs ont une clé valide, ce qui est "
        "le cas d'un SIRET ou d'un SIREN corrects.",
        "",
        f"Feuille « {FEUILLE_DOUBLONS} » : les colonnes qui ne doivent pas se répéter",
        "Une ligne par groupe de colonnes : clé (l'identifiant d'une ligne, une seule "
        "par table), unique (une répétition est une anomalie), présumé unique (une "
        "répétition est à arbitrer). Le type se change dans le menu de la colonne type. "
        "Les groupes en vigueur sont au statut documenté. L'outil propose les colonnes "
        "toujours renseignées et sans répétition (statut observé) : valide les retient "
        "avec le type affiché.",
        "Pour ajouter un groupe de plusieurs colonnes : sur une ligne vide en bas de la "
        "feuille, choisir la table et le type dans les menus, écrire les colonnes "
        "séparées par une virgule (par exemple « raison_sociale, cp ») et choisir le "
        "statut documenté.",
        "",
        f"Feuille « {FEUILLE_TABLES} » : les lignes identiques hors de la clé",
        "oui : deux lignes pareilles en tout sauf la clé sont signalées. non : elles "
        "sont normales dans cette table (deux enregistrements peuvent être pareils). Ne "
        "s'applique qu'à une table qui a une clé.",
        "",
        f"Feuille « {FEUILLE_MOTIFS} » : lecture seule",
        "Les motifs (expressions régulières) sont écrits par le consultant et repris "
        "tels quels.",
        "",
        "Règles à respecter pour que le formulaire puisse être réimporté",
        "Ne modifier que les cellules non verrouillées. Ne pas renommer, déplacer ni "
        "supprimer de feuilles, de colonnes ou de lignes. Avant l'import : enregistrer "
        "puis fermer le classeur dans Excel. Un formulaire ne s'applique qu'une fois : "
        "après son import, ou après une modification du fichier des règles, il est à "
        "réexporter.",
        REPERE_EMPREINTE,
    ]


def _remplir(feuille, colonnes, lignes, modifiables, lignes_ajout=0, saisie_ajout=()) -> None:
    """Écrit l'en-tête et les lignes, puis protège la feuille hors cellules modifiables."""
    for indice, nom in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=1, column=indice, value=nom)
        cellule.font = Font(bold=True)
        lettre = cellule.column_letter
        feuille.column_dimensions[lettre].width = LARGEURS[nom]
        if nom in COLONNES_MASQUEES:
            feuille.column_dimensions[lettre].hidden = True
    for numero, ligne in enumerate(lignes, start=2):
        for indice, nom in enumerate(colonnes, start=1):
            cellule = feuille.cell(row=numero, column=indice)
            valeur = ligne.get(nom)
            if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
                cellule.value = valeur
            else:
                _ecrire_texte(cellule, valeur)
            if nom in modifiables:
                cellule.protection = Protection(locked=False)
    fin = len(lignes) + 1
    for numero in range(fin + 1, fin + lignes_ajout + 1):
        for indice, nom in enumerate(colonnes, start=1):
            if nom in saisie_ajout:
                cellule = feuille.cell(row=numero, column=indice)
                cellule.number_format = "@"
                cellule.protection = Protection(locked=False)
    derniere = get_column_letter(len(colonnes))
    feuille.freeze_panes = "C2"
    feuille.auto_filter.ref = f"A1:{derniere}{max(fin + lignes_ajout, 2)}"
    feuille.protection.sheet = True
    feuille.protection.autoFilter = False
    feuille.protection.formatColumns = False


def _colonne(colonnes, nom) -> int:
    return colonnes.index(nom) + 1


def _ajouter_menu(feuille, menu, colonnes, nom, premiere, derniere) -> None:
    if derniere >= premiere:
        lettre = get_column_letter(_colonne(colonnes, nom))
        menu.add(f"{lettre}{premiere}:{lettre}{derniere}")


def exporter_formulaire(
    tables: dict,
    dictionnaire: dict,
    contenu: dict,
    sensibles: dict,
    periode_fin: str,
    chemin: str | Path | None = None,
    *,
    empreinte: str,
    ecraser: bool = False,
) -> Path:
    """Écrit le formulaire des règles de contrôle et rend son chemin.

    `contenu` est celui de config/controles.yaml (`controles.lire_fichier_controles`)
    et `empreinte` son empreinte (`empreinte_controles`) : elle est inscrite dans le
    formulaire, et l'import le refusera si le fichier a changé depuis. Sans
    `ecraser=True`, refuse de remplacer un formulaire existant : une revue y est
    peut-être en cours.
    """
    chemin = Path(chemin) if chemin is not None else chemin_formulaire(periode_fin)
    if chemin.exists() and not ecraser:
        raise FileExistsError(
            f"{chemin} existe déjà : export refusé (ecraser=True pour forcer)"
        )
    lignes = {
        "decimales": lignes_decimales(tables, dictionnaire, contenu),
        "bornes": lignes_bornes(tables, dictionnaire, contenu, sensibles),
        "cle_luhn": lignes_luhn(tables, dictionnaire, contenu),
        "doublons": lignes_doublons(tables, dictionnaire, contenu),
    }
    noms_tables = [table for table in tables if table in dictionnaire]

    classeur = Workbook()
    notice = classeur.active
    notice.title = FEUILLE_MODE_EMPLOI
    notice.column_dimensions["A"].width = 120
    rubriques = ("Ce qui", "Feuille", "Règles à", "Formulaire des")
    for numero, texte in enumerate(_mode_emploi(periode_fin, lignes), start=1):
        cellule = notice.cell(row=numero, column=1)
        _ecrire_texte(cellule, texte)
        cellule.alignment = Alignment(wrap_text=True, vertical="top")
        if texte == REPERE_EMPREINTE:
            _ecrire_texte(notice.cell(row=numero, column=2), empreinte)
        if texte.startswith(rubriques):
            cellule.font = Font(bold=True)
    notice.protection.sheet = True

    for nom in TYPES_FORMULAIRE:
        colonnes = COLONNES_TYPE[nom]
        feuille = classeur.create_sheet(FEUILLES_TYPE[nom])
        _remplir(feuille, colonnes, lignes[nom], MODIFIABLES[nom])
        derniere = len(lignes[nom]) + 1
        _ajouter_menu(feuille, _menu(feuille, STATUTS), colonnes, "statut", 2, derniere)
        if nom == "bornes":
            menu = _menu(feuille, (OUI, NON))
            for exclusion in EXCLUSIONS.values():
                _ajouter_menu(feuille, menu, colonnes, exclusion, 2, derniere)

    feuille = classeur.create_sheet(FEUILLE_DOUBLONS)
    doublons = lignes["doublons"]
    _remplir(
        feuille, COLONNES_DOUBLONS, doublons, MODIFIABLES_DOUBLONS, LIGNES_AJOUT,
        SAISIE_AJOUT_DOUBLONS,
    )
    derniere = len(doublons) + 1 + LIGNES_AJOUT
    _ajouter_menu(feuille, _menu(feuille, STATUTS), COLONNES_DOUBLONS, "statut", 2, derniere)
    _ajouter_menu(
        feuille, _menu(feuille, tuple(GROUPES_DOUBLONS.values())), COLONNES_DOUBLONS,
        "type", 2, derniere,
    )
    if noms_tables:
        plage = f"{quote_sheetname(FEUILLE_LISTES)}!$A$1:$A${len(noms_tables)}"
        _ajouter_menu(
            feuille, _menu(feuille, formule=plage), COLONNES_DOUBLONS, "table",
            len(doublons) + 2, derniere,
        )

    feuille = classeur.create_sheet(FEUILLE_TABLES)
    lignes_t = lignes_tables(tables, dictionnaire, contenu)
    _remplir(feuille, COLONNES_TABLES, lignes_t, ("identiques_hors_clé",))
    _ajouter_menu(
        feuille, _menu(feuille, (OUI, NON)), COLONNES_TABLES, "identiques_hors_clé", 2,
        len(lignes_t) + 1,
    )

    feuille = classeur.create_sheet(FEUILLE_MOTIFS)
    _remplir(feuille, COLONNES_MOTIFS, lignes_motifs(contenu), ())

    listes = classeur.create_sheet(FEUILLE_LISTES)
    for numero, table in enumerate(noms_tables, start=1):
        _ecrire_texte(listes.cell(row=numero, column=1), table)
    listes.protection.sheet = True
    listes.sheet_state = "hidden"

    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    return chemin
