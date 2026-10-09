"""Tests de l'export, de l'import et de l'application du classeur de revue (frigeo.revue)."""

from copy import deepcopy

import pytest
from openpyxl import load_workbook

from frigeo.dictionnaire import (
    charger_dictionnaire,
    charger_meta,
    dossier_sauvegardes,
    ecrire_dictionnaire,
    lire_instant,
)
from frigeo.revue import (
    COLONNES_FORMES,
    COLONNES_RAPPORT_REVUE,
    COLONNES_REGLES,
    COLONNES_VALEURS,
    LIGNES_AJOUT,
    REPERE_REVUES,
    appliquer_revue,
    exporter_revue,
    importer_revue,
    lignes_formes,
    lignes_regles,
    lignes_valeurs,
    reste_a_traiter,
)

PIEGES = ["076", "27", "2026-03-01", "=1+1", "Non", "Gisors  ", "a b"]


def _regle(regle, a_arbitrer, statut="observé"):
    return {
        "regle": regle,
        "proposition": regle,
        "statut": statut,
        "motif": "motif fictif",
        "a_arbitrer": a_arbitrer,
        "commentaire": "",
        "revu_le": None,
    }


def _valeur(valeur, origine, effectif):
    return {
        "valeur": valeur,
        "origine": origine,
        "effectif": effectif,
        "statut": "observé",
        "commentaire": "",
        "revu_le": None,
        "remplacement": None,
    }


def _colonne(valeurs=None, a_arbitrer=None, vides=0, hors_nature=0):
    """Colonne avec une liste fermée si `valeurs` est donnée (effectif 10 par valeur)."""
    a_arbitrer = a_arbitrer or {}
    liste = _regle(
        "liste fermée" if valeurs else "aucune",
        {"valeurs": len(a_arbitrer)} if valeurs else {},
    )
    liste["liste"] = [_valeur(v, "liste proposée", 10) for v in valeurs or []] + [
        _valeur(v, "à arbitrer", n) for v, n in a_arbitrer.items()
    ]
    return {
        "obligatoire": _regle("obligatoire", {"vides": vides}),
        "nature": _regle("texte", {"hors_nature": hors_nature}),
        "valeurs": liste,
        # Aucun format proposé : voir `_avec_formats` pour les formats et leurs formes.
        "format": {**_regle("aucune", {}), "liste": []},
    }


def _dictionnaire():
    duree = _colonne()
    duree["obligatoire"] = _regle("à décider", {"vides": 1})
    duree["nature"] = _regle(
        "à décider", {"nature_dominante": "entier (texte)", "hors_nature": 40}
    )
    return {
        "clients": {
            "statut": _colonne(["Actif", "Inactif"], {"actif": 3}, vides=5),
            "code": _colonne(PIEGES),
            "type_commerce": _colonne(),
        },
        "interventions": {"duree": duree},
    }


def _lire(feuille):
    lignes = list(feuille.iter_rows(values_only=True))
    return [dict(zip(lignes[0], ligne)) for ligne in lignes[1:] if any(v is not None for v in ligne)]


def test_lignes_regles():
    lignes = lignes_regles(_dictionnaire())
    cles = [(l["table"], l["colonne"], l["regle"]) for l in lignes]
    # Quatre lignes par colonne, y compris sans liste ni format proposé.
    assert len(lignes) == 4 * 4
    assert {l["regle"] for l in lignes} == {"obligatoire", "nature", "valeurs", "format"}
    assert ("clients", "statut", "valeurs") in cles
    par_cle = dict(zip(cles, lignes))
    sans_liste = par_cle[("clients", "type_commerce", "valeurs")]
    assert (sans_liste["proposition"], sans_liste["regle_en_vigueur"]) == ("aucune", "aucune")
    assert (sans_liste["statut"], sans_liste["regle_retenue"]) == ("observé", None)
    assert par_cle[("clients", "statut", "obligatoire")]["nb_a_arbitrer"] == 5
    assert par_cle[("clients", "statut", "valeurs")]["proposition"] == "liste de 2 valeurs"
    assert par_cle[("clients", "statut", "valeurs")]["nb_a_arbitrer"] == 1
    # Le nom de la nature dominante n'est pas compté comme un écart.
    assert par_cle[("interventions", "duree", "nature")]["nb_a_arbitrer"] == 40


def test_lignes_valeurs():
    lignes = lignes_valeurs(_dictionnaire())
    assert len(lignes) == 3 + len(PIEGES)
    assert {l["colonne"] for l in lignes} == {"statut", "code"}
    # L'effectif et l'origine viennent du dictionnaire.
    assert [l["effectif"] for l in lignes[:3]] == [10, 10, 3]
    assert [l["origine"] for l in lignes[:3]] == ["liste proposée", "liste proposée", "à arbitrer"]
    remarques = {l["valeur"]: l["remarque"] for l in lignes}
    assert remarques["Gisors  "] == "espaces en bord"
    assert remarques["a\u00a0b"] == "espace insécable"
    assert remarques["Actif"] == ""


def test_statut_propre_a_chaque_valeur():
    dictionnaire = _dictionnaire()
    liste = dictionnaire["clients"]["statut"]["valeurs"]["liste"]
    liste[2].update(statut="invalide", revu_le="2026-10-08")
    lignes = lignes_valeurs(dictionnaire)
    assert [l["statut"] for l in lignes[:3]] == ["observé", "observé", "invalide"]


def test_classeur_ecrit_et_relu(tmp_path):
    chemin = exporter_revue(_dictionnaire(), "2026-03", tmp_path / "revue" / "revue.xlsx", nb_revues=0)
    classeur = load_workbook(chemin)
    assert classeur.sheetnames == ["mode_emploi", "regles", "valeurs", "formes", "listes"]
    regles, valeurs = classeur["regles"], classeur["valeurs"]
    assert tuple(c.value for c in regles[1]) == COLONNES_REGLES
    assert tuple(c.value for c in valeurs[1]) == COLONNES_VALEURS
    assert tuple(c.value for c in classeur["formes"][1]) == COLONNES_FORMES
    assert len(_lire(regles)) == 16
    assert {l["statut"] for l in _lire(regles)} == {"observé"}
    notice = " ".join(str(c.value) for c in classeur["mode_emploi"]["A"] if c.value)
    assert "2026-03" in notice
    assert "16 lignes dans la feuille « regles » (quatre par colonne" in notice
    assert "0 lignes dans la feuille « formes »" in notice
    assert "réimporté" in notice


def test_valeurs_conservees_comme_texte(tmp_path):
    chemin = exporter_revue(_dictionnaire(), "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    lues = [l for l in _lire(load_workbook(chemin)["valeurs"]) if l["colonne"] == "code"]
    assert [l["valeur"] for l in lues] == PIEGES
    assert all(isinstance(l["valeur"], str) for l in lues)


def test_protection_et_menus(tmp_path):
    chemin = exporter_revue(_dictionnaire(), "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    classeur = load_workbook(chemin)
    regles, valeurs = classeur["regles"], classeur["valeurs"]
    assert regles.protection.sheet and valeurs.protection.sheet
    assert classeur["mode_emploi"].protection.sheet
    # Feuille des règles : seules statut, regle_retenue et commentaire sont libres.
    libres = {c.value for c, d in zip(regles[1], regles[2]) if not d.protection.locked}
    assert libres == {"statut", "regle_retenue", "commentaire"}
    # Sur une ligne « valeurs », regle_retenue sert à déclarer une liste fermée.
    ligne = next(l for l in regles.iter_rows(min_row=2) if l[2].value == "valeurs")
    assert not ligne[COLONNES_REGLES.index("regle_retenue")].protection.locked
    libres = {c.value for c, d in zip(valeurs[1], valeurs[2]) if not d.protection.locked}
    assert libres == {"statut", "remplacement", "commentaire"}
    # Lignes d'ajout en bas de la feuille des valeurs : pas de remplacement à saisir.
    ajout = valeurs[3 + len(PIEGES) + 2]
    libres = {c.value for c, d in zip(valeurs[1], ajout) if not d.protection.locked}
    assert libres == {"table", "colonne", "valeur", "statut", "commentaire"}
    assert valeurs.max_row == 1 + 3 + len(PIEGES) + LIGNES_AJOUT
    menus = [m.formula1 for m in regles.data_validations.dataValidation]
    assert '"observé,valide,invalide,documenté"' in menus
    assert '"obligatoire,facultatif,toujours vide"' in menus
    assert '"liste fermée"' in menus
    assert '"formes fermées"' in menus
    assert all(len(m) <= 255 for m in menus)
    # Feuille des valeurs : menu des statuts, puis menus table et colonne des ajouts.
    assert len(valeurs.data_validations.dataValidation) == 3
    # Excel refuse une valeur saisie hors d'un menu.
    for feuille in (regles, valeurs):
        assert all(m.showErrorMessage for m in feuille.data_validations.dataValidation)
    # La ligne d'en-tête et les colonnes table et colonne restent visibles.
    assert regles.freeze_panes == valeurs.freeze_panes == "C2"


def test_colonne_sans_liste_absente_de_la_feuille_des_valeurs(tmp_path):
    chemin = exporter_revue(_dictionnaire(), "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    colonnes = {l["colonne"] for l in _lire(load_workbook(chemin)["valeurs"])}
    assert "type_commerce" not in colonnes


def test_export_refuse_d_ecraser_un_classeur_existant(tmp_path):
    chemin = exporter_revue(_dictionnaire(), "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    with pytest.raises(FileExistsError, match="existe déjà"):
        exporter_revue(_dictionnaire(), "2026-03", chemin, nb_revues=0)
    exporter_revue(_dictionnaire(), "2026-03", chemin, nb_revues=0, ecraser=True)


def _exporter(tmp_path):
    return exporter_revue(_dictionnaire(), "2026-03", tmp_path / "revue.xlsx", nb_revues=0)


def _ligne(feuille, **cles):
    """Numéro de la ligne Excel dont les cellules nommées ont les valeurs données."""
    entete = [c.value for c in feuille[1]]
    for ligne in feuille.iter_rows(min_row=2):
        if all(ligne[entete.index(nom)].value == valeur for nom, valeur in cles.items()):
            return ligne[0].row
    raise KeyError(cles)


def _modifier(chemin, nom_feuille, ligne, **valeurs):
    classeur = load_workbook(chemin)
    feuille = classeur[nom_feuille]
    entete = [c.value for c in feuille[1]]
    numero = ligne if isinstance(ligne, int) else _ligne(feuille, **ligne)
    for nom, valeur in valeurs.items():
        feuille.cell(row=numero, column=entete.index(nom) + 1).value = valeur
    classeur.save(chemin)


def _erreur_import(chemin):
    with pytest.raises(ValueError) as erreur:
        importer_revue(chemin, _dictionnaire(), nb_revues=0)
    return str(erreur.value)


def test_import_sans_modification(tmp_path):
    revue = importer_revue(_exporter(tmp_path), _dictionnaire(), nb_revues=0)
    assert len(revue.regles) == 16
    assert revue.formes.empty
    assert len(revue.valeurs) == 3 + len(PIEGES)
    assert set(revue.regles["statut"]) == set(revue.valeurs["statut"]) == {"observé"}
    assert list(revue.valeurs.loc[revue.valeurs["colonne"] == "code", "valeur"]) == PIEGES
    assert revue.avertissements == []


def test_import_des_decisions_du_metier(tmp_path):
    chemin = _exporter(tmp_path)
    regle = dict(table="clients", colonne="statut")
    _modifier(chemin, "regles", dict(regle, regle="obligatoire"), statut="valide")
    # Statut saisi en forme décomposée, avec une espace en trop : il est normalisé.
    _modifier(chemin, "regles", dict(regle, regle="nature"), statut=" documenté",
              regle_retenue="entier (texte)")
    _modifier(chemin, "regles", dict(regle, regle="valeurs"), statut="invalide",
              commentaire="clé vers un référentiel")
    _modifier(chemin, "valeurs", dict(regle, valeur="actif"), statut="invalide")
    # Ajouts : dans une colonne avec liste et dans une colonne sans liste.
    fin = 1 + 3 + len(PIEGES)
    _modifier(chemin, "valeurs", fin + 1, table="clients", colonne="statut",
              valeur="En sommeil", statut="documenté")
    _modifier(chemin, "valeurs", fin + 5, table="clients", colonne="type_commerce",
              valeur="Boulangerie", statut="documenté", commentaire="chapitre 4")

    revue = importer_revue(chemin, _dictionnaire(), nb_revues=0)
    regles = revue.regles.set_index(["colonne", "regle"])
    assert regles.loc[("statut", "obligatoire"), "statut"] == "valide"
    assert regles.loc[("statut", "nature"), "statut"] == "documenté"
    assert regles.loc[("statut", "nature"), "regle_retenue"] == "entier (texte)"
    ajouts = revue.valeurs[revue.valeurs["origine"] == "ajoutée"]
    assert list(ajouts["valeur"]) == ["En sommeil", "Boulangerie"]
    assert list(ajouts["ligne_excel"]) == [fin + 1, fin + 5]
    assert revue.avertissements == [
        "clients.statut : liste écartée, mais 1 valeurs revues "
        "(statuts conservés, sans effet tant que la liste est écartée)"
    ]


def test_problemes_de_la_feuille_regles_listes_en_une_fois(tmp_path):
    chemin = _exporter(tmp_path)
    statut = dict(table="clients", colonne="statut")
    duree = dict(table="interventions", colonne="duree")
    _modifier(chemin, "regles", dict(statut, regle="obligatoire"), statut=None)
    _modifier(chemin, "regles", dict(statut, regle="nature"), statut="ok")
    _modifier(chemin, "regles", dict(duree, regle="obligatoire"), statut="valide")
    _modifier(chemin, "regles", dict(duree, regle="nature"), statut="documenté")
    _modifier(chemin, "regles", dict(table="clients", colonne="code", regle="obligatoire"),
              statut="documenté", regle_retenue="presque toujours vide")
    _modifier(chemin, "regles", dict(table="clients", colonne="code", regle="nature"),
              statut="valide", regle_retenue="texte")
    _modifier(chemin, "regles", dict(statut, regle="valeurs"), statut="invalide")
    _modifier(chemin, "regles", dict(table="clients", colonne="code", regle="valeurs"),
              statut="documenté")
    classeur = load_workbook(chemin)
    feuille = classeur["regles"]
    feuille.delete_rows(_ligne(feuille, table="clients", colonne="type_commerce", regle="nature"))
    classeur.save(chemin)

    message = _erreur_import(chemin)
    assert "9 problèmes" in message
    for attendu in (
        "statut vide",
        "statut inconnu 'ok'",
        "ne peut pas être validée",
        "regle_retenue obligatoire avec le statut documenté",
        "'presque toujours vide' hors du menu",
        "regle_retenue réservée au statut documenté",
        "commentaire obligatoire avec le statut invalide",
        "une liste est déjà proposée, il n'y a rien à déclarer",
        "('clients', 'type_commerce', 'nature') absente du classeur",
    ):
        assert attendu in message
    assert "regles, ligne " in message


def test_classeur_perime_refuse(tmp_path):
    chemin = _exporter(tmp_path)
    dictionnaire = _dictionnaire()
    dictionnaire["clients"]["statut"]["obligatoire"]["proposition"] = "facultatif"
    dictionnaire["clients"]["statut"]["valeurs"]["liste"].append(
        _valeur("Radié", "à arbitrer", 1)
    )
    with pytest.raises(ValueError) as erreur:
        importer_revue(chemin, dictionnaire, nb_revues=0)
    message = str(erreur.value)
    assert "classeur périmé, à réexporter" in message
    assert "('clients', 'statut', 'Radié') absente du classeur" in message


def test_problemes_de_la_feuille_valeurs(tmp_path):
    chemin = _exporter(tmp_path)
    fin = 1 + 3 + len(PIEGES)
    _modifier(chemin, "valeurs", dict(table="clients", colonne="statut", valeur="Actif"),
              statut="documenté")
    _modifier(chemin, "valeurs", fin + 1, table="clients", colonne="statut",
              valeur="Archivé", statut="valide")
    _modifier(chemin, "valeurs", fin + 2, table="clients", colonne="inconnue",
              valeur="X", statut="documenté")
    _modifier(chemin, "valeurs", fin + 3, table="clients", colonne="statut",
              valeur="Inactif", statut="documenté")
    _modifier(chemin, "valeurs", fin + 4, table="clients", colonne="code",
              valeur=27, statut="documenté")
    _modifier(chemin, "valeurs", fin + 5, table="clients", colonne="statut",
              statut="documenté")
    message = _erreur_import(chemin)
    assert "6 problèmes" in message
    for attendu in (
        "statut 'documenté' impossible sur une valeur observée",
        "une valeur ajoutée doit porter le statut documenté",
        "colonne clients.inconnue absente du dictionnaire",
        "valeur 'Inactif' déjà présente pour clients.statut",
        "valeur 27 lue comme int",
        "valeur vide",
    ):
        assert attendu in message


@pytest.mark.parametrize("cas", ["feuille", "entete"])
def test_structure_modifiee_refusee(tmp_path, cas):
    chemin = _exporter(tmp_path)
    classeur = load_workbook(chemin)
    if cas == "feuille":
        classeur["valeurs"].title = "Valeurs"
        attendu = "feuille « valeurs » absente"
    else:
        classeur["regles"]["D1"] = "proposée"
        attendu = "en-tête modifié"
    classeur.save(chemin)
    assert attendu in _erreur_import(chemin)


# Décisions enregistrées, remplacements, classeur périmé et application (morceau 4c).

JOUR = "2026-10-09 10:15:00"
STATUT = dict(table="clients", colonne="statut")
FIN = 1 + 3 + len(PIEGES)  # dernière ligne exportée de la feuille des valeurs


def _importer(chemin, dictionnaire=None, nb_revues=0):
    return importer_revue(chemin, dictionnaire or _dictionnaire(), nb_revues=nb_revues)


def _decisions(chemin):
    """Saisit dans le classeur un jeu de décisions couvrant tous les statuts."""
    _modifier(chemin, "regles", dict(STATUT, regle="obligatoire"), statut="valide")
    _modifier(chemin, "regles", dict(STATUT, regle="nature"), statut="documenté",
              regle_retenue="entier (texte)", commentaire="code numérique")
    _modifier(chemin, "regles", dict(table="clients", colonne="code", regle="valeurs"),
              statut="invalide", commentaire="identifiant libre")
    _modifier(chemin, "regles", dict(table="interventions", colonne="duree", regle="obligatoire"),
              statut="documenté", regle_retenue="facultatif")
    _modifier(chemin, "valeurs", dict(STATUT, valeur="Actif"), statut="valide")
    _modifier(chemin, "valeurs", dict(STATUT, valeur="actif"), statut="invalide",
              remplacement="Actif", commentaire="casse")
    _modifier(chemin, "valeurs", FIN + 1, table="clients", colonne="statut",
              valeur="En sommeil", statut="documenté", commentaire="chapitre 4")
    _modifier(chemin, "valeurs", FIN + 2, table="clients", colonne="type_commerce",
              valeur="Boulangerie", statut="documenté")


def _apres_premiere_revue(tmp_path):
    """Dictionnaire obtenu en appliquant le jeu de décisions, et l'entrée de journal."""
    chemin = _exporter(tmp_path)
    _decisions(chemin)
    resultat, rapport, entree = appliquer_revue(_dictionnaire(), _importer(chemin), chemin, JOUR)
    return resultat, rapport, entree


def _element(dictionnaire, valeur, colonne="statut"):
    liste = dictionnaire["clients"][colonne]["valeurs"]["liste"]
    return next(e for e in liste if e["valeur"] == valeur)


def test_repere_des_revues(tmp_path):
    chemin = exporter_revue(_dictionnaire(), "2026-03", tmp_path / "revue.xlsx", nb_revues=2)
    assert len(_importer(chemin, nb_revues=2).regles) == 16
    # Une revue a été appliquée depuis l'export : le classeur annulerait ses décisions.
    with pytest.raises(ValueError, match="exporté après 2 revues, le dictionnaire en compte 3"):
        _importer(chemin, nb_revues=3)
    classeur = load_workbook(chemin)
    notice = classeur["mode_emploi"]
    ligne = next(c.row for c in notice["A"] if c.value == REPERE_REVUES)
    assert notice.cell(row=ligne, column=2).value == 2
    notice.cell(row=ligne, column=2).value = "deux"
    classeur.save(chemin)
    with pytest.raises(ValueError, match="repère des revues absent ou modifié"):
        _importer(chemin, nb_revues=2)
    with pytest.raises(ValueError, match="nb_revues"):
        exporter_revue(_dictionnaire(), "2026-03", tmp_path / "autre.xlsx", nb_revues=-1)


def test_application_sans_decision_ne_change_rien(tmp_path):
    chemin = _exporter(tmp_path)
    dictionnaire = _dictionnaire()
    resultat, rapport, entree = appliquer_revue(dictionnaire, _importer(chemin), chemin, JOUR)
    assert resultat == dictionnaire
    assert entree is None
    assert rapport.empty and tuple(rapport.columns) == COLONNES_RAPPORT_REVUE


def test_application_des_decisions(tmp_path):
    dictionnaire = _dictionnaire()
    temoin = deepcopy(dictionnaire)
    chemin = _exporter(tmp_path)
    _decisions(chemin)
    revue = _importer(chemin)
    resultat, rapport, entree = appliquer_revue(dictionnaire, revue, chemin, JOUR)
    assert dictionnaire == temoin  # l'argument n'est pas modifié

    statut = resultat["clients"]["statut"]
    # valide : la règle est la proposition du moment.
    assert (statut["obligatoire"]["statut"], statut["obligatoire"]["regle"]) == ("valide", "obligatoire")
    assert statut["obligatoire"]["revu_le"] == JOUR
    # documenté : la règle est la règle retenue, la proposition reste.
    nature = statut["nature"]
    assert (nature["statut"], nature["regle"], nature["proposition"]) == (
        "documenté", "entier (texte)", "texte",
    )
    assert nature["commentaire"] == "code numérique"
    # invalide : plus aucune règle, la liste et ses valeurs sont conservées.
    code = resultat["clients"]["code"]["valeurs"]
    assert (code["statut"], code["regle"], code["proposition"]) == (
        "invalide", "aucune", "liste fermée",
    )
    assert [e["valeur"] for e in code["liste"]] == PIEGES
    duree = resultat["interventions"]["duree"]["obligatoire"]
    assert (duree["regle"], duree["proposition"]) == ("facultatif", "à décider")
    # Une ligne laissée à « observé » n'est pas touchée.
    assert statut["valeurs"] == {**temoin["clients"]["statut"]["valeurs"], "liste": statut["valeurs"]["liste"]}

    assert _element(resultat, "Actif")["statut"] == "valide"
    assert _element(resultat, "Inactif") == _element(temoin, "Inactif")
    actif = _element(resultat, "actif")
    assert (actif["statut"], actif["remplacement"], actif["commentaire"], actif["revu_le"]) == (
        "invalide", "Actif", "casse", JOUR,
    )
    assert _element(resultat, "En sommeil") == {
        "valeur": "En sommeil",
        "origine": "ajoutée",
        "effectif": 0,
        "statut": "documenté",
        "commentaire": "chapitre 4",
        "revu_le": JOUR,
        "remplacement": None,
    }
    # Une valeur peut être ajoutée dans une colonne sans liste proposée.
    assert _element(resultat, "Boulangerie", "type_commerce")["origine"] == "ajoutée"
    assert resultat["clients"]["type_commerce"]["valeurs"]["regle"] == "aucune"

    assert entree == {
        "date": JOUR,
        "classeur": "revue.xlsx",
        "regles_changees": 4,
        "valeurs_changees": 2,
        "valeurs_ajoutees": 2,
        "formes_changees": 0,
        "formes_ajoutees": 0,
        "commentaires_changes": 0,
    }
    assert len(rapport) == 8
    ligne = rapport[rapport["valeur"] == "actif"].iloc[0]
    assert (ligne["evenement"], ligne["avant"], ligne["apres"]) == (
        "décision changée", "observé", "invalide, remplacée par 'Actif'",
    )
    ligne = rapport[(rapport["colonne"] == "statut") & (rapport["regle"] == "nature")].iloc[0]
    assert (ligne["avant"], ligne["apres"]) == (
        "observé, règle texte", "documenté, règle entier (texte)",
    )
    assert set(rapport.loc[rapport["valeur"].isin(["En sommeil", "Boulangerie"]), "evenement"]) == {
        "valeur ajoutée"
    }


def test_dictionnaire_revu_ecrit_avec_son_journal(tmp_path):
    fichier = tmp_path / "dictionnaire.yaml"
    ecrire_dictionnaire(_dictionnaire(), "2026-03", fichier)
    resultat, _, entree = _apres_premiere_revue(tmp_path)
    ecrire_dictionnaire(resultat, "2026-03", fichier, revue=entree)
    assert charger_dictionnaire(fichier) == resultat
    assert charger_meta(fichier)["revues"] == [entree]


def test_export_affiche_les_decisions_enregistrees(tmp_path):
    resultat, _, _ = _apres_premiere_revue(tmp_path)
    regles = {(l["table"], l["colonne"], l["regle"]): l for l in lignes_regles(resultat)}
    nature = regles[("clients", "statut", "nature")]
    assert (nature["statut"], nature["regle_en_vigueur"], nature["regle_retenue"], nature["commentaire"]) == (
        "documenté", "entier (texte)", "entier (texte)", "code numérique",
    )
    valide = regles[("clients", "statut", "obligatoire")]
    assert (valide["regle_en_vigueur"], valide["regle_retenue"], valide["commentaire"]) == (
        "obligatoire", None, None,
    )
    assert regles[("clients", "code", "valeurs")]["regle_en_vigueur"] == "aucune"
    valeurs = {(l["colonne"], l["valeur"]): l for l in lignes_valeurs(resultat)}
    assert valeurs[("statut", "actif")]["remplacement"] == "Actif"
    assert valeurs[("statut", "actif")]["commentaire"] == "casse"
    assert valeurs[("statut", "En sommeil")]["origine"] == "ajoutée"
    assert ("type_commerce", "Boulangerie") in valeurs

    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    classeur = load_workbook(chemin)
    lues = {(l["colonne"], l["valeur"]): l for l in _lire(classeur["valeurs"])}
    assert lues[("statut", "actif")]["remplacement"] == "Actif"
    assert lues[("statut", "En sommeil")]["statut"] == "documenté"
    feuille = classeur["regles"]
    assert feuille.cell(row=2, column=COLONNES_REGLES.index("regle_en_vigueur") + 1).protection.locked


def test_liste_ecartee_qui_n_est_plus_proposee_reste_dans_le_classeur():
    dictionnaire = _dictionnaire()
    code = dictionnaire["clients"]["code"]["valeurs"]
    code.update(statut="invalide", regle="aucune", proposition="aucune", revu_le=JOUR,
                commentaire="identifiant libre")
    ligne = next(l for l in lignes_regles(dictionnaire) if (l["colonne"], l["regle"]) == ("code", "valeurs"))
    assert (ligne["proposition"], ligne["statut"]) == ("aucune", "invalide")
    # Sans décision, une colonne sans liste proposée garde sa ligne « valeurs ».
    cles = [(l["colonne"], l["regle"]) for l in lignes_regles(dictionnaire)]
    assert ("type_commerce", "valeurs") in cles


def test_classeur_reexporte_sans_modification_ne_change_rien(tmp_path):
    resultat, _, _ = _apres_premiere_revue(tmp_path)
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    revue = _importer(chemin, resultat, nb_revues=1)
    encore, rapport, entree = appliquer_revue(resultat, revue, chemin, "2026-11-02 09:00:00")
    assert encore == resultat
    assert entree is None and rapport.empty


def test_deuxieme_revue_modifie_annule_et_retire(tmp_path):
    resultat, _, _ = _apres_premiere_revue(tmp_path)
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    # Retour à « observé » : la décision est annulée.
    _modifier(chemin, "regles", dict(STATUT, regle="nature"), statut="observé",
              regle_retenue=None, commentaire=None)
    # Règle retenue changée sans changer de statut.
    _modifier(chemin, "regles", dict(table="interventions", colonne="duree", regle="obligatoire"),
              regle_retenue="obligatoire")
    # Commentaire seul modifié.
    _modifier(chemin, "regles", dict(STATUT, regle="obligatoire"), commentaire="confirmé")
    _modifier(chemin, "valeurs", dict(STATUT, valeur="Actif"), commentaire="valeur de référence")
    # Remplacement retiré, valeur ajoutée retirée.
    _modifier(chemin, "valeurs", dict(STATUT, valeur="actif"), remplacement=None)
    _modifier(chemin, "valeurs", dict(STATUT, valeur="En sommeil"), statut="observé")
    second = "2026-11-02 09:00:00"
    revue = _importer(chemin, resultat, nb_revues=1)
    final, rapport, entree = appliquer_revue(resultat, revue, "C:\\Users\\x\\revue2.xlsx", second)

    nature = final["clients"]["statut"]["nature"]
    assert (nature["statut"], nature["regle"], nature["revu_le"], nature["commentaire"]) == (
        "observé", "texte", None, "",
    )
    duree = final["interventions"]["duree"]["obligatoire"]
    assert (duree["statut"], duree["regle"], duree["revu_le"]) == ("documenté", "obligatoire", second)
    # Un commentaire seul ne déplace pas la date de la décision.
    obligatoire = final["clients"]["statut"]["obligatoire"]
    assert (obligatoire["commentaire"], obligatoire["revu_le"]) == ("confirmé", JOUR)
    assert (_element(final, "Actif")["commentaire"], _element(final, "Actif")["revu_le"]) == (
        "valeur de référence", JOUR,
    )
    actif = _element(final, "actif")
    assert (actif["statut"], actif["remplacement"], actif["revu_le"]) == ("invalide", None, second)
    valeurs = [e["valeur"] for e in final["clients"]["statut"]["valeurs"]["liste"]]
    assert "En sommeil" not in valeurs

    assert entree == {
        "date": second,
        "classeur": "revue2.xlsx",
        "regles_changees": 2,
        "valeurs_changees": 2,
        "valeurs_ajoutees": 0,
        "formes_changees": 0,
        "formes_ajoutees": 0,
        "commentaires_changes": 2,
    }
    evenements = sorted(rapport["evenement"])
    assert evenements == ["commentaire modifié"] * 2 + ["décision changée"] * 3 + ["valeur retirée"]
    ligne = rapport[rapport["evenement"] == "valeur retirée"].iloc[0]
    assert (ligne["valeur"], ligne["avant"], ligne["apres"]) == ("En sommeil", "documenté", "")


def test_retour_a_observe_selon_l_effectif(tmp_path):
    dictionnaire = _dictionnaire()
    liste = dictionnaire["clients"]["statut"]["valeurs"]["liste"]
    # Valeur ajoutée par le métier, observée depuis dans les données.
    liste.append({**_valeur("En sommeil", "ajoutée", 4), "statut": "documenté", "revu_le": JOUR})
    # Valeur revue qui n'est plus observée.
    liste[2].update(statut="invalide", effectif=0, revu_le=JOUR)
    chemin = exporter_revue(dictionnaire, "2026-03", tmp_path / "revue.xlsx", nb_revues=1)
    _modifier(chemin, "valeurs", dict(STATUT, valeur="En sommeil"), statut="observé")
    _modifier(chemin, "valeurs", dict(STATUT, valeur="actif"), statut="observé")
    revue = _importer(chemin, dictionnaire, nb_revues=1)
    resultat, rapport, entree = appliquer_revue(dictionnaire, revue, chemin, "2026-11-02 09:00:00")
    sommeil = _element(resultat, "En sommeil")
    assert (sommeil["origine"], sommeil["statut"], sommeil["effectif"], sommeil["revu_le"]) == (
        "à arbitrer", "observé", 4, None,
    )
    assert "actif" not in [e["valeur"] for e in resultat["clients"]["statut"]["valeurs"]["liste"]]
    assert entree["valeurs_changees"] == 2


def test_remplacements_refuses(tmp_path):
    chemin = _exporter(tmp_path)
    _modifier(chemin, "valeurs", dict(STATUT, valeur="Actif"), statut="valide", remplacement="Inactif")
    # Cible restée à « observé », visée par deux lignes.
    _modifier(chemin, "valeurs", dict(STATUT, valeur="actif"), statut="invalide", remplacement="Inactif")
    code = dict(table="clients", colonne="code")
    _modifier(chemin, "valeurs", dict(code, valeur="076"), statut="invalide", remplacement="076")
    _modifier(chemin, "valeurs", dict(code, valeur="27"), statut="invalide", remplacement="Actif")
    _modifier(chemin, "valeurs", dict(code, valeur="Non"), statut="invalide", remplacement=27)
    _modifier(chemin, "valeurs", FIN + 1, table="clients", colonne="statut",
              valeur="En sommeil", statut="documenté", remplacement="Actif")
    _modifier(chemin, "valeurs", FIN + 2, table="clients", colonne="statut",
              valeur="Actf", statut="documenté")
    message = _erreur_import(chemin)
    assert "6 problèmes" in message
    for attendu in (
        "remplacement réservé au statut invalide",
        "une valeur ne peut pas être son propre remplacement",
        "remplacement 'Actif' absent des valeurs de clients.code",
        "remplacement 27 lu comme int",
        "remplacement interdit sur une valeur ajoutée",
        "valeur 'Inactif' au statut 'observé' alors qu'elle sert de remplacement (lignes 4)",
    ):
        assert attendu in message


def test_remplacement_vers_une_valeur_ajoutee_dans_le_meme_classeur(tmp_path):
    chemin = _exporter(tmp_path)
    _modifier(chemin, "valeurs", dict(STATUT, valeur="actif"), statut="invalide",
              remplacement="En sommeil")
    _modifier(chemin, "valeurs", FIN + 1, table="clients", colonne="statut",
              valeur="En sommeil", statut="documenté")
    resultat, _, _ = appliquer_revue(_dictionnaire(), _importer(chemin), chemin, JOUR)
    assert _element(resultat, "actif")["remplacement"] == "En sommeil"


def test_invalider_ou_retirer_une_cible_bloque_l_import(tmp_path):
    resultat, _, _ = _apres_premiere_revue(tmp_path)
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    _modifier(chemin, "valeurs", dict(STATUT, valeur="Actif"), statut="invalide")
    with pytest.raises(ValueError) as erreur:
        _importer(chemin, resultat, nb_revues=1)
    message = str(erreur.value)
    assert "1 problèmes" in message
    assert "valeur 'Actif' au statut 'invalide' alors qu'elle sert de remplacement" in message


def test_statuts_d_une_valeur_ajoutee_deja_enregistree(tmp_path):
    resultat, _, _ = _apres_premiere_revue(tmp_path)
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    _modifier(chemin, "valeurs", dict(STATUT, valeur="En sommeil"), statut="invalide")
    with pytest.raises(ValueError, match="impossible sur une valeur ajoutée"):
        _importer(chemin, resultat, nb_revues=1)


def test_regle_validee_conservee_quand_la_proposition_change(tmp_path):
    resultat, _, _ = _apres_premiere_revue(tmp_path)
    # Régénération : la proposition ne permet plus de trancher.
    resultat["clients"]["statut"]["obligatoire"]["proposition"] = "à décider"
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    revue = _importer(chemin, resultat, nb_revues=1)
    encore, _, entree = appliquer_revue(resultat, revue, chemin, "2026-11-02 09:00:00")
    assert entree is None
    assert encore["clients"]["statut"]["obligatoire"]["regle"] == "obligatoire"
    # Une nouvelle validation de cette proposition reste refusée.
    _modifier(chemin, "regles", dict(table="interventions", colonne="duree", regle="nature"),
              statut="valide")
    with pytest.raises(ValueError, match="ne peut pas être validée"):
        _importer(chemin, resultat, nb_revues=1)


def test_rien_a_valider_sans_liste_proposee(tmp_path):
    dictionnaire = _dictionnaire()
    dictionnaire["clients"]["code"]["valeurs"].update(
        statut="invalide", regle="aucune", proposition="aucune", revu_le=JOUR, commentaire="libre"
    )
    chemin = exporter_revue(dictionnaire, "2026-03", tmp_path / "revue.xlsx", nb_revues=1)
    _modifier(chemin, "regles", dict(table="clients", colonne="code", regle="valeurs"), statut="valide")
    with pytest.raises(ValueError, match="il n'y a rien à valider"):
        _importer(chemin, dictionnaire, nb_revues=1)


def test_regle_en_vigueur_modifiee_dans_le_classeur(tmp_path):
    chemin = _exporter(tmp_path)
    _modifier(chemin, "regles", dict(STATUT, regle="obligatoire"), regle_en_vigueur="facultatif")
    assert "regle_en_vigueur 'facultatif' différente du dictionnaire" in _erreur_import(chemin)


# Horodatage des décisions et protection à l'écriture (morceau 4d).


def test_instant_de_l_application(tmp_path):
    chemin = _exporter(tmp_path)
    _modifier(chemin, "regles", dict(STATUT, regle="obligatoire"), statut="valide")
    revue = _importer(chemin)
    # Par défaut, l'instant présent, à la seconde.
    resultat, _, entree = appliquer_revue(_dictionnaire(), revue, chemin)
    assert len(entree["date"]) == 19 and lire_instant(entree["date"]) is not None
    assert resultat["clients"]["statut"]["obligatoire"]["revu_le"] == entree["date"]
    # Une date seule vaut minuit ce jour-là.
    _, _, entree = appliquer_revue(_dictionnaire(), revue, chemin, "2026-11-02")
    assert entree["date"] == "2026-11-02 00:00:00"
    with pytest.raises(ValueError, match="02/11/2026"):
        appliquer_revue(_dictionnaire(), revue, chemin, "02/11/2026")


def test_deux_revues_ecrites_puis_regeneration_refusee(tmp_path):
    fichier = tmp_path / "config" / "dictionnaire.yaml"
    ecrire_dictionnaire(_dictionnaire(), "2026-03", fichier)
    resultat, _, entree = _apres_premiere_revue(tmp_path)
    ecrire_dictionnaire(resultat, "2026-03", fichier, revue=entree)
    # Rien à protéger avant la première revue : aucune copie.
    assert not dossier_sauvegardes(fichier).exists()
    apres_premiere = fichier.read_bytes()

    # Seconde revue : décision annulée, règle retenue changée, valeur ajoutée retirée.
    classeur = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    _modifier(classeur, "regles", dict(STATUT, regle="nature"), statut="observé",
              regle_retenue=None, commentaire=None)
    _modifier(classeur, "regles", dict(table="interventions", colonne="duree", regle="obligatoire"),
              regle_retenue="obligatoire")
    _modifier(classeur, "regles", dict(STATUT, regle="obligatoire"), commentaire="confirmé")
    _modifier(classeur, "valeurs", dict(STATUT, valeur="En sommeil"), statut="observé")
    revue = _importer(classeur, resultat, nb_revues=1)
    # Un classeur appliqué à une heure antérieure à la dernière revue est refusé.
    final, _, ancienne = appliquer_revue(resultat, revue, classeur, "2026-10-09 10:14:59")
    with pytest.raises(ValueError, match="doit être postérieure à la dernière revue"):
        ecrire_dictionnaire(final, "2026-03", fichier, revue=ancienne)
    # Le même jour, une heure plus tard, il passe.
    final, _, seconde = appliquer_revue(resultat, revue, classeur, "2026-10-09 11:15:00")
    ecrire_dictionnaire(final, "2026-03", fichier, revue=seconde)
    assert charger_dictionnaire(fichier) == final
    assert charger_meta(fichier)["revues"] == [entree, seconde]
    copies = list(dossier_sauvegardes(fichier).iterdir())
    assert len(copies) == 1 and copies[0].read_bytes() == apres_premiere

    # Un dictionnaire fraîchement généré ne remplace pas le dictionnaire revu.
    with pytest.raises(ValueError, match="écriture refusée") as erreur:
        ecrire_dictionnaire(_dictionnaire(), "2026-04", fichier)
    message = str(erreur.value)
    assert "clients.statut, valeur 'actif' : statut 'invalide' devenu 'observé'" in message
    assert "clients.type_commerce, valeur 'Boulangerie' : décision (documenté) absente" in message
    assert charger_dictionnaire(fichier) == final


def test_mode_emploi_decrit_les_cas_particuliers(tmp_path):
    classeur = load_workbook(_exporter(tmp_path))
    cellules = [c for c in classeur["mode_emploi"]["A"] if c.value]
    textes = [c.value for c in cellules]
    debut = textes.index("Cas particuliers")
    fin = next(i for i, t in enumerate(textes) if t.startswith("Règles à respecter"))
    assert cellules[debut].font.bold
    cas = textes[debut + 1 : fin]
    assert [t.split(" :")[0] for t in cas] == [
        "Annuler une décision (règle, valeur ou forme)",
        "Retirer une valeur ou une forme ajoutée (origine « ajoutée »)",
        "Écarter toute une liste ou tout un format",
        "Statut invalide sur une règle",
        "Changer une règle déjà validée quand la proposition a changé",
        "Déclarer correcte une valeur ou une forme absente des données",
        "Déclarer une liste que l'outil n'a pas proposée (proposition « aucune »)",
        "Déclarer un format que l'outil n'a pas proposé (proposition « aucune »)",
        "Colonne sensible (motif « colonne sensible », effectif « masqué »)",
        "Remplacer une valeur erronée",
        "Valeur mal formée",
        "Modifier seulement un commentaire",
        "Reprendre une revue",
        "Avant l'import",
    ]
    # Chaque consigne n'est donnée qu'une fois dans la feuille.
    assert sum("ne s'applique qu'une fois" in t for t in textes) == 1
    assert sum("etirer une valeur ou une forme ajoutée" in t for t in textes) == 1


def test_mode_emploi_compte_ce_qui_reste_a_traiter(tmp_path):
    def reste(chemin):
        textes = [c.value for c in load_workbook(chemin)["mode_emploi"]["A"] if c.value]
        debut = textes.index("Ce qui reste à traiter")
        return " ".join(textes[debut + 1 : debut + 7])

    # Avant toute revue : tout est à « observé ».
    notice = reste(_exporter(tmp_path))
    assert "« regles » : 10 lignes à « observé » sur 10" in notice
    assert "2 ont une proposition « à décider »" in notice
    assert "4 ont des écarts à arbitrer" in notice
    # Deux colonnes sans liste proposée, quatre sans format proposé : ces lignes ne
    # sont pas à revoir.
    assert "S'y ajoutent 2 lignes « valeurs » et 4 lignes « format » à « observé »" in notice
    assert "« formes » : 0 lignes à « observé » sur 0, dont 0 d'origine « à arbitrer »" in notice
    assert "« valeurs » : 10 lignes à « observé » sur 10, dont 1 d'origine « à arbitrer »" in notice
    assert "S'y ajoutent 0 valeurs de listes écartées" in notice

    # Après la première revue : quatre règles décidées, dont une liste écartée.
    resultat, _, _ = _apres_premiere_revue(tmp_path / "premiere")
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    notice = reste(chemin)
    assert "« regles » : 6 lignes à « observé » sur 10" in notice
    assert "1 ont une proposition « à décider »" in notice
    assert "2 ont des écarts à arbitrer" in notice
    # Les sept valeurs de la liste écartée ne sont plus à revoir.
    assert "« valeurs » : 1 lignes à « observé » sur 12, dont 0 d'origine « à arbitrer »" in notice
    assert "S'y ajoutent 7 valeurs de listes écartées" in notice


def test_reste_a_traiter_sans_exporter(tmp_path):
    assert reste_a_traiter(_dictionnaire()) == {
        "regles": 10,
        "regles_sans_liste": 2,
        "regles_sans_format": 4,
        "regles_a_traiter": 10,
        "regles_a_decider": 2,
        "regles_avec_ecarts": 4,
        "valeurs": 10,
        "valeurs_a_traiter": 10,
        "valeurs_a_arbitrer": 1,
        "valeurs_listes_ecartees": 0,
        "valeurs_non_declarees": 0,
        "formes": 0,
        "formes_a_traiter": 0,
        "formes_a_arbitrer": 0,
        "formes_formats_ecartes": 0,
        "formes_non_declarees": 0,
        "complet": False,
    }
    resultat, _, _ = _apres_premiere_revue(tmp_path)
    reste = reste_a_traiter(resultat)
    assert (reste["regles_a_traiter"], reste["valeurs_a_traiter"]) == (6, 1)
    assert (reste["valeurs_listes_ecartees"], reste["complet"]) == (7, False)
    # Tout décider : la revue est complète, les valeurs de la liste écartée mises à part.
    for colonnes in resultat.values():
        for regles in colonnes.values():
            for nom, regle in regles.items():
                if regle["statut"] == "observé":
                    regle.update(statut="invalide", regle="aucune", revu_le=JOUR)
    reste = reste_a_traiter(resultat)
    assert (reste["regles_a_traiter"], reste["valeurs_a_traiter"], reste["complet"]) == (0, 0, True)
    assert reste["valeurs_listes_ecartees"] == 8


# Liste fermée déclarée par le métier sur une colonne sans liste proposée (morceau 5a).

COMMERCE = dict(table="clients", colonne="type_commerce")


def _declarer(chemin, valeurs=("Boulangerie", "Boucherie"), **champs):
    """Déclare dans le classeur une liste fermée sur clients.type_commerce."""
    champs = {"statut": "documenté", "regle_retenue": "liste fermée", **champs}
    _modifier(chemin, "regles", dict(COMMERCE, regle="valeurs"), **champs)
    for rang, valeur in enumerate(valeurs, start=1):
        _modifier(chemin, "valeurs", FIN + rang, **COMMERCE, valeur=valeur, statut="documenté")


def _apres_declaration(tmp_path):
    chemin = _exporter(tmp_path)
    _declarer(chemin)
    return appliquer_revue(_dictionnaire(), _importer(chemin), chemin, JOUR)


def test_declarer_une_liste_non_proposee(tmp_path):
    resultat, rapport, entree = _apres_declaration(tmp_path)
    regle = resultat["clients"]["type_commerce"]["valeurs"]
    assert (regle["regle"], regle["proposition"]) == ("liste fermée", "aucune")
    assert (regle["statut"], regle["revu_le"]) == ("documenté", JOUR)
    assert [(v["valeur"], v["origine"], v["statut"], v["effectif"]) for v in regle["liste"]] == [
        ("Boulangerie", "ajoutée", "documenté", 0),
        ("Boucherie", "ajoutée", "documenté", 0),
    ]
    assert (entree["regles_changees"], entree["valeurs_ajoutees"]) == (1, 2)
    ligne = rapport[rapport["regle"].eq("valeurs") & rapport["valeur"].eq("")].iloc[0]
    assert (ligne["avant"], ligne["apres"]) == (
        "observé, règle aucune", "documenté, règle liste fermée",
    )
    # La liste déclarée compte désormais parmi les lignes revues.
    reste = reste_a_traiter(resultat)
    assert (reste["regles"], reste["regles_a_traiter"], reste["regles_sans_liste"]) == (11, 10, 1)

    # Le dictionnaire s'écrit, et le classeur suivant affiche la déclaration.
    fichier = tmp_path / "dictionnaire.yaml"
    ecrire_dictionnaire(_dictionnaire(), "2026-03", fichier)
    ecrire_dictionnaire(resultat, "2026-03", fichier, revue=entree)
    assert charger_dictionnaire(fichier) == resultat
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    lue = next(l for l in _lire(load_workbook(chemin)["regles"])
               if (l["colonne"], l["regle"]) == ("type_commerce", "valeurs"))
    assert (lue["proposition"], lue["regle_en_vigueur"], lue["regle_retenue"]) == (
        "aucune", "liste fermée", "liste fermée",
    )
    _, rapport, entree = appliquer_revue(resultat, _importer(chemin, resultat, 1), chemin, JOUR)
    assert entree is None and rapport.empty


def test_declarations_refusees(tmp_path):
    # Sans valeur, sans règle retenue, ou avec une autre règle retenue.
    chemin = _exporter(tmp_path)
    _declarer(chemin, valeurs=())
    assert "liste déclarée sans aucune valeur au statut valide ou documenté pour " \
        "clients.type_commerce" in _erreur_import(chemin)
    _declarer(chemin, regle_retenue=None)
    assert "regle_retenue obligatoire avec le statut documenté" in _erreur_import(chemin)
    _declarer(chemin, regle_retenue="aucune")
    assert "regle_retenue 'aucune' hors du menu de la règle valeurs" in _erreur_import(chemin)

    # Une liste déjà proposée se valide, elle ne se déclare pas : un seul problème.
    chemin = _exporter(tmp_path / "proposee")
    _modifier(chemin, "regles", dict(STATUT, regle="valeurs"), statut="documenté",
              regle_retenue="liste fermée")
    message = _erreur_import(chemin)
    assert "1 problèmes" in message and "il n'y a rien à déclarer (choisir valide)" in message

    # Rien à écarter sur une colonne sans liste proposée ni valeur.
    chemin = _exporter(tmp_path / "ecartee")
    _modifier(chemin, "regles", dict(COMMERCE, regle="valeurs"), statut="invalide",
              commentaire="texte libre")
    assert "il n'y a rien à écarter" in _erreur_import(chemin)


def test_annuler_ou_vider_une_liste_declaree(tmp_path):
    declare, _, _ = _apres_declaration(tmp_path)
    # Retirer toutes les valeurs d'une liste déclarée bloque l'import.
    chemin = exporter_revue(declare, "2026-03", tmp_path / "vide.xlsx", nb_revues=1)
    for valeur in ("Boulangerie", "Boucherie"):
        _modifier(chemin, "valeurs", dict(COMMERCE, valeur=valeur), statut="observé")
    with pytest.raises(ValueError, match="liste déclarée sans aucune valeur"):
        _importer(chemin, declare, nb_revues=1)

    # Retour à « observé » : plus de liste en vigueur, les valeurs ajoutées restent.
    chemin = exporter_revue(declare, "2026-03", tmp_path / "annule.xlsx", nb_revues=1)
    _modifier(chemin, "regles", dict(COMMERCE, regle="valeurs"), statut="observé",
              regle_retenue=None)
    resultat, _, entree = appliquer_revue(declare, _importer(chemin, declare, 1), chemin, JOUR)
    regle = resultat["clients"]["type_commerce"]["valeurs"]
    assert (regle["regle"], regle["statut"], regle["revu_le"]) == ("aucune", "observé", None)
    assert [v["statut"] for v in regle["liste"]] == ["documenté", "documenté"]
    assert (entree["regles_changees"], entree["valeurs_changees"]) == (1, 0)

    # Écarter une liste déclarée : elle a des valeurs, la décision est admise.
    chemin = exporter_revue(declare, "2026-03", tmp_path / "ecarte.xlsx", nb_revues=1)
    _modifier(chemin, "regles", dict(COMMERCE, regle="valeurs"), statut="invalide",
              regle_retenue=None, commentaire="finalement libre")
    resultat, _, _ = appliquer_revue(declare, _importer(chemin, declare, 1), chemin, JOUR)
    assert resultat["clients"]["type_commerce"]["valeurs"]["regle"] == "aucune"


# Colonnes candidates et effectif masqué des colonnes sensibles (morceau 5b).


def _avec_candidate():
    """Dictionnaire dont clients.type_commerce porte trois valeurs observées, sans liste proposée."""
    dictionnaire = _dictionnaire()
    dictionnaire["clients"]["type_commerce"]["valeurs"]["liste"] = [
        _valeur("Boulangerie", "observée", 12),
        _valeur("Boucherie", "observée", 6),
        _valeur("boulangerie", "observée", 1),
    ]
    return dictionnaire


def test_valeurs_candidates_hors_du_reste_a_traiter(tmp_path):
    dictionnaire = _avec_candidate()
    reste = reste_a_traiter(dictionnaire)
    assert (reste["valeurs"], reste["valeurs_a_traiter"]) == (13, 10)
    assert (reste["valeurs_non_declarees"], reste["valeurs_listes_ecartees"]) == (3, 0)
    chemin = exporter_revue(dictionnaire, "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    classeur = load_workbook(chemin)
    lues = [l for l in _lire(classeur["valeurs"]) if l["colonne"] == "type_commerce"]
    assert [(l["valeur"], l["effectif"], l["origine"]) for l in lues] == [
        ("Boulangerie", 12, "observée"), ("Boucherie", 6, "observée"), ("boulangerie", 1, "observée"),
    ]
    notice = " ".join(c.value for c in classeur["mode_emploi"]["A"] if c.value)
    assert "et 3 valeurs d'origine « observée », à revoir seulement si vous déclarez" in notice


def test_declarer_une_liste_en_validant_les_valeurs_observees(tmp_path):
    dictionnaire = _avec_candidate()
    chemin = exporter_revue(dictionnaire, "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    # Revoir des valeurs sans déclarer la liste : admis, avec un avertissement.
    _modifier(chemin, "valeurs", dict(COMMERCE, valeur="Boulangerie"), statut="valide")
    assert _importer(chemin, dictionnaire).avertissements == [
        "clients.type_commerce : 1 valeurs revues, mais liste non déclarée "
        "(statuts conservés, sans effet tant que la liste n'est pas déclarée)"
    ]
    # Déclaration : aucune valeur à ressaisir, seule la valeur absente des données s'ajoute.
    _declarer(chemin, valeurs=())
    _modifier(chemin, "valeurs", FIN + 3 + 1, **COMMERCE, valeur="Poissonnerie",
              statut="documenté")
    _modifier(chemin, "valeurs", dict(COMMERCE, valeur="Boucherie"), statut="valide")
    _modifier(chemin, "valeurs", dict(COMMERCE, valeur="boulangerie"), statut="invalide",
              remplacement="Boulangerie")
    revue = _importer(chemin, dictionnaire)
    assert revue.avertissements == []
    resultat, _, entree = appliquer_revue(dictionnaire, revue, chemin, JOUR)
    regle = resultat["clients"]["type_commerce"]["valeurs"]
    assert (regle["regle"], regle["statut"]) == ("liste fermée", "documenté")
    assert [(v["valeur"], v["origine"], v["statut"], v["effectif"], v["remplacement"])
            for v in regle["liste"]] == [
        ("Boulangerie", "observée", "valide", 12, None),
        ("Boucherie", "observée", "valide", 6, None),
        ("boulangerie", "observée", "invalide", 1, "Boulangerie"),
        ("Poissonnerie", "ajoutée", "documenté", 0, None),
    ]
    assert (entree["valeurs_changees"], entree["valeurs_ajoutees"]) == (3, 1)
    reste = reste_a_traiter(resultat)
    assert (reste["valeurs_non_declarees"], reste["valeurs_a_traiter"]) == (0, 10)


def test_valeur_ajoutee_dans_une_colonne_sensible(tmp_path):
    dictionnaire = _dictionnaire()
    dictionnaire["clients"]["type_commerce"]["valeurs"]["motif"] = "colonne sensible"
    chemin = exporter_revue(dictionnaire, "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    _declarer(chemin, valeurs=("Retraite", "Démission"))
    resultat, _, entree = appliquer_revue(dictionnaire, _importer(chemin, dictionnaire), chemin, JOUR)
    liste = resultat["clients"]["type_commerce"]["valeurs"]["liste"]
    # L'effectif est masqué : le dictionnaire ne dit pas si la valeur est observée.
    assert [(v["valeur"], v["effectif"]) for v in liste] == [("Retraite", None), ("Démission", None)]
    fichier = tmp_path / "dictionnaire.yaml"
    ecrire_dictionnaire(dictionnaire, "2026-03", fichier)
    ecrire_dictionnaire(resultat, "2026-03", fichier, revue=entree)
    assert charger_dictionnaire(fichier) == resultat

    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    lues = [l for l in _lire(load_workbook(chemin)["valeurs"]) if l["colonne"] == "type_commerce"]
    assert [l["effectif"] for l in lues] == ["masqué", "masqué"]
    # Remise à « observé » : la valeur est retirée, comme une valeur d'effectif nul.
    _modifier(chemin, "valeurs", dict(COMMERCE, valeur="Démission"), statut="observé")
    resultat, rapport, _ = appliquer_revue(resultat, _importer(chemin, resultat, 1), chemin, JOUR)
    assert [v["valeur"] for v in resultat["clients"]["type_commerce"]["valeurs"]["liste"]] == ["Retraite"]
    assert list(rapport["evenement"]) == ["valeur retirée"]


# Menus des lignes d'ajout et feuille technique des listes (morceau 5c).


def test_menus_table_et_colonne_des_lignes_d_ajout(tmp_path):
    chemin = _exporter(tmp_path)
    classeur = load_workbook(chemin)
    listes, valeurs = classeur["listes"], classeur["valeurs"]
    # Feuille technique masquée et verrouillée : une colonne par table.
    assert listes.sheet_state == "hidden" and listes.protection.sheet
    assert [c.value for c in listes[1]] == ["clients", "interventions"]
    assert [c.value for c in listes["A"]] == ["clients", "statut", "code", "type_commerce"]
    assert [c.value for c in listes["B"]][:2] == ["interventions", "duree"]
    # Les menus ne couvrent que les lignes d'ajout, en bas de la feuille.
    menus = {str(m.sqref): m.formula1 for m in valeurs.data_validations.dataValidation}
    premiere, derniere = FIN + 1, FIN + LIGNES_AJOUT
    assert menus[f"A{premiere}:A{derniere}"] == "'listes'!$A$1:$B$1"
    colonnes = menus[f"B{premiere}:B{derniere}"]
    # Le menu des colonnes suit la table choisie sur la même ligne.
    assert f"MATCH($A{premiere},'listes'!$1:$1,0)" in colonnes
    assert colonnes.startswith("OFFSET('listes'!$A$1,1,")
    # La feuille technique ne gêne ni l'import ni l'application.
    _modifier(chemin, "valeurs", premiere, table="clients", colonne="type_commerce",
              valeur="Boulangerie", statut="documenté")
    assert list(_importer(chemin).valeurs["valeur"])[-1] == "Boulangerie"


# La règle de format dans le classeur : ligne « format » et feuille des formes (morceau 6b).

REF = dict(table="clients", colonne="ref")
FIN_FORMES = 1 + 3  # dernière ligne exportée de la feuille des formes de `_avec_formats`


def _forme(forme, origine, effectif):
    return {
        "forme": forme,
        "origine": origine,
        "effectif": effectif,
        "statut": "observé",
        "commentaire": "",
        "revu_le": None,
    }


def _avec_formats():
    """Dictionnaire avec un format proposé (clients.ref) et une colonne candidate à un
    format (clients.type_commerce)."""
    dictionnaire = _dictionnaire()
    ref = _colonne()
    ref["format"] = {
        **_regle("formes fermées", {"formes": 1}),
        "liste": [_forme("AA-9999", "liste proposée", 40), _forme("AA9999", "à arbitrer", 2)],
    }
    dictionnaire["clients"]["ref"] = ref
    dictionnaire["clients"]["type_commerce"]["format"]["liste"] = [
        _forme("Aaaa", "observée", 12)
    ]
    return dictionnaire


def _exporter_formats(tmp_path, nom="revue.xlsx"):
    return exporter_revue(_avec_formats(), "2026-03", tmp_path / nom, nb_revues=0)


def _formes(dictionnaire, colonne="ref"):
    return dictionnaire["clients"][colonne]["format"]["liste"]


def test_ligne_format_et_lignes_des_formes():
    dictionnaire = _avec_formats()
    regles = {(l["colonne"], l["regle"]): l for l in lignes_regles(dictionnaire)}
    ligne = regles[("ref", "format")]
    assert (ligne["proposition"], ligne["regle_en_vigueur"], ligne["nb_a_arbitrer"]) == (
        "1 forme", "formes fermées", 1,
    )
    # Colonne candidate : des formes observées, mais aucun format proposé.
    assert regles[("type_commerce", "format")]["proposition"] == "aucune"
    _formes(dictionnaire).insert(1, _forme("AA-99999", "liste proposée", 30))
    regles = {(l["colonne"], l["regle"]): l for l in lignes_regles(dictionnaire)}
    assert regles[("ref", "format")]["proposition"] == "2 formes"

    lignes = lignes_formes(_avec_formats())
    assert [(l["colonne"], l["forme"], l["effectif"], l["origine"]) for l in lignes] == [
        ("type_commerce", "Aaaa", 12, "observée"),
        ("ref", "AA-9999", 40, "liste proposée"),
        ("ref", "AA9999", 2, "à arbitrer"),
    ]
    # Une forme n'a pas de remplacement.
    assert all(tuple(l) == COLONNES_FORMES for l in lignes)
    # Ce qui ne se voit pas à l'écran est signalé, comme pour une valeur.
    dictionnaire = _avec_formats()
    _formes(dictionnaire).append(_forme("AA-9999 ", "à arbitrer", 1))
    assert lignes_formes(dictionnaire)[-1]["remarque"] == "espaces en bord"
    # Colonne sensible : l'effectif d'une forme déclarée est masqué.
    dictionnaire["clients"]["statut"]["format"]["motif"] = "colonne sensible"
    dictionnaire["clients"]["statut"]["format"]["liste"] = [
        {**_forme("Aaaaa", "ajoutée", None), "statut": "documenté", "revu_le": JOUR}
    ]
    assert lignes_formes(dictionnaire)[0]["effectif"] == "masqué"


def test_feuille_des_formes_exportee(tmp_path):
    dictionnaire = _avec_formats()
    # Une forme faite de chiffres reste un texte dans le classeur.
    _formes(dictionnaire).append(_forme("99999", "à arbitrer", 1))
    chemin = exporter_revue(dictionnaire, "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    classeur = load_workbook(chemin)
    formes = classeur["formes"]
    assert tuple(c.value for c in formes[1]) == COLONNES_FORMES
    assert [l["forme"] for l in _lire(formes)] == ["Aaaa", "AA-9999", "AA9999", "99999"]
    assert formes.protection.sheet
    libres = {c.value for c, d in zip(formes[1], formes[2]) if not d.protection.locked}
    assert libres == {"statut", "commentaire"}
    ajout = formes[1 + 4 + 2]
    libres = {c.value for c, d in zip(formes[1], ajout) if not d.protection.locked}
    assert libres == {"table", "colonne", "forme", "statut", "commentaire"}
    assert formes.max_row == 1 + 4 + LIGNES_AJOUT
    assert formes.freeze_panes == "C2"
    # Menu des statuts, puis menus table et colonne des lignes d'ajout.
    menus = {str(m.sqref): m.formula1 for m in formes.data_validations.dataValidation}
    assert len(menus) == 3 and all(
        m.showErrorMessage for m in formes.data_validations.dataValidation
    )
    assert menus[f"A6:A{5 + LIGNES_AJOUT}"] == "'listes'!$A$1:$B$1"
    assert "MATCH($A6,'listes'!$1:$1,0)" in menus[f"B6:B{5 + LIGNES_AJOUT}"]
    # Sur une ligne « format », regle_retenue sert à déclarer des formes fermées.
    regles = classeur["regles"]
    ligne = next(l for l in regles.iter_rows(min_row=2) if l[2].value == "format")
    assert not ligne[COLONNES_REGLES.index("regle_retenue")].protection.locked
    notice = [c for c in classeur["mode_emploi"]["A"] if c.value]
    rubrique = next(c for c in notice if c.value.startswith("Feuille « formes » : les formes"))
    assert rubrique.font.bold
    textes = " ".join(c.value for c in notice)
    assert "9 pour un chiffre, A pour une majuscule, a pour une minuscule" in textes
    assert "4 lignes dans la feuille « formes »" in textes


def _revoir_le_format(chemin):
    """Valide le format de clients.ref, arbitre ses formes et en ajoute une."""
    _modifier(chemin, "regles", dict(REF, regle="format"), statut="valide")
    _modifier(chemin, "formes", dict(REF, forme="AA-9999"), statut="valide")
    _modifier(chemin, "formes", dict(REF, forme="AA9999"), statut="invalide",
              commentaire="tiret oublié")
    _modifier(chemin, "formes", FIN_FORMES + 1, **REF, forme="AA-99999", statut="documenté",
              commentaire="nouvelle numérotation")


def test_revue_des_formes(tmp_path):
    dictionnaire = _avec_formats()
    temoin = deepcopy(dictionnaire)
    chemin = _exporter_formats(tmp_path)
    # Sans modification : trois formes lues, rien à appliquer.
    revue = _importer(chemin, dictionnaire)
    assert list(revue.formes["forme"]) == ["Aaaa", "AA-9999", "AA9999"]
    assert appliquer_revue(dictionnaire, revue, chemin, JOUR)[2] is None

    _revoir_le_format(chemin)
    revue = _importer(chemin, dictionnaire)
    assert revue.avertissements == []
    assert list(revue.formes["ligne_excel"]) == [2, 3, 4, FIN_FORMES + 1]
    resultat, rapport, entree = appliquer_revue(dictionnaire, revue, chemin, JOUR)
    assert dictionnaire == temoin  # l'argument n'est pas modifié

    regle = resultat["clients"]["ref"]["format"]
    assert (regle["statut"], regle["regle"], regle["revu_le"]) == ("valide", "formes fermées", JOUR)
    assert _formes(resultat) == [
        {**_forme("AA-9999", "liste proposée", 40), "statut": "valide", "revu_le": JOUR},
        {**_forme("AA9999", "à arbitrer", 2), "statut": "invalide", "revu_le": JOUR,
         "commentaire": "tiret oublié"},
        # Forme ajoutée : effectif 0, pas de champ de remplacement.
        {**_forme("AA-99999", "ajoutée", 0), "statut": "documenté", "revu_le": JOUR,
         "commentaire": "nouvelle numérotation"},
    ]
    # Les valeurs et les formes non revues ne sont pas touchées.
    assert _formes(resultat, "type_commerce") == _formes(temoin, "type_commerce")
    assert resultat["clients"]["statut"] == temoin["clients"]["statut"]
    assert entree == {
        "date": JOUR,
        "classeur": "revue.xlsx",
        "regles_changees": 1,
        "valeurs_changees": 0,
        "valeurs_ajoutees": 0,
        "formes_changees": 2,
        "formes_ajoutees": 1,
        "commentaires_changes": 0,
    }
    # Dans le rapport, la colonne « valeur » porte la forme.
    assert tuple(rapport.columns) == COLONNES_RAPPORT_REVUE
    assert set(rapport["regle"]) == {"format"}
    evenements = dict(zip(rapport["valeur"], rapport["evenement"]))
    assert evenements == {
        "": "décision changée",
        "AA-9999": "décision changée",
        "AA9999": "décision changée",
        "AA-99999": "forme ajoutée",
    }
    ligne = rapport[rapport["valeur"] == "AA9999"].iloc[0]
    assert (ligne["avant"], ligne["apres"]) == ("observé", "invalide")

    # Le dictionnaire s'écrit avec son journal, et le classeur suivant affiche la revue.
    fichier = tmp_path / "dictionnaire.yaml"
    ecrire_dictionnaire(dictionnaire, "2026-03", fichier)
    ecrire_dictionnaire(resultat, "2026-03", fichier, revue=entree)
    assert charger_dictionnaire(fichier) == resultat
    assert charger_meta(fichier)["revues"] == [entree]
    suivant = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    lues = {l["forme"]: l for l in _lire(load_workbook(suivant)["formes"])}
    assert (lues["AA9999"]["statut"], lues["AA9999"]["commentaire"]) == ("invalide", "tiret oublié")
    assert (lues["AA-99999"]["origine"], lues["AA-99999"]["effectif"]) == ("ajoutée", 0)
    encore, rapport, entree = appliquer_revue(
        resultat, _importer(suivant, resultat, 1), suivant, "2026-11-02 09:00:00"
    )
    assert encore == resultat and entree is None and rapport.empty

    # Seconde revue : forme ajoutée retirée, décision annulée, commentaire seul modifié.
    _modifier(suivant, "formes", dict(REF, forme="AA-99999"), statut="observé")
    _modifier(suivant, "formes", dict(REF, forme="AA9999"), statut="observé", commentaire=None)
    _modifier(suivant, "formes", dict(REF, forme="AA-9999"), commentaire="forme de référence")
    second = "2026-11-02 09:00:00"
    final, rapport, entree = appliquer_revue(
        resultat, _importer(suivant, resultat, 1), suivant, second
    )
    assert [f["forme"] for f in _formes(final)] == ["AA-9999", "AA9999"]
    # Un commentaire seul ne déplace pas la date de la décision.
    assert (_formes(final)[0]["commentaire"], _formes(final)[0]["revu_le"]) == (
        "forme de référence", JOUR,
    )
    # Une forme observée dans les données redevient une forme à revoir.
    assert _formes(final)[1] == _forme("AA9999", "à arbitrer", 2)
    assert (entree["formes_changees"], entree["formes_ajoutees"]) == (2, 0)
    assert (entree["commentaires_changes"], entree["valeurs_changees"]) == (1, 0)
    assert sorted(rapport["evenement"]) == [
        "commentaire modifié", "décision changée", "forme retirée",
    ]
    ecrire_dictionnaire(final, "2026-03", fichier, revue=entree)
    assert charger_dictionnaire(fichier) == final


def test_formes_refusees(tmp_path):
    chemin = _exporter_formats(tmp_path)
    fin = FIN_FORMES
    # Une valeur saisie à la place de sa forme.
    _modifier(chemin, "formes", fin + 1, **REF, forme="CL-00017", statut="documenté")
    _modifier(chemin, "formes", fin + 2, **REF, forme="AA-999", statut="valide")
    _modifier(chemin, "formes", fin + 3, table="clients", colonne="inconnue",
              forme="AAA", statut="documenté")
    _modifier(chemin, "formes", fin + 4, **REF, forme="AA9999", statut="documenté")
    _modifier(chemin, "formes", fin + 5, **REF, statut="documenté")
    _modifier(chemin, "formes", fin + 6, **REF, forme=999, statut="documenté")
    # Une valeur dont la forme figure déjà dans la feuille : la ligne est à vider.
    _modifier(chemin, "formes", fin + 7, **REF, forme="CL-0001", statut="documenté")
    _modifier(chemin, "formes", dict(REF, forme="AA-9999"), statut="documenté")
    with pytest.raises(ValueError) as erreur:
        _importer(chemin, _avec_formats())
    message = str(erreur.value)
    assert "8 problèmes" in message
    for attendu in (
        "formes, ligne 11 : 'CL-0001' n'est pas une forme (9 pour un chiffre, A pour une "
        "majuscule, a pour une minuscule, les autres caractères tels quels), la forme "
        "correspondante est 'AA-9999', déjà présente pour cette colonne (vider la ligne)",
        "formes, ligne 5 : 'CL-00017' n'est pas une forme (9 pour un chiffre, A pour une "
        "majuscule, a pour une minuscule, les autres caractères tels quels), la forme "
        "correspondante est 'AA-99999'",
        "une forme ajoutée doit porter le statut documenté",
        "colonne clients.inconnue absente du dictionnaire",
        "forme 'AA9999' déjà présente pour clients.ref",
        "forme vide",
        "forme 999 lue comme int",
        "statut 'documenté' impossible sur une forme observée",
    ):
        assert attendu in message

    # Ligne exportée modifiée ou supprimée : le classeur ne correspond plus au dictionnaire.
    chemin = _exporter_formats(tmp_path, "modifie.xlsx")
    _modifier(chemin, "formes", dict(REF, forme="AA9999"), forme="AA_9999")
    with pytest.raises(ValueError) as erreur:
        _importer(chemin, _avec_formats())
    message = str(erreur.value)
    assert "forme 'AA_9999' de clients.ref inconnue du dictionnaire" in message
    assert "formes : forme ('clients', 'ref', 'AA9999') absente du classeur" in message

    # Une forme ajoutée déjà enregistrée ne se garde ou ne se retire que telle quelle.
    chemin = _exporter_formats(tmp_path, "premiere.xlsx")
    _revoir_le_format(chemin)
    resultat, _, _ = appliquer_revue(
        _avec_formats(), _importer(chemin, _avec_formats()), chemin, JOUR
    )
    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    _modifier(chemin, "formes", dict(REF, forme="AA-99999"), statut="invalide")
    with pytest.raises(ValueError, match="impossible sur une forme ajoutée"):
        _importer(chemin, resultat, nb_revues=1)


COMMERCE_FORMAT = dict(table="clients", colonne="type_commerce", regle="format")


def test_declarer_un_format_non_propose(tmp_path):
    dictionnaire = _avec_formats()
    chemin = _exporter_formats(tmp_path)
    # Revoir une forme observée sans déclarer le format : admis, avec un avertissement.
    _modifier(chemin, "formes", dict(COMMERCE, forme="Aaaa"), statut="valide")
    assert _importer(chemin, dictionnaire).avertissements == [
        "clients.type_commerce : 1 formes revues, mais format non déclaré "
        "(statuts conservés, sans effet tant que le format n'est pas déclaré)"
    ]
    # Déclaration : la forme observée est validée sans ressaisie, une autre est ajoutée.
    _modifier(chemin, "regles", COMMERCE_FORMAT, statut="documenté",
              regle_retenue="formes fermées")
    _modifier(chemin, "formes", FIN_FORMES + 1, **COMMERCE, forme="Aaaa-aaaa",
              statut="documenté")
    revue = _importer(chemin, dictionnaire)
    assert revue.avertissements == []
    resultat, rapport, entree = appliquer_revue(dictionnaire, revue, chemin, JOUR)
    regle = resultat["clients"]["type_commerce"]["format"]
    assert (regle["regle"], regle["proposition"], regle["statut"]) == (
        "formes fermées", "aucune", "documenté",
    )
    assert [(f["forme"], f["origine"], f["statut"], f["effectif"]) for f in regle["liste"]] == [
        ("Aaaa", "observée", "valide", 12),
        ("Aaaa-aaaa", "ajoutée", "documenté", 0),
    ]
    assert (entree["regles_changees"], entree["formes_changees"], entree["formes_ajoutees"]) == (
        1, 1, 1,
    )
    ligne = rapport[rapport["regle"].eq("format") & rapport["valeur"].eq("")].iloc[0]
    assert (ligne["avant"], ligne["apres"]) == (
        "observé, règle aucune", "documenté, règle formes fermées",
    )
    # Le classeur suivant affiche la déclaration et ne change rien.
    suivant = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    lue = next(l for l in _lire(load_workbook(suivant)["regles"])
               if (l["colonne"], l["regle"]) == ("type_commerce", "format"))
    assert (lue["proposition"], lue["regle_en_vigueur"], lue["regle_retenue"]) == (
        "aucune", "formes fermées", "formes fermées",
    )
    assert appliquer_revue(resultat, _importer(suivant, resultat, 1), suivant, JOUR)[2] is None

    # Retirer toutes les formes d'un format déclaré bloque l'import.
    _modifier(suivant, "formes", dict(COMMERCE, forme="Aaaa"), statut="observé")
    _modifier(suivant, "formes", dict(COMMERCE, forme="Aaaa-aaaa"), statut="observé")
    with pytest.raises(ValueError, match="format déclaré sans aucune forme"):
        _importer(suivant, resultat, nb_revues=1)

    # Écarter le format déclaré : ses formes restent, sans effet, et c'est signalé.
    ecarte = exporter_revue(resultat, "2026-03", tmp_path / "ecarte.xlsx", nb_revues=1)
    _modifier(ecarte, "regles", COMMERCE_FORMAT, statut="invalide", regle_retenue=None,
              commentaire="libellé libre")
    revue = _importer(ecarte, resultat, nb_revues=1)
    assert revue.avertissements == [
        "clients.type_commerce : format écarté, mais 1 formes revues "
        "(statuts conservés, sans effet tant que le format est écarté)"
    ]
    final, _, _ = appliquer_revue(resultat, revue, ecarte, "2026-11-02 09:00:00")
    regle = final["clients"]["type_commerce"]["format"]
    assert (regle["regle"], regle["statut"], len(regle["liste"])) == ("aucune", "invalide", 2)


def test_declarations_de_format_refusees(tmp_path):
    def erreur(chemin):
        with pytest.raises(ValueError) as levee:
            _importer(chemin, _avec_formats())
        return str(levee.value)

    # Format déclaré sans aucune forme valide ou documentée, puis avec une autre règle.
    chemin = _exporter_formats(tmp_path)
    _modifier(chemin, "regles", COMMERCE_FORMAT, statut="documenté",
              regle_retenue="formes fermées")
    assert (
        "format déclaré sans aucune forme au statut valide ou documenté pour "
        "clients.type_commerce (ajouter ses formes en bas de la feuille « formes »)"
    ) in erreur(chemin)
    _modifier(chemin, "regles", COMMERCE_FORMAT, regle_retenue="liste fermée")
    assert "regle_retenue 'liste fermée' hors du menu de la règle format" in erreur(chemin)

    # Un format déjà proposé se valide, il ne se déclare pas : un seul problème.
    chemin = _exporter_formats(tmp_path, "propose.xlsx")
    _modifier(chemin, "regles", dict(REF, regle="format"), statut="documenté",
              regle_retenue="formes fermées")
    message = erreur(chemin)
    assert "1 problèmes" in message
    assert "un format est déjà proposé, il n'y a rien à déclarer (choisir valide)" in message

    # Rien à valider sans format proposé, rien à écarter sans aucune forme.
    chemin = _exporter_formats(tmp_path, "rien.xlsx")
    _modifier(chemin, "regles", COMMERCE_FORMAT, statut="valide")
    _modifier(chemin, "regles", dict(STATUT, regle="format"), statut="invalide",
              commentaire="libellé")
    message = erreur(chemin)
    assert "2 problèmes" in message
    assert "aucun format n'est proposé, il n'y a rien à valider" in message
    assert "aucun format n'est proposé et la colonne n'a aucune forme, il n'y a rien à écarter" in message

    # Écarter un format proposé reste admis, avec un commentaire.
    chemin = _exporter_formats(tmp_path, "ecarte.xlsx")
    _modifier(chemin, "regles", dict(REF, regle="format"), statut="invalide")
    assert "commentaire obligatoire avec le statut invalide" in erreur(chemin)
    _modifier(chemin, "regles", dict(REF, regle="format"), commentaire="texte libre régulier")
    resultat, _, entree = appliquer_revue(
        _avec_formats(), _importer(chemin, _avec_formats()), chemin, JOUR
    )
    regle = resultat["clients"]["ref"]["format"]
    assert (regle["regle"], regle["proposition"]) == ("aucune", "formes fermées")
    assert [f["forme"] for f in regle["liste"]] == ["AA-9999", "AA9999"]
    assert entree["regles_changees"] == 1


def test_reste_a_traiter_des_formes(tmp_path):
    dictionnaire = _avec_formats()
    reste = reste_a_traiter(dictionnaire)
    # Cinq colonnes : trois sans liste proposée, quatre sans format proposé.
    assert (reste["regles"], reste["regles_sans_liste"], reste["regles_sans_format"]) == (13, 3, 4)
    assert (reste["regles_a_traiter"], reste["regles_avec_ecarts"]) == (13, 5)
    # La forme de la colonne candidate n'est à revoir que si le format est déclaré.
    assert (reste["formes"], reste["formes_a_traiter"], reste["formes_a_arbitrer"]) == (3, 2, 1)
    assert (reste["formes_non_declarees"], reste["formes_formats_ecartes"]) == (1, 0)
    notice = " ".join(
        c.value for c in load_workbook(_exporter_formats(tmp_path))["mode_emploi"]["A"] if c.value
    )
    assert "« formes » : 2 lignes à « observé » sur 3, dont 1 d'origine « à arbitrer »" in notice
    assert "S'y ajoutent 0 formes de formats écartés" in notice
    assert "et 1 formes d'origine « observée », à revoir seulement si vous déclarez le format" in notice

    # Format écarté : ses formes non revues ne sont plus à revoir.
    ecarte = _avec_formats()
    ecarte["clients"]["ref"]["format"].update(
        statut="invalide", regle="aucune", revu_le=JOUR, commentaire="texte libre"
    )
    reste = reste_a_traiter(ecarte)
    assert (reste["formes_a_traiter"], reste["formes_formats_ecartes"]) == (0, 2)
    assert (reste["regles_a_traiter"], reste["complet"]) == (12, False)

    # Tout revoir : la revue est complète quand il ne reste ni règle, ni valeur, ni forme.
    complet = _avec_formats()
    for colonnes in complet.values():
        for regles in colonnes.values():
            for regle in regles.values():
                if regle["statut"] == "observé":
                    regle.update(statut="invalide", regle="aucune", revu_le=JOUR)
    complet["clients"]["ref"]["format"].update(statut="valide", regle="formes fermées")
    reste = reste_a_traiter(complet)
    assert (reste["regles_a_traiter"], reste["valeurs_a_traiter"]) == (0, 0)
    assert (reste["formes_a_traiter"], reste["complet"]) == (2, False)
    for forme in _formes(complet):
        forme.update(statut="valide", revu_le=JOUR)
    assert reste_a_traiter(complet)["complet"]


def test_forme_ajoutee_dans_une_colonne_sensible(tmp_path):
    dictionnaire = _dictionnaire()
    dictionnaire["clients"]["type_commerce"]["format"]["motif"] = "colonne sensible"
    chemin = exporter_revue(dictionnaire, "2026-03", tmp_path / "revue.xlsx", nb_revues=0)
    _modifier(chemin, "regles", COMMERCE_FORMAT, statut="documenté",
              regle_retenue="formes fermées")
    _modifier(chemin, "formes", 2, **COMMERCE, forme="9 99 99 99 999 999 99",
              statut="documenté")
    _modifier(chemin, "formes", 3, **COMMERCE, forme="999999999999999", statut="documenté")
    resultat, _, entree = appliquer_revue(dictionnaire, _importer(chemin, dictionnaire), chemin, JOUR)
    liste = resultat["clients"]["type_commerce"]["format"]["liste"]
    # L'effectif est masqué : le dictionnaire ne dit pas si la forme est observée.
    assert [(f["forme"], f["effectif"]) for f in liste] == [
        ("9 99 99 99 999 999 99", None), ("999999999999999", None),
    ]
    fichier = tmp_path / "dictionnaire.yaml"
    ecrire_dictionnaire(dictionnaire, "2026-03", fichier)
    ecrire_dictionnaire(resultat, "2026-03", fichier, revue=entree)
    assert charger_dictionnaire(fichier) == resultat

    chemin = exporter_revue(resultat, "2026-03", tmp_path / "revue2.xlsx", nb_revues=1)
    lues = _lire(load_workbook(chemin)["formes"])
    assert [l["effectif"] for l in lues] == ["masqué", "masqué"]
    # Remise à « observé » : la forme est retirée, comme une forme d'effectif nul.
    _modifier(chemin, "formes", dict(COMMERCE, forme="999999999999999"), statut="observé")
    resultat, rapport, _ = appliquer_revue(resultat, _importer(chemin, resultat, 1), chemin, JOUR)
    assert [f["forme"] for f in resultat["clients"]["type_commerce"]["format"]["liste"]] == [
        "9 99 99 99 999 999 99"
    ]
    assert list(rapport["evenement"]) == ["forme retirée"]


def test_classeur_exporte_avant_la_feuille_des_formes_refuse(tmp_path):
    chemin = _exporter(tmp_path)
    classeur = load_workbook(chemin)
    del classeur["formes"]
    classeur.save(chemin)
    message = _erreur_import(chemin)
    assert "feuille « formes » absente" in message
    assert "à réexporter" in message
    # Classeur à trois lignes par colonne : la ligne « format » manque.
    chemin = _exporter(tmp_path / "ancien")
    classeur = load_workbook(chemin)
    feuille = classeur["regles"]
    feuille.delete_rows(_ligne(feuille, table="clients", colonne="code", regle="format"))
    classeur.save(chemin)
    assert "('clients', 'code', 'format') absente du classeur" in _erreur_import(chemin)
