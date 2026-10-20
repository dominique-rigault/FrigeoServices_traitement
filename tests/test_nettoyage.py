"""Tests des règles de nettoyage (frigeo.nettoyage)."""

from collections import Counter
from copy import deepcopy
from datetime import date, datetime

import pandas as pd
import pytest

import frigeo.decisions as module_decisions

from frigeo.controles import _REPERTOIRE, VALEUR_MASQUEE, constats_encodage, controler
from frigeo.dictionnaire import MOTIF_SENSIBLE
from frigeo.doublons import detecter_doublons
from frigeo.nettoyage import (
    COLONNES_BILAN_NETTOYAGE,
    COLONNES_CATALOGUE_NETTOYAGE,
    COLONNES_JOURNAL,
    MAX_PASSES,
    RAISON_DESACCORD,
    RAISON_PASSES,
    RAISON_REMPLACEMENT,
    Reparation,
    catalogue_nettoyage,
    nettoyer,
    reparer_encodage,
)

AUCUNE_SENSIBLE = {}
AUCUNE_DECLARATION = {}
INSECABLE = chr(160)
REMPLACEMENT = chr(0xFFFD)
TOUS = ("cp1252", "iso8859_15", "latin_1")


@pytest.fixture(autouse=True)
def sans_fichier_de_decisions(tmp_path, monkeypatch):
    """Les décisions enregistrées sur le poste ne sont jamais lues par ces tests."""
    monkeypatch.setattr(
        module_decisions, "chemin_decisions", lambda: tmp_path / "decisions_absent.yaml"
    )


def _mal(texte, encodage="cp1252", fois=1):
    """Texte écrit en UTF-8 puis relu dans l'encodage, `fois` fois de suite."""
    for _ in range(fois):
        texte = texte.encode("utf-8").decode(encodage)
    return texte


def _regle(regle, motif="motif fictif"):
    return {
        "regle": regle,
        "proposition": regle,
        "statut": "observé",
        "motif": motif,
        "a_arbitrer": {},
        "commentaire": "",
        "revu_le": None,
    }


def _colonne(nature="texte", motif_valeurs="motif fictif"):
    """Règles d'une colonne, sans contrôle; `motif_valeurs` permet de marquer la
    colonne sensible dans le dictionnaire."""
    valeurs = _regle("aucune", motif_valeurs)
    valeurs["liste"] = []
    format_ = _regle("aucune")
    format_["liste"] = []
    return {
        "obligatoire": _regle("facultatif"),
        "nature": _regle(nature),
        "valeurs": valeurs,
        "format": format_,
    }


def _table(fichier="t.csv", premiere_ligne=2, periode=None, **colonnes) -> pd.DataFrame:
    """Table fictive avec son lignage (première ligne de données en ligne 2) et, si
    demandé, sa période d'export."""
    nombre = len(next(iter(colonnes.values())))
    donnees = pd.DataFrame(colonnes, dtype=object)
    donnees.insert(0, "num_ligne_source", range(premiere_ligne, premiere_ligne + nombre))
    donnees.insert(0, "feuille_source", None)
    donnees.insert(0, "fichier_source", fichier)
    if periode is not None:
        donnees["periode_export"] = periode
    return donnees


def _dictionnaire(tables, **particulieres):
    """Dictionnaire qui décrit exactement les tables; `particulieres` remplace les
    règles d'une colonne (« table.colonne »)."""
    lignage = {"fichier_source", "feuille_source", "num_ligne_source", "periode_export"}
    return {
        table: {
            colonne: particulieres.get(f"{table}.{colonne}", _colonne())
            for colonne in donnees.columns
            if colonne not in lignage
        }
        for table, donnees in tables.items()
    }


def _lignes(tableau, *colonnes):
    return [tuple(ligne) for ligne in tableau[list(colonnes)].itertuples(index=False)]


# --- reparer_encodage ---------------------------------------------------------


def test_reparer_encodage_une_passe():
    assert reparer_encodage(_mal("Café du Port")) == Reparation("Café du Port", None, 1, TOUS)
    # L'apostrophe typographique n'est repérée qu'en Windows-1252.
    assert reparer_encodage(_mal("l’atelier")) == Reparation("l’atelier", None, 1, ("cp1252",))
    assert reparer_encodage(_mal("Crêperie", "latin_1")) == Reparation(
        "Crêperie", None, 1, TOUS
    )


def test_reparer_encodage_reunit_les_encodages():
    # `Ã©` est repéré par les trois encodages, `â€™` par Windows-1252 seul : les
    # réparations ne se contredisent pas, elles sont réunies en une passe.
    texte = _mal("Café l’atelier")
    assert reparer_encodage(texte) == Reparation("Café l’atelier", None, 1, TOUS)
    # Le reste du texte est gardé tel quel, même un caractère correct que
    # l'encodage ne connaît pas.
    # (`Ã‰` n'est repéré qu'en Windows-1252, où 0x89 est le signe pour mille.)
    assert reparer_encodage("→ " + _mal("Élise") + " ’") == Reparation(
        "→ Élise ’", None, 1, ("cp1252",)
    )


def test_reparer_encodage_plusieurs_passes():
    # Au plus 5 passes, valeur validée par Dominique le 19/10/2026.
    assert MAX_PASSES == 5
    for fois in range(2, 6):
        assert reparer_encodage(_mal("Crêperie Élise", fois=fois)) == Reparation(
            "Crêperie Élise", None, fois, TOUS
        )
    assert reparer_encodage(_mal("é", fois=6)) == Reparation(None, RAISON_PASSES, 5, TOUS)
    assert RAISON_PASSES == "séquence toujours présente après 5 passes"


def test_reparer_encodage_refus():
    # `Ã€` donne `À` en Windows-1252 et `ä` en ISO-8859-15.
    assert reparer_encodage("Ã€") == Reparation(
        None, RAISON_DESACCORD, 1, ("cp1252", "iso8859_15")
    )
    assert reparer_encodage("Caf" + _mal("é") + " Ã€") == Reparation(
        None, RAISON_DESACCORD, 1, ("cp1252", "iso8859_15", "latin_1")
    )
    # Une cellule qui contient le caractère de remplacement n'est pas touchée, même
    # si elle porte aussi une séquence réparable.
    assert reparer_encodage("perdu " + REMPLACEMENT) == Reparation(
        None, RAISON_REMPLACEMENT, 0, ()
    )
    assert reparer_encodage(_mal("Café ") + REMPLACEMENT) == Reparation(
        None, RAISON_REMPLACEMENT, 0, ()
    )


def test_reparer_encodage_textes_corrects_inchanges():
    for texte in ("Café", "CHÂTEAU", "SÃO PAULO", "café" + INSECABLE + "»", "abc", ""):
        assert reparer_encodage(texte) == Reparation(texte, None, 0, ()), texte


def test_aucune_reparation_fausse():
    # Chaque caractère du répertoire, écrit en UTF-8 puis relu dans chacun des trois
    # encodages, est soit rendu à l'identique, soit refusé pour désaccord : jamais
    # remplacé par un autre caractère.
    issues = Counter()
    for caractere in sorted(_REPERTOIRE):
        for encodage in TOUS:
            try:
                relu = _mal(caractere, encodage)
            except UnicodeDecodeError:
                continue
            reparation = reparer_encodage("a" + relu + "b")
            if reparation.texte is None:
                assert reparation.raison == RAISON_DESACCORD, (caractere, encodage)
                issues[(encodage, "refus")] += 1
            else:
                assert reparation.texte == "a" + caractere + "b", (caractere, encodage)
                assert constats_encodage(reparation.texte) == ()
                issues[(encodage, "réparé")] += 1
    assert issues == {
        ("cp1252", "réparé"): 109,
        ("cp1252", "refus"): 8,
        ("iso8859_15", "réparé"): 115,
        ("iso8859_15", "refus"): 8,
        ("latin_1", "réparé"): 123,
    }


# --- nettoyer -----------------------------------------------------------------


def _exemple():
    clients = _table(
        "clients.csv",
        code=["C1", "C2", "C3", "C4", "C5", "C6", "C7"],
        nom=[
            "Café",
            _mal("Café"),
            None,
            "  ",
            "perdu " + REMPLACEMENT,
            "Ã€",
            _mal("Crêperie", fois=2),
        ],
        ville=["Rouen", "Évreux", _mal("Évreux"), "Rouen", 12, datetime(2025, 3, 14), "Rouen"],
    )
    clients.index = range(10, 17)
    factures = _table("factures_2025.xlsx", periode="2025", num=["F1", "F2"],
                      libelle=[_mal("Forfait déplacement"), "Main d'oeuvre"],
                      montant=[31.85, 46.8])
    return {"clients": clients, "factures": factures}


def test_nettoyer_ne_modifie_pas_les_tables_recues():
    tables = _exemple()
    avant = deepcopy(tables)
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    for table, donnees in tables.items():
        pd.testing.assert_frame_equal(donnees, avant[table])
        nettoyee = resultat.tables[table]
        assert nettoyee is not donnees
        assert list(nettoyee.columns) == list(donnees.columns)
        assert list(nettoyee.index) == list(donnees.index)
        assert list(nettoyee.dtypes) == list(donnees.dtypes)
    clients = resultat.tables["clients"]
    # Seules les cellules corrigées changent.
    def pareilles(a, b):
        return a is b or a == b or (pd.isna(a) and pd.isna(b))

    differentes = [
        (index, colonne)
        for colonne in clients.columns
        for index in clients.index
        if not pareilles(clients.at[index, colonne], avant["clients"].at[index, colonne])
    ]
    assert differentes == [(11, "nom"), (16, "nom"), (12, "ville")]
    assert clients.at[11, "nom"] == "Café"
    assert clients.at[16, "nom"] == "Crêperie"
    assert clients.at[12, "ville"] == "Évreux"
    assert clients.at[14, "ville"] == 12
    assert clients.at[15, "ville"] == datetime(2025, 3, 14)
    assert resultat.tables["factures"].at[0, "libelle"] == "Forfait déplacement"
    assert resultat.tables["factures"].at[0, "montant"] == 31.85


def test_nettoyer_journal_et_bilan():
    tables = _exemple()
    declarations = {"clients": {"cle": ["code"]}, "factures": {"cle": ["num"]}}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, declarations)
    journal, bilan = resultat.journal, resultat.bilan
    assert list(journal.columns) == list(COLONNES_JOURNAL)
    assert list(bilan.columns) == list(COLONNES_BILAN_NETTOYAGE)
    # Ordre : tables, puis lignes, puis colonnes.
    assert _lignes(
        journal, "regle_nettoyage", "table", "fichier_source", "feuille_source",
        "num_ligne_source", "periode_export", "cle", "colonne", "action", "source",
    ) == [
        ("encodage", "clients", "clients.csv", None, 3, None, "C2", "nom",
         "valeur modifiée", "règle générique"),
        ("encodage", "clients", "clients.csv", None, 4, None, "C3", "ville",
         "valeur modifiée", "règle générique"),
        ("encodage", "clients", "clients.csv", None, 8, None, "C7", "nom",
         "valeur modifiée", "règle générique"),
        ("encodage", "factures", "factures_2025.xlsx", None, 2, "2025", "F1", "libelle",
         "valeur modifiée", "règle générique"),
    ]
    assert _lignes(journal, "valeur_avant", "valeur_apres", "detail") == [
        (_mal("Café"), "Café", "repéré en cp1252, iso8859_15, latin_1; 1 passe"),
        (_mal("Évreux"), "Évreux", "repéré en cp1252; 1 passe"),
        (_mal("Crêperie", fois=2), "Crêperie",
         "repéré en cp1252, iso8859_15, latin_1; 2 passes"),
        (_mal("Forfait déplacement"), "Forfait déplacement",
         "repéré en cp1252, iso8859_15, latin_1; 1 passe"),
    ]
    # Une ligne par colonne qui a au moins une cellule visée : les cellules que la
    # règle encodage de l'étape 3 signale, caractère de remplacement compris.
    assert _lignes(bilan, *COLONNES_BILAN_NETTOYAGE) == [
        ("encodage", "clients", "nom", 4, 2, 2,
         "caractère de remplacement 1; encodages en désaccord 1"),
        ("encodage", "clients", "ville", 1, 1, 0, None),
        ("encodage", "factures", "libelle", 1, 1, 0, None),
    ]


def test_nettoyer_cellules_visees_identiques():
    # Une même valeur répétée est réparée une fois par cellule, chacune au journal.
    tables = {"t": _table(nom=[_mal("Café"), "x", _mal("Café")], code=["a", "b", "c"])}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, {"t": {"cle": ["code"]}})
    assert list(resultat.tables["t"]["nom"]) == ["Café", "x", "Café"]
    assert _lignes(resultat.journal, "num_ligne_source", "cle") == [(2, "a"), (4, "c")]
    assert _lignes(resultat.bilan, "visees", "corrigees") == [(2, 2)]


def test_nettoyer_colonnes_sensibles_masquees():
    tables = {
        "t": _table(
            code=["a", "b"],
            nom=[_mal("Café"), "x"],
            adresse=[_mal("rue Élise"), "y"],
            ville=[_mal("Évreux"), "z"],
        )
    }
    dictionnaire = _dictionnaire(tables, **{"t.adresse": _colonne(motif_valeurs=MOTIF_SENSIBLE)})
    sensibles = {"t": frozenset({"nom"})}
    resultat = nettoyer(tables, dictionnaire, sensibles, {"t": {"cle": ["code"]}})
    # Les cellules sensibles sont corrigées comme les autres, mais leurs valeurs sont
    # masquées au journal (classification ou dictionnaire).
    assert list(resultat.tables["t"].loc[0, ["nom", "adresse", "ville"]]) == [
        "Café", "rue Élise", "Évreux"
    ]
    assert _lignes(resultat.journal, "colonne", "cle", "valeur_avant", "valeur_apres") == [
        ("nom", "a", VALEUR_MASQUEE, VALEUR_MASQUEE),
        ("adresse", "a", VALEUR_MASQUEE, VALEUR_MASQUEE),
        ("ville", "a", _mal("Évreux"), "Évreux"),
    ]
    # Une clé dont une colonne est sensible est masquée.
    resultat = nettoyer(tables, dictionnaire, {"t": frozenset({"code"})}, {"t": {"cle": ["code"]}})
    assert set(resultat.journal["cle"]) == {VALEUR_MASQUEE}
    assert _lignes(resultat.journal, "colonne", "valeur_apres") == [
        ("nom", "Café"), ("adresse", VALEUR_MASQUEE), ("ville", "Évreux"),
    ]


def test_nettoyer_cle_composee_ou_absente():
    tables = {"t": _table(a=["1", "1"], b=["x", "y"], nom=["n", _mal("é")])}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, {"t": {"cle": ["a", "b"]}})
    assert _lignes(resultat.journal, "cle", "periode_export") == [("1 | y", None)]
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    assert _lignes(resultat.journal, "cle") == [(None,)]


def test_nettoyer_sans_cellule_visee():
    tables = {"t": _table(nom=["Café", "CHÂTEAU", None], n=[1, 2, 3])}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    assert resultat.journal.empty and list(resultat.journal.columns) == list(COLONNES_JOURNAL)
    assert resultat.bilan.empty
    assert list(resultat.bilan.columns) == list(COLONNES_BILAN_NETTOYAGE)
    pd.testing.assert_frame_equal(resultat.tables["t"], tables["t"])


def test_nettoyer_verifie_sensibilite_et_dictionnaire():
    tables = {"t": _table(nom=[_mal("Café")])}
    with pytest.raises(RuntimeError, match="sensibilité incohérente"):
        nettoyer(tables, _dictionnaire(tables), {"t": frozenset({"absente"})}, {})
    dictionnaire = _dictionnaire(tables)
    del dictionnaire["t"]["nom"]
    with pytest.raises(ValueError, match="à régénérer"):
        nettoyer(tables, dictionnaire, AUCUNE_SENSIBLE, {})
    dictionnaire = _dictionnaire(tables)
    del dictionnaire["t"]["nom"]["format"]
    with pytest.raises(ValueError, match="Dictionnaire invalide"):
        nettoyer(tables, dictionnaire, AUCUNE_SENSIBLE, {})


def test_catalogue_nettoyage():
    catalogue = catalogue_nettoyage()
    assert list(catalogue.columns) == list(COLONNES_CATALOGUE_NETTOYAGE)
    assert _lignes(catalogue, "rang", "regle_nettoyage", "source") == [
        (1, "encodage", "règle générique"),
        (2, "remplacement", "dictionnaire"),
        (3, "nature", "règle générique"),
        (4, "doublons", "règle générique"),
        (5, "decision", "décision du métier"),
    ]
    assert "caractère de remplacement" in catalogue.at[0, "ne_corrige_jamais"]
    assert "hors liste" in catalogue.at[1, "ne_corrige_jamais"]
    assert "date impossible" in catalogue.at[2, "ne_corrige_jamais"]
    assert "identiques hors clé" in catalogue.at[3, "ne_corrige_jamais"]
    assert "décidé de garder" in catalogue.at[4, "ne_corrige_jamais"]


# --- remplacement -------------------------------------------------------------


def _valeur(valeur, statut="valide", origine="liste proposée", effectif=1, remplacement=None,
            revu_le="2026-10-08 13:05:11"):
    return {
        "valeur": valeur,
        "origine": origine,
        "effectif": effectif,
        "statut": statut,
        "commentaire": "",
        "revu_le": None if statut == "observé" else revu_le,
        "remplacement": remplacement,
    }


def _liste(*elements, regle="liste fermée", motif="motif fictif"):
    """Colonne dont la règle `valeurs` porte la liste donnée, revue par le métier."""
    colonne = _colonne()
    colonne["valeurs"] = {
        **_regle(regle, motif),
        "statut": "valide",
        "revu_le": "2026-10-08 13:05:11",
        "liste": list(elements),
    }
    return colonne


LISTE_CATEGORIES = (
    _valeur("Dépannage"),
    _valeur("Maintenance"),
    _valeur("depannage", "invalide", remplacement="Dépannage"),
    _valeur("Dépanage", "invalide", remplacement="Dépannage", revu_le="2026-10-09 18:50:53"),
    _valeur("DEP", "invalide"),
    _valeur("12", "invalide", remplacement="Maintenance"),
    _valeur("Maint", "observé", origine="à arbitrer"),
)


def test_remplacement_des_valeurs_invalidees():
    tables = {
        "t": _table(
            code=["a", "b", "c", "d", "e", "f", "g", "h", "i", "j"],
            categorie=["Dépannage", "depannage", "DEP", "Autre", "Maint", None, "  ", 12,
                       "Dépanage", "depannage"],
        )
    }
    dictionnaire = _dictionnaire(tables, **{"t.categorie": _liste(*LISTE_CATEGORIES)})
    resultat = nettoyer(tables, dictionnaire, AUCUNE_SENSIBLE, {"t": {"cle": ["code"]}})
    assert list(resultat.tables["t"]["categorie"]) == [
        "Dépannage", "Dépannage", "DEP", "Autre", "Maint", None, "  ", 12, "Dépannage",
        "Dépannage",
    ]
    assert list(tables["t"]["categorie"])[1] == "depannage"
    assert _lignes(
        resultat.journal, "regle_nettoyage", "num_ligne_source", "cle", "colonne", "action",
        "source", "valeur_avant", "valeur_apres", "detail",
    ) == [
        ("remplacement", 3, "b", "categorie", "valeur modifiée", "dictionnaire", "depannage",
         "Dépannage", "valeur invalidée le 2026-10-08"),
        ("remplacement", 10, "i", "categorie", "valeur modifiée", "dictionnaire", "Dépanage",
         "Dépannage", "valeur invalidée le 2026-10-09"),
        ("remplacement", 11, "j", "categorie", "valeur modifiée", "dictionnaire", "depannage",
         "Dépannage", "valeur invalidée le 2026-10-08"),
    ]
    # Valeurs visées : invalides ou hors liste; « Maint », à arbitrer, ne l'est pas,
    # ni les cellules vides. La cellule 12 (nombre) n'est pas remplacée par un texte.
    assert _lignes(resultat.bilan, *COLONNES_BILAN_NETTOYAGE) == [
        ("remplacement", "t", "categorie", 6, 3, 3,
         "valeur invalidée sans remplacement 1; valeur hors liste 1; valeur non textuelle 1"),
    ]
    # Mêmes cellules que les anomalies de la règle valeurs à l'étape 3.
    ecarts, _ = controler(tables, dictionnaire, AUCUNE_SENSIBLE, {})
    visees = ecarts[(ecarts["regle"] == "valeurs") & (ecarts["issue"] == "anomalie")]
    assert len(visees) == 6
    # Après nettoyage, il ne reste que les trois non corrigées.
    ecarts, _ = controler(resultat.tables, dictionnaire, AUCUNE_SENSIBLE, {})
    restantes = ecarts[(ecarts["regle"] == "valeurs") & (ecarts["issue"] == "anomalie")]
    assert list(restantes["num_ligne_source"]) == [4, 5, 9]


def test_remplacement_seulement_pour_une_liste_fermee():
    # Une liste écartée (règle « aucune ») ne donne ni contrôle ni remplacement, même
    # si elle garde des valeurs invalidées avec leur remplacement.
    tables = {"t": _table(categorie=["depannage", "Autre"])}
    colonne = _liste(*LISTE_CATEGORIES, regle="aucune")
    colonne["valeurs"]["statut"] = "observé"
    colonne["valeurs"]["revu_le"] = None
    dictionnaire = _dictionnaire(tables, **{"t.categorie": colonne})
    resultat = nettoyer(tables, dictionnaire, AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    assert resultat.journal.empty and resultat.bilan.empty
    assert list(resultat.tables["t"]["categorie"]) == ["depannage", "Autre"]


def test_remplacement_apres_reparation_de_l_encodage():
    # La règle remplacement s'applique à la valeur laissée par la règle encodage.
    # Le journal suit l'ordre des règles, puis des lignes.
    tables = {"t": _table(categorie=["depannage", _mal("Dépanage"), "Dépannage"], n=[1, 2, 3])}
    dictionnaire = _dictionnaire(tables, **{"t.categorie": _liste(*LISTE_CATEGORIES)})
    resultat = nettoyer(tables, dictionnaire, AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    assert list(resultat.tables["t"]["categorie"]) == ["Dépannage"] * 3
    assert _lignes(
        resultat.journal, "regle_nettoyage", "num_ligne_source", "source", "valeur_avant",
        "valeur_apres",
    ) == [
        ("encodage", 3, "règle générique", _mal("Dépanage"), "Dépanage"),
        ("remplacement", 2, "dictionnaire", "depannage", "Dépannage"),
        ("remplacement", 3, "dictionnaire", "Dépanage", "Dépannage"),
    ]
    assert _lignes(resultat.bilan, "regle_nettoyage", "visees", "corrigees") == [
        ("encodage", 1, 1),
        ("remplacement", 2, 2),
    ]


def test_remplacement_sur_colonne_sensible_masque():
    sensibles_liste = [
        _valeur("Démission", "documenté", effectif=None, origine="ajoutée"),
        _valeur("demission", "invalide", effectif=None, origine="observée",
                remplacement="Démission"),
    ]
    tables = {"t": _table(code=["a", "b"], motif=["demission", "Démission"])}
    colonne = _liste(*sensibles_liste, motif=MOTIF_SENSIBLE)
    dictionnaire = _dictionnaire(tables, **{"t.motif": colonne})
    resultat = nettoyer(tables, dictionnaire, AUCUNE_SENSIBLE, {"t": {"cle": ["code"]}})
    assert list(resultat.tables["t"]["motif"]) == ["Démission", "Démission"]
    assert _lignes(resultat.journal, "cle", "valeur_avant", "valeur_apres", "detail") == [
        ("a", VALEUR_MASQUEE, VALEUR_MASQUEE, "valeur invalidée le 2026-10-08"),
    ]


# --- nature -------------------------------------------------------------------

COUPURE = date(2026, 3, 31)


def _nature(valeurs, nature, sensibles=AUCUNE_SENSIBLE, date_coupure=COUPURE):
    """Nettoie une colonne « v » de la nature donnée; rend (valeurs nettoyées,
    journal, bilan). La colonne « n » rend les lignes toutes différentes : sans
    elle, deux valeurs converties en la même deviendraient des doublons exacts."""
    tables = {"t": _table(v=list(valeurs), n=list(range(len(valeurs))))}
    dictionnaire = _dictionnaire(tables, **{"t.v": _colonne(nature)})
    resultat = nettoyer(tables, dictionnaire, sensibles, AUCUNE_DECLARATION, date_coupure)
    return list(resultat.tables["t"]["v"]), resultat.journal, resultat.bilan, (
        tables, dictionnaire, resultat.tables
    )


def test_nature_dates_en_texte():
    valeurs, journal, bilan, _ = _nature(
        [
            "14/03/2025",
            "2025-03-14",
            "14.03.2025",
            "14-03-2025",
            "20250314",
            "14/03/25",
            "14.03.25",
            datetime(2025, 3, 14),
            "31/02/2025",
            "15/06/98",
            datetime(2025, 3, 14, 10, 30),
            45,
            "1h30",
            "14/3/2025",
            "10:30",
            "14/03-2025",
        ],
        "date JJ/MM/AAAA",
    )
    assert valeurs[:8] == ["14/03/2025"] * 8
    assert valeurs[8:] == [
        "31/02/2025", "15/06/98", datetime(2025, 3, 14, 10, 30), 45, "1h30", "14/3/2025",
        "10:30", "14/03-2025",
    ]
    assert set(journal["source"]) == {"règle générique"}
    assert _lignes(journal, "regle_nettoyage", "num_ligne_source", "valeur_avant", "detail") == [
        ("nature", 3, "2025-03-14", "forme lue 9999-99-99"),
        ("nature", 4, "14.03.2025", "forme lue 99.99.9999"),
        ("nature", 5, "14-03-2025", "forme lue 99-99-9999"),
        ("nature", 6, "20250314", "forme lue 99999999"),
        ("nature", 7, "14/03/25", "forme lue 99/99/99, siècle 20 ajouté"),
        ("nature", 8, "14.03.25", "forme lue 99.99.99, siècle 20 ajouté"),
        ("nature", 9, "2025-03-14 00:00:00", "forme lue 9999-99-99 99:99:99"),
    ]
    assert _lignes(bilan, *COLONNES_BILAN_NETTOYAGE) == [
        ("nature", "t", "v", 15, 7, 8,
         "date impossible 1; année sur deux chiffres, date future 1; date avec une heure 1; "
         "autre famille de nature 2; forme non reconnue 3"),
    ]


def test_nature_date_native_et_date_de_coupure():
    valeurs, journal, bilan, _ = _nature(
        ["04/09/2019", "2019-09-04", "04-09-19", "31/03/26", "01/04/26", datetime(2020, 1, 1)],
        "date native",
    )
    assert valeurs == [
        datetime(2019, 9, 4), datetime(2019, 9, 4), datetime(2019, 9, 4), datetime(2026, 3, 31),
        "01/04/26", datetime(2020, 1, 1),
    ]
    assert all(type(v) is datetime for v in valeurs if not isinstance(v, str))
    assert _lignes(bilan, "visees", "corrigees", "raisons") == [
        (5, 4, "année sur deux chiffres, date future 1")
    ]
    # Sans date de coupure, une année sur deux chiffres n'est jamais complétée.
    valeurs, journal, bilan, _ = _nature(["04-09-19", "2019-09-04"], "date native",
                                         date_coupure=None)
    assert valeurs == ["04-09-19", datetime(2019, 9, 4)]
    assert _lignes(bilan, "visees", "corrigees", "raisons") == [
        (2, 1, "année sur deux chiffres, date de coupure non fournie 1")
    ]


def test_nature_nombre_natif():
    valeurs, journal, bilan, _ = _nature(
        [
            1.5,
            "181,45 €",
            "181.45",
            "1 234,50",
            "1" + INSECABLE + "234,50 €",
            "1" + chr(0x202F) + "234",
            "-12,5",
            "95",
            "1.234",
            "45 min",
            "1h30",
            "14/03/2025",
            "12 $",
            "12,5€",
            "abc",
        ],
        "nombre natif",
    )
    assert valeurs == [
        1.5, 181.45, 181.45, 1234.5, 1234.5, 1234.0, -12.5, 95.0, "1.234", "45 min", "1h30",
        "14/03/2025", "12 $", 12.5, "abc",
    ]
    assert all(type(v) is float for v in valeurs if not isinstance(v, str))
    assert _lignes(journal, "num_ligne_source", "detail") == [
        (3, "forme lue 999,99 €"),
        (4, "forme lue 999.99"),
        (5, "forme lue 9 999,99"),
        (6, "forme lue 9" + INSECABLE + "999,99 €"),
        (7, "forme lue 9" + chr(0x202F) + "999"),
        (8, "forme lue -99,9"),
        (9, "forme lue 99"),
        (15, "forme lue 99,9€"),
    ]
    assert _lignes(bilan, "visees", "corrigees", "raisons") == [
        (14, 8, "valeur avec une unité 2; séparateur ambigu 1; autre famille de nature 1; "
                "forme non reconnue 2"),
    ]


def test_nature_nombres_en_texte():
    valeurs, *_ = _nature(
        ["12,30", "999.99", "9 999,99", "95", 95.5, "12,30 €", "-0.50", 7],
        "décimal virgule (texte)",
    )
    assert valeurs == ["12,30", "999,99", "9999,99", "95,0", "95,5", "12,30", "-0,50", "7,0"]
    valeurs, *_ = _nature(["12.30", "12,30", "95"], "décimal point (texte)")
    assert valeurs == ["12.30", "12.30", "95.0"]
    valeurs, _, bilan, _ = _nature(
        ["12", "99 999", "12,0", "12,00 €", "12,5", 7.0, 7.5, "1h30"], "entier (texte)"
    )
    assert valeurs == ["12", "99999", "12", "12", "12,5", "7", 7.5, "1h30"]
    assert _lignes(bilan, "visees", "corrigees", "raisons") == [
        (7, 4, "valeur avec une unité 1; partie décimale non nulle 2"),
    ]


def test_nature_heure_et_colonne_texte():
    valeurs, journal, bilan, _ = _nature(["08:00", "7:30", "45", "07h30"], "heure HH:MM")
    assert valeurs == ["08:00", "7:30", "45", "07h30"]
    assert journal.empty
    assert _lignes(bilan, "visees", "corrigees", "raisons") == [
        (3, 0, "autre famille de nature 1; forme non reconnue 2"),
    ]
    # Une colonne de texte n'a pas de contrôle de nature : rien n'est visé.
    valeurs, journal, bilan, _ = _nature(["2025-03-14", "1 234,50"], "texte")
    assert valeurs == ["2025-03-14", "1 234,50"] and journal.empty and bilan.empty


def test_nature_coherente_avec_le_controle():
    # Après nettoyage, seules les cellules non converties restent hors nature.
    for valeurs, nature in (
        (["14.03.2025", "31/02/2025", "14/03/25", datetime(2025, 3, 14), "x"], "date JJ/MM/AAAA"),
        (["181,45 €", "1.234", 3, "9 999,99"], "nombre natif"),
        (["999.99", "9 999,99", "1h30"], "décimal virgule (texte)"),
        (["99 999", "12,5", "12,0"], "entier (texte)"),
    ):
        _, _, bilan, (tables, dictionnaire, nettoyees) = _nature(valeurs, nature)
        avant, _ = controler(tables, dictionnaire, AUCUNE_SENSIBLE, {})
        apres, _ = controler(nettoyees, dictionnaire, AUCUNE_SENSIBLE, {})
        hors = (avant["regle"] == "nature").sum()
        assert hors == bilan.at[0, "visees"], nature
        assert (apres["regle"] == "nature").sum() == bilan.at[0, "non_corrigees"], nature


def test_nature_colonne_sensible_masquee():
    valeurs, journal, _, _ = _nature(["14.03.2025"], "date JJ/MM/AAAA", {"t": frozenset({"v"})})
    assert valeurs == ["14/03/2025"]
    assert _lignes(journal, "valeur_avant", "valeur_apres", "detail") == [
        (VALEUR_MASQUEE, VALEUR_MASQUEE, "forme masquée")
    ]
    _, journal, _, _ = _nature(["14.03.25"], "date JJ/MM/AAAA", {"t": frozenset({"v"})})
    assert list(journal["detail"]) == ["forme masquée, siècle 20 ajouté"]


# --- doublons -----------------------------------------------------------------


def test_doublons_exacts_retires_premiere_occurrence_gardee():
    clients = _table(
        "clients.csv",
        code=["C1", "C2", "C1", "C3", "C3", "C4", "C5", "C1"],
        nom=["Café", "Bar", "Café", "Épicerie", "Epicerie ", "Bar", "Bar", "Café"],
        ville=["Rouen", "Évreux", "Rouen", "Vernon", "vernon", "Évreux", "Évreux", "Rouen"],
    )
    clients.index = range(10, 18)
    factures = pd.concat(
        [
            _table("factures_2025.xlsx", periode="2025", num=["F1", "F2"], total=[10, 20]),
            _table("factures_2026-01.xlsx", periode="2026-01", num=["F2"], total=[20]),
        ],
        ignore_index=True,
    )
    factures["feuille_source"] = "Entetes"
    tables = {"clients": clients, "factures": factures}
    declarations = {"clients": {"cle": ["code"]}, "factures": {"cle": ["num"]}}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, declarations)
    # Lignes 4 et 9 : copies de la ligne 2; ligne 6 : copie de la ligne 5 après
    # normalisation (casse, accents, espaces). C4 et C5 sont identiques à C2 hors clé :
    # à arbitrer, gardées. La facture F2 est en double entre deux exports.
    assert list(resultat.tables["clients"].index) == [10, 11, 13, 15, 16]
    assert list(resultat.tables["clients"]["num_ligne_source"]) == [2, 3, 5, 7, 8]
    assert list(resultat.tables["factures"]["fichier_source"]) == [
        "factures_2025.xlsx", "factures_2025.xlsx"
    ]
    assert _lignes(
        resultat.journal, "regle_nettoyage", "table", "fichier_source", "num_ligne_source",
        "periode_export", "cle", "colonne", "action", "source", "valeur_avant",
        "valeur_apres", "detail",
    ) == [
        ("doublons", "clients", "clients.csv", 4, None, "C1", None, "ligne retirée",
         "règle générique", None, None, "doublon exact de la ligne 2 (clients.csv)"),
        ("doublons", "clients", "clients.csv", 6, None, "C3", None, "ligne retirée",
         "règle générique", None, None,
         "doublon exact de la ligne 5 (clients.csv), après normalisation"),
        ("doublons", "clients", "clients.csv", 9, None, "C1", None, "ligne retirée",
         "règle générique", None, None, "doublon exact de la ligne 2 (clients.csv)"),
        ("doublons", "factures", "factures_2026-01.xlsx", 2, "2026-01", "F2", None,
         "ligne retirée", "règle générique", None, None,
         "doublon exact de la ligne 3 (factures_2025.xlsx, Entetes)"),
    ]
    assert _lignes(resultat.bilan, *COLONNES_BILAN_NETTOYAGE) == [
        ("doublons", "clients", None, 3, 3, 0, None),
        ("doublons", "factures", None, 1, 1, 0, None),
    ]
    # Les tables reçues ne sont pas modifiées.
    assert len(tables["clients"]) == 8 and len(tables["factures"]) == 3


def test_doublons_cherches_apres_les_autres_regles():
    # Deux lignes qui ne deviennent identiques qu'après réparation de l'encodage.
    tables = {"t": _table(code=["a", "a"], nom=["Café", _mal("Café")])}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, {"t": {"cle": ["code"]}})
    assert list(resultat.tables["t"]["num_ligne_source"]) == [2]
    assert _lignes(resultat.journal, "regle_nettoyage", "num_ligne_source", "action") == [
        ("encodage", 3, "valeur modifiée"),
        ("doublons", 3, "ligne retirée"),
    ]
    # Après nettoyage, plus aucun doublon exact; une clé répétée sur des lignes
    # différentes reste (décision du métier).
    tables = {"t": _table(code=["a", "a", "b", "b"], nom=["x", "y", "z", "z"])}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, {"t": {"cle": ["code"]}})
    doublons, _ = detecter_doublons(resultat.tables, {"t": {"cle": ["code"]}}, AUCUNE_SENSIBLE)
    assert set(doublons["constat"]) == {"valeur en double"}
    assert list(resultat.tables["t"]["code"]) == ["a", "a", "b"]


def test_doublons_lignage_ambigu_non_retire():
    # Deux lignes identiques qui portent le même lignage ne sont pas retirées.
    donnees = _table(nom=["x", "x", "y", "y"])
    donnees["num_ligne_source"] = [2, 2, 3, 4]
    tables = {"t": donnees}
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    assert list(resultat.tables["t"]["num_ligne_source"]) == [2, 2, 3]
    assert _lignes(resultat.bilan, "visees", "corrigees", "non_corrigees", "raisons") == [
        (2, 1, 1, "lignage ambigu 1")
    ]



# --- règle decision -----------------------------------------------------------

AUCUNE_DECISION = {"meta": {"structure": 1, "imports": []}, "decisions": {}}


def _empreinte_cellule(valeur):
    return module_decisions._empreinte([module_decisions._forme(valeur)])


def _empreinte_ligne(**valeurs):
    return module_decisions._empreinte(
        f"{c}={module_decisions._forme(v)}" for c, v in valeurs.items()
    )


def _decision(ligne, champ, choix, empreinte, valeur=None, nature=None, table="t",
              decide_le="2026-10-20 10:00:00", fichier="t.csv"):
    return {
        "sorte": "ligne" if choix == "retirer" else "cellule",
        "table": table,
        "fichier_source": fichier,
        "feuille_source": None,
        "num_ligne_source": ligne,
        "champ": champ,
        "constat": "constat",
        "empreinte": empreinte,
        "decision": choix,
        "valeur_retenue": valeur,
        "nature": nature,
        "commentaire": None,
        "decide_le": decide_le,
        "classeur": "arbitrage_2026-03.xlsx",
    }


def _decisions(**par_id):
    return {"meta": {"structure": 1, "imports": []}, "decisions": par_id}


def _nettoyer_avec(tables, decisions, dictionnaire=None, sensibles=AUCUNE_SENSIBLE,
                   declarations=AUCUNE_DECLARATION):
    dictionnaire = dictionnaire or _dictionnaire(tables)
    return nettoyer(tables, dictionnaire, sensibles, declarations, None, decisions)


def _du_regle(tableau, regle="decision"):
    return tableau[tableau["regle_nettoyage"] == regle]


def test_decision_corriger():
    tables = {"t": _table(code=["a", "b", "c"], d=["01/01/2025", "31/02/2025", "x"])}
    dictionnaire = _dictionnaire(tables, **{"t.d": _colonne("date JJ/MM/AAAA")})
    decisions = _decisions(
        a1=_decision(3, "d", "corriger", _empreinte_cellule("31/02/2025"), "28/02/2025",
                     "date JJ/MM/AAAA"),
    )
    resultat = _nettoyer_avec(tables, decisions, dictionnaire)
    assert list(resultat.tables["t"]["d"]) == ["01/01/2025", "28/02/2025", "x"]
    assert tables["t"].at[1, "d"] == "31/02/2025"
    assert _lignes(_du_regle(resultat.journal), "table", "num_ligne_source", "colonne",
                   "action", "source", "valeur_avant", "valeur_apres", "detail") == [
        ("t", 3, "d", "valeur modifiée", "décision du métier", "31/02/2025", "28/02/2025",
         "décision du 2026-10-20, classeur arbitrage_2026-03.xlsx, anomalie a1"),
    ]
    assert _lignes(_du_regle(resultat.bilan), "table", "colonne", "visees", "corrigees",
                   "non_corrigees", "raisons") == [("t", "d", 1, 1, 0, None)]
    # Sans décision, rien.
    resultat = _nettoyer_avec(tables, AUCUNE_DECISION, dictionnaire)
    assert _du_regle(resultat.bilan).empty and _du_regle(resultat.journal).empty


def test_decision_apres_les_regles_automatiques():
    # L'empreinte est celle de la valeur laissée par les règles 1 à 4 (encodage réparé).
    tables = {"t": _table(code=["a", "b"], nom=[_mal("Café"), "Thé"])}
    decisions = _decisions(
        a1=_decision(2, "nom", "corriger", _empreinte_cellule("Café"), "Café du Port", "texte"),
        a2=_decision(3, "nom", "corriger", _empreinte_cellule("Thé vert"), "Thé noir", "texte"),
    )
    resultat = _nettoyer_avec(tables, decisions)
    assert list(resultat.tables["t"]["nom"]) == ["Café du Port", "Thé"]
    assert _lignes(resultat.journal, "regle_nettoyage", "num_ligne_source", "valeur_apres") == [
        ("encodage", 2, "Café"),
        ("decision", 2, "Café du Port"),
    ]
    assert _lignes(_du_regle(resultat.bilan), "visees", "corrigees", "non_corrigees",
                   "raisons") == [(2, 1, 1, "source modifiée depuis la décision 1")]


def test_decision_nature_changee():
    tables = {"t": _table(code=["a", "b"], d=["x", "y"], n=[1.5, "z"])}
    dictionnaire = _dictionnaire(
        tables, **{"t.d": _colonne("date JJ/MM/AAAA"), "t.n": _colonne("nombre natif")}
    )
    decisions = _decisions(
        # Enregistrée dans la nature AAAAMMJJ, relue en JJ/MM/AAAA.
        a1=_decision(2, "d", "corriger", _empreinte_cellule("x"), "20250228", "date AAAAMMJJ"),
        a2=_decision(3, "n", "corriger", _empreinte_cellule("z"), "abc", "texte"),
    )
    resultat = _nettoyer_avec(tables, decisions, dictionnaire)
    assert list(resultat.tables["t"]["d"]) == ["28/02/2025", "y"]
    assert list(resultat.tables["t"]["n"]) == [1.5, "z"]
    assert _lignes(_du_regle(resultat.bilan), "colonne", "corrigees", "raisons") == [
        ("d", 1, None),
        ("n", 0, "valeur retenue illisible dans la nature en vigueur 1"),
    ]


def test_decision_type_de_la_colonne_garde():
    tables = {"t": _table(code=["a", "b"])}
    tables["t"]["m"] = [1.5, 2.5]
    tables["t"]["j"] = pd.to_datetime(["2025-01-01", "2025-01-02"])
    dictionnaire = _dictionnaire(
        tables, **{"t.m": _colonne("nombre natif"), "t.j": _colonne("date native")}
    )
    decisions = _decisions(
        a1=_decision(2, "m", "corriger", _empreinte_cellule(1.5), 3.0, "nombre natif"),
        a2=_decision(3, "j", "corriger", _empreinte_cellule(pd.Timestamp(2025, 1, 2)),
                     datetime(2025, 2, 1), "date native"),
    )
    resultat = _nettoyer_avec(tables, decisions, dictionnaire)
    assert list(resultat.tables["t"]["m"]) == [3.0, 2.5]
    assert resultat.tables["t"]["m"].dtype == "float64"
    assert resultat.tables["t"].at[1, "j"] == pd.Timestamp(2025, 2, 1)
    assert resultat.tables["t"]["j"].dtype == tables["t"]["j"].dtype


def test_decision_retirer_apres_les_corrections():
    tables = {"t": _table(code=["a", "a", "b"], nom=["x", "y", "z"])}
    decisions = _decisions(
        c1=_decision(3, "nom", "corriger", _empreinte_cellule("y"), "w", "texte"),
        # Empreinte de la ligne corrigée, comme à l'export.
        r1=_decision(3, "code", "retirer", _empreinte_ligne(code="a", nom="w")),
        # Empreinte d'une ligne qui a changé depuis.
        r2=_decision(4, "code", "retirer", _empreinte_ligne(code="b", nom="autre")),
    )
    resultat = _nettoyer_avec(tables, decisions)
    assert list(resultat.tables["t"]["num_ligne_source"]) == [2, 4]
    assert list(resultat.tables["t"].index) == [0, 2]
    assert _lignes(_du_regle(resultat.journal), "num_ligne_source", "colonne", "action") == [
        (3, "nom", "valeur modifiée"),
        (3, None, "ligne retirée"),
    ]
    assert _lignes(_du_regle(resultat.bilan), "table", "colonne", "visees", "corrigees",
                   "raisons") == [
        ("t", "nom", 1, 1, None),
        ("t", None, 2, 1, "source modifiée depuis la décision 1"),
    ]


def test_decision_retirer_apres_les_doublons_et_deux_fois_la_meme_ligne():
    # Une ligne retirée comme doublon exact ne l'est pas une seconde fois; une ligne
    # visée par deux décisions n'est retirée qu'une fois.
    tables = {"t": _table(code=["a", "a", "b"], nom=["x", "x", "z"])}
    decisions = _decisions(
        r1=_decision(3, "code", "retirer", _empreinte_ligne(code="a", nom="x")),
        r2=_decision(4, "code", "retirer", _empreinte_ligne(code="b", nom="z")),
        r3=_decision(4, "nom", "retirer", _empreinte_ligne(code="b", nom="z")),
    )
    resultat = _nettoyer_avec(tables, decisions)
    assert list(resultat.tables["t"]["num_ligne_source"]) == [2]
    assert _lignes(resultat.journal, "regle_nettoyage", "num_ligne_source", "detail") == [
        ("doublons", 3, "doublon exact de la ligne 2 (t.csv)"),
        ("decision", 4, "décision du 2026-10-20, classeur arbitrage_2026-03.xlsx, anomalie r2"),
    ]
    assert _lignes(_du_regle(resultat.bilan), "visees", "corrigees", "raisons") == [
        (3, 2, "ligne source introuvable 1")
    ]


def test_decision_introuvable_et_garder():
    donnees = _table(code=["a", "b", "c"], nom=["x", "y", "z"])
    donnees["num_ligne_source"] = [2, 3, 3]
    tables = {"t": donnees}
    decisions = _decisions(
        a1=_decision(9, "nom", "corriger", _empreinte_cellule("x"), "v", "texte"),
        a2=_decision(3, "nom", "corriger", _empreinte_cellule("y"), "v", "texte"),
        a3=_decision(2, "absente", "corriger", _empreinte_cellule("x"), "v", "texte"),
        a4=_decision(2, "nom", "corriger", _empreinte_cellule("x"), "v", "texte", table="u"),
        a5=_decision(2, "nom", "garder", _empreinte_cellule("x")),
        a6=_decision(2, "code", "corriger", _empreinte_cellule("a"), "v", "texte",
                     fichier="autre.csv"),
    )
    resultat = _nettoyer_avec(tables, decisions)
    assert resultat.tables["t"].equals(tables["t"])
    assert _du_regle(resultat.journal).empty
    assert _lignes(_du_regle(resultat.bilan), "table", "colonne", "visees", "corrigees",
                   "raisons") == [
        ("t", "code", 1, 0, "ligne source introuvable 1"),
        ("t", "nom", 2, 0, "ligne source introuvable 2"),
        ("t", "absente", 1, 0, "ligne source introuvable 1"),
        ("u", "nom", 1, 0, "ligne source introuvable 1"),
    ]


def test_decision_la_plus_ancienne_d_abord():
    tables = {"t": _table(code=["a"], nom=["x"])}
    decisions = _decisions(
        b=_decision(2, "nom", "corriger", _empreinte_cellule("x"), "récente", "texte",
                    decide_le="2026-11-20 10:00:00"),
        a=_decision(2, "nom", "corriger", _empreinte_cellule("x"), "ancienne", "texte",
                    decide_le="2026-10-20 10:00:00"),
    )
    resultat = _nettoyer_avec(tables, decisions)
    assert list(resultat.tables["t"]["nom"]) == ["ancienne"]
    assert _lignes(_du_regle(resultat.bilan), "visees", "corrigees", "raisons") == [
        (2, 1, "source modifiée depuis la décision 1")
    ]


def test_decision_colonne_sensible_masquee():
    tables = {"t": _table(code=["a"], adresse=["1 rue"])}
    decisions = _decisions(
        a1=_decision(2, "adresse", "corriger", _empreinte_cellule("1 rue"), "2 rue", "texte"),
    )
    resultat = _nettoyer_avec(tables, decisions, sensibles={"t": frozenset({"adresse"})},
                              declarations={"t": {"cle": ["adresse"]}})
    assert list(resultat.tables["t"]["adresse"]) == ["2 rue"]
    assert _lignes(resultat.journal, "cle", "valeur_avant", "valeur_apres") == [
        (VALEUR_MASQUEE, VALEUR_MASQUEE, VALEUR_MASQUEE)
    ]


def test_decisions_lues_par_defaut_et_invalides(tmp_path, monkeypatch):
    tables = {"t": _table(code=["a"], nom=["x"])}
    fichier = tmp_path / "decisions.yaml"
    monkeypatch.setattr(module_decisions, "chemin_decisions", lambda: fichier)
    # Fichier absent : aucune décision.
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    assert _du_regle(resultat.bilan).empty
    import yaml

    contenu = _decisions(
        a1=_decision(2, "nom", "corriger", _empreinte_cellule("x"), "y", "texte")
    )
    fichier.write_text(yaml.safe_dump(contenu, allow_unicode=True), encoding="utf-8")
    resultat = nettoyer(tables, _dictionnaire(tables), AUCUNE_SENSIBLE, AUCUNE_DECLARATION)
    assert list(resultat.tables["t"]["nom"]) == ["y"]
    contenu["decisions"]["a1"]["decision"] = "retirer"
    with pytest.raises(ValueError, match="Décisions invalides"):
        _nettoyer_avec(tables, contenu)


def test_cycle_complet_d_arbitrage(tmp_path):
    """Export, décisions du métier, import, puis nettoyage avec les décisions : ce
    qui est corrigé ou retiré disparaît, ce qui est gardé n'est plus proposé."""
    from openpyxl import load_workbook

    from frigeo.anomalies import reunir_anomalies
    from frigeo.decisions import (
        COLONNES_ARBITRAGE,
        ETAT_VALABLE,
        charger_decisions,
        ecrire_decisions,
        elements_a_arbitrer,
        exporter_arbitrage,
        importer_arbitrage,
    )
    from frigeo.gestion import controler_gestion

    tables = {
        "t": _table(code=["a", "a", "b", "c"], nom=["x", "y", "z", "w"],
                    d=["31/02/2025", "01/01/2025", "32/01/2025", "02/01/2025"])
    }
    dictionnaire = _dictionnaire(tables, **{"t.d": _colonne("date JJ/MM/AAAA")})
    declarations = {"t": {"cle": ["code"]}}
    fichier = tmp_path / "decisions.yaml"

    def nettoyer_et_controler():
        nettoyees = nettoyer(tables, dictionnaire, AUCUNE_SENSIBLE, declarations, None,
                             charger_decisions(fichier)).tables
        ecarts, bilan = controler(nettoyees, dictionnaire, AUCUNE_SENSIBLE, {})
        doublons, bilan_doublons = detecter_doublons(nettoyees, declarations, AUCUNE_SENSIBLE)
        ecarts_g, bilan_g = controler_gestion(
            nettoyees, dictionnaire, [], declarations, doublons, AUCUNE_SENSIBLE
        )
        anomalies = reunir_anomalies(nettoyees, ecarts, doublons, ecarts_g, bilan,
                                     bilan_doublons, bilan_g, declarations, AUCUNE_SENSIBLE)
        return nettoyees, elements_a_arbitrer(nettoyees, anomalies, charger_decisions(fichier))

    _, elements = nettoyer_et_controler()
    assert _lignes(elements, "sorte", "num_ligne_source") == [
        ("cellule", 2), ("cellule", 4), ("ligne", 2), ("ligne", 3)
    ]
    chemin = exporter_arbitrage(elements, "2026-03", tmp_path / "a.xlsx", nb_imports=0)
    classeur = load_workbook(chemin)
    feuille = classeur["arbitrage"]
    colonnes = {nom: i for i, nom in enumerate(COLONNES_ARBITRAGE, start=1)}
    choix = {2: ("corriger", "28.02.2025"), 3: ("garder", None), 5: ("retirer", None)}
    for numero, (decision, valeur) in choix.items():
        feuille.cell(row=numero, column=colonnes["decision"]).value = decision
        feuille.cell(row=numero, column=colonnes["valeur_retenue"]).value = valeur
    classeur.save(chemin)
    ecrire_decisions(importer_arbitrage(chemin, dictionnaire, nb_imports=0), "2026-03", fichier)

    nettoyees, elements = nettoyer_et_controler()
    assert list(nettoyees["t"]["num_ligne_source"]) == [2, 4, 5]
    assert list(nettoyees["t"]["d"]) == ["28/02/2025", "32/01/2025", "02/01/2025"]
    # Il ne reste que la date gardée, dont la décision est valable : rien à exporter.
    assert _lignes(elements, "sorte", "num_ligne_source", "etat_decision") == [
        ("cellule", 4, ETAT_VALABLE)
    ]


def test_decision_index_des_tables_non_continu():
    # Le journal retrouve la ligne de la table reçue par son étiquette, pas par sa
    # position dans la table nettoyée.
    donnees = _table(code=["a", "a", "b"], nom=["x", "x", "y"])
    donnees.index = [10, 11, 12]
    tables = {"t": donnees}
    decisions = _decisions(
        a1=_decision(4, "nom", "corriger", _empreinte_cellule("y"), "z", "texte"),
    )
    resultat = _nettoyer_avec(tables, decisions, declarations={"t": {"cle": ["code"]}})
    assert list(resultat.tables["t"].index) == [10, 12]
    assert list(resultat.tables["t"]["nom"]) == ["x", "z"]
    assert _lignes(resultat.journal, "regle_nettoyage", "num_ligne_source", "cle") == [
        ("doublons", 3, "a"),
        ("decision", 4, "b"),
    ]
