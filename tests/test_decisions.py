"""Tests des décisions du métier ligne par ligne (frigeo.decisions)."""

from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

import frigeo.dictionnaire as module_dictionnaire
from frigeo.anomalies import COLONNES_ANOMALIES, reunir_anomalies
from frigeo.controles import controler
from frigeo.decisions import (
    COLONNE_DECISION_METIER,
    COLONNES_ARBITRAGE,
    COLONNES_DECISIONS_LUES,
    COLONNES_ELEMENTS,
    ETAT_CADUQUE,
    ETAT_INTROUVABLE,
    ETAT_SANS_DECISION,
    ETAT_VALABLE,
    FEUILLE_ARBITRAGE,
    SORTE_CELLULE,
    SORTE_ECART,
    SORTE_LIGNE,
    ajouter_decisions,
    charger_decisions,
    ecrire_decisions,
    elements_a_arbitrer,
    exporter_arbitrage,
    importer_arbitrage,
    lire_valeur_retenue,
    repere_arbitrage,
)
from frigeo.doublons import detecter_doublons
from frigeo.gestion import controler_gestion

AUCUNE_DECISION = {"meta": {"structure": 1, "imports": []}, "decisions": {}}


def _table(fichier="t.csv", premiere_ligne=2, **colonnes) -> pd.DataFrame:
    """Table fictive avec son lignage (première ligne de données en ligne 2)."""
    nombre = len(next(iter(colonnes.values())))
    donnees = pd.DataFrame(colonnes, dtype=object)
    donnees.insert(0, "num_ligne_source", range(premiere_ligne, premiere_ligne + nombre))
    donnees.insert(0, "feuille_source", None)
    donnees.insert(0, "fichier_source", fichier)
    return donnees


def _tables():
    return {
        "t": _table(
            id=["A1", "A2", "A2", "A4"],
            d=["01/01/2025", "31/02/2025", "03/01/2025", "04/01/2052"],
            m=[10.5, None, 3.0, 4.0],
            nom=["x", "y", "z", "w"],
        )
    }


def _anomalie(etape, issue, ligne, champ, table="t", controle=None, valeur=None, **autres):
    """Ligne de la table des anomalies; `ligne` None pour un écart sans lignage."""
    return {
        "id_anomalie": autres.pop("id_anomalie", f"{etape}-{champ}-{ligne}-{issue}"),
        "etape": etape,
        "controle": controle or f"{table} {champ}",
        "table": table,
        "fichier_source": None if ligne is None else "t.csv",
        "feuille_source": None,
        "num_ligne_source": ligne,
        "periode_export": None,
        "cle": autres.pop("cle", None),
        "champ": champ,
        "valeur": valeur,
        "regle": autres.pop("regle", "règle"),
        "constat": autres.pop("constat", "constat"),
        "issue": issue,
        "statut_regle": None,
        "detail": autres.pop("detail", None),
        "proposition": autres.pop("proposition", None),
        "gravite": None,
        "indicateurs": None,
    }


def _anomalies(*lignes):
    anomalies = pd.DataFrame(list(lignes), columns=list(COLONNES_ANOMALIES), dtype=object)
    anomalies["num_ligne_source"] = pd.array(
        [None if n is None else int(n) for n in anomalies["num_ligne_source"]], dtype="Int64"
    )
    return anomalies


def _exemple():
    """Une anomalie de chaque cas, avec les sortes attendues (None : non retenue)."""
    return _anomalies(
        _anomalie("format", "anomalie", 3, "d", valeur="masqué", id_anomalie="f1"),
        _anomalie("format", "à arbitrer", 2, "nom", id_anomalie="f2"),
        _anomalie("doublons", "anomalie", 3, "id", id_anomalie="d1", detail="groupe 1"),
        _anomalie("doublons", "anomalie", 4, "id", id_anomalie="d2", detail="groupe 1"),
        _anomalie("doublons", "à arbitrer", 2, "nom", id_anomalie="d3"),
        _anomalie("gestion", "à arbitrer", 5, "d", id_anomalie="g1",
                  proposition="04/01/2025", constat="date hors période"),
        _anomalie("gestion", "à arbitrer", 2, "d, m", id_anomalie="g2"),
        _anomalie("gestion", "à arbitrer", None, "m", id_anomalie="g3",
                  valeur="12,5", cle="2025"),
        _anomalie("gestion", "à arbitrer", 2, "d, autre.colonne", id_anomalie="g4"),
        _anomalie("gestion", "anomalie", 2, "d", id_anomalie="g5"),
        _anomalie("gestion", "information", 2, "d", id_anomalie="g6"),
        _anomalie("format", "anomalie", 99, "d", id_anomalie="f3"),
    )


def _dictionnaire(**natures):
    """Dictionnaire réduit à la nature des colonnes de « t » (texte par défaut)."""
    colonnes = ("id", "d", "m", "nom")
    return {"t": {c: {"nature": {"regle": natures.get(c, "texte")}} for c in colonnes}}


DICTIONNAIRE = _dictionnaire(d="date JJ/MM/AAAA", m="nombre natif")


# --- elements_a_arbitrer -----------------------------------------------------


def test_elements_sortes_ordre_et_valeurs():
    elements = elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION)
    assert list(elements.columns) == list(COLONNES_ELEMENTS)
    assert list(zip(elements["id_anomalie"], elements["sorte"], elements["etat_decision"])) == [
        ("f1", SORTE_CELLULE, ETAT_SANS_DECISION),
        ("d1", SORTE_LIGNE, ETAT_SANS_DECISION),
        ("d2", SORTE_LIGNE, ETAT_SANS_DECISION),
        ("d3", SORTE_LIGNE, ETAT_SANS_DECISION),
        ("g1", SORTE_CELLULE, ETAT_SANS_DECISION),
        ("g2", SORTE_ECART, ETAT_SANS_DECISION),
        ("g3", SORTE_ECART, ETAT_SANS_DECISION),
        ("g4", SORTE_ECART, ETAT_SANS_DECISION),
        ("f3", SORTE_CELLULE, ETAT_INTROUVABLE),
    ]
    par_id = elements.set_index("id_anomalie")
    # La valeur est lue dans la table, même si la table des anomalies la masque.
    assert par_id.at["f1", "valeur"] == "31/02/2025"
    assert par_id.at["g1", "valeur"] == "04/01/2052"
    assert par_id.at["g2", "valeur"] == "01/01/2025 | 10.5"
    # Une cellule vide est affichée vide; un champ qui nomme une table liée garde
    # la valeur de la table des anomalies; un écart sans lignage aussi.
    assert par_id.at["d1", "valeur"] == "A2"
    assert par_id.at["g4", "valeur"] is None
    assert par_id.at["g3", "valeur"] == "12,5"
    assert par_id.at["g1", "proposition"] == "04/01/2025"
    assert par_id.at["d1", "detail"] == "groupe 1"
    assert pd.isna(par_id.at["f3", "empreinte"])
    assert elements["num_ligne_source"].dtype == "Int64"
    # Rien n'est modifié.
    assert _exemple().equals(_exemple())


def test_empreintes():
    tables = _tables()
    elements = elements_a_arbitrer(tables, _exemple(), AUCUNE_DECISION).set_index("id_anomalie")
    assert elements["empreinte"].dropna().str.len().eq(16).all()
    # Même exécution, mêmes empreintes.
    deuxieme = elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION)
    assert list(deuxieme["empreinte"].dropna()) == list(elements["empreinte"].dropna())

    def apres(table, position, colonne, valeur):
        modifiees = _tables()
        modifiees[table].at[position, colonne] = valeur
        return elements_a_arbitrer(modifiees, _exemple(), AUCUNE_DECISION).set_index(
            "id_anomalie"
        )["empreinte"]

    # La cellule : seule sa valeur compte.
    assert apres("t", 1, "d", "28/02/2025")["f1"] != elements.at["f1", "empreinte"]
    assert apres("t", 1, "nom", "autre")["f1"] == elements.at["f1", "empreinte"]
    # La ligne : toute colonne hors lignage compte, le lignage ne compte pas.
    assert apres("t", 1, "nom", "autre")["d1"] != elements.at["d1", "empreinte"]
    avec_chargement = _tables()
    avec_chargement["t"]["date_chargement"] = ["a", "b", "c", "d"]
    assert elements_a_arbitrer(avec_chargement, _exemple(), AUCUNE_DECISION).set_index(
        "id_anomalie"
    ).at["d1", "empreinte"] == elements.at["d1", "empreinte"]
    # Un lignage changé : la ligne n'est plus retrouvée.
    assert apres("t", 1, "feuille_source", "f")["d1"] is None
    # Un texte, un nombre et une date de même écriture n'ont pas la même empreinte.
    assert apres("t", 0, "m", "10.5")["g2"] != elements.at["g2", "empreinte"]
    assert apres("t", 0, "m", 10.5)["g2"] == elements.at["g2", "empreinte"]
    natif = apres("t", 0, "d", datetime(2025, 1, 1))["g2"]
    assert natif != elements.at["g2", "empreinte"]


def test_date_native_affichee_aaaa_mm_jj():
    tables = {"t": _table(id=["A1"], d=[pd.Timestamp(2025, 3, 14)], m=[1.0], nom=["x"])}
    anomalies = _anomalies(_anomalie("format", "anomalie", 2, "d", id_anomalie="f1"))
    elements = elements_a_arbitrer(tables, anomalies, AUCUNE_DECISION)
    assert elements.at[0, "valeur"] == "2025-03-14"


def test_lignage_en_double_introuvable():
    t = _table(id=["A1", "A2"], d=["x", "y"], m=[1.0, 2.0], nom=["a", "b"])
    t["num_ligne_source"] = [2, 2]
    anomalies = _anomalies(_anomalie("format", "anomalie", 2, "d", id_anomalie="f1"))
    elements = elements_a_arbitrer({"t": t}, anomalies, AUCUNE_DECISION)
    assert elements.at[0, "etat_decision"] == ETAT_INTROUVABLE


def test_etat_des_decisions():
    elements = elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION).set_index(
        "id_anomalie"
    )

    def decision(identifiant, sorte, empreinte):
        return {
            "sorte": sorte,
            "empreinte": empreinte,
            "decision": "garder",
        }

    decisions = {
        "meta": {"structure": 1, "imports": []},
        "decisions": {
            "f1": decision("f1", SORTE_CELLULE, elements.at["f1", "empreinte"]),
            "d1": decision("d1", SORTE_LIGNE, "autre"),
            "g1": decision("g1", SORTE_ECART, elements.at["g1", "empreinte"]),
            "inconnu": decision("inconnu", SORTE_CELLULE, "x"),
        },
    }
    etats = elements_a_arbitrer(_tables(), _exemple(), decisions).set_index("id_anomalie")[
        "etat_decision"
    ]
    assert etats["f1"] == ETAT_VALABLE
    assert etats["d1"] == ETAT_CADUQUE
    # Sorte changée : caduque.
    assert etats["g1"] == ETAT_CADUQUE
    assert etats["d2"] == ETAT_SANS_DECISION


def test_table_des_anomalies_inattendue():
    with pytest.raises(ValueError, match="colonnes absentes id_anomalie"):
        elements_a_arbitrer(_tables(), _exemple().drop(columns="id_anomalie"), AUCUNE_DECISION)


def test_coherence_avec_les_controles():
    """Sur des contrôles réels : une cellule par anomalie de format, une ligne par
    doublon, rien pour les écarts à arbitrer du format."""
    tables = {
        "t": _table(id=["A1", "A2", "A2", "A4"], v=["a", "", "b", "c"], q=["1", "2", "3", "4"])
    }
    regle_obligatoire = {
        "regle": "obligatoire", "proposition": "obligatoire", "statut": "valide",
        "motif": "m", "a_arbitrer": {}, "commentaire": "", "revu_le": "2026-10-01 10:00:00",
    }

    def regle(nom, valeur="aucune"):
        return {"regle": valeur, "proposition": valeur, "statut": "observé", "motif": "m",
                "a_arbitrer": {}, "commentaire": "", "revu_le": None,
                **({"liste": []} if nom in ("valeurs", "format") else {})}

    def colonne(obligatoire=False):
        return {
            "obligatoire": regle_obligatoire if obligatoire else regle("obligatoire", "facultatif"),
            "nature": regle("nature", "texte"),
            "valeurs": regle("valeurs"),
            "format": regle("format"),
        }

    dictionnaire = {"t": {"id": colonne(), "v": colonne(True), "q": colonne()}}
    declarations = {"t": {"cle": ["id"]}}
    ecarts, bilan = controler(tables, dictionnaire, {}, {})
    doublons, bilan_doublons = detecter_doublons(tables, declarations, {})
    ecarts_gestion, bilan_gestion = controler_gestion(
        tables, dictionnaire, [], declarations, doublons, {}
    )
    anomalies = reunir_anomalies(
        tables, ecarts, doublons, ecarts_gestion, bilan, bilan_doublons, bilan_gestion,
        declarations, {},
    )
    elements = elements_a_arbitrer(tables, anomalies, AUCUNE_DECISION)
    attendues = int((ecarts["issue"] == "anomalie").sum()) + len(doublons)
    assert attendues == 3
    assert len(elements) == attendues
    assert list(elements["sorte"]) == [SORTE_CELLULE, SORTE_LIGNE, SORTE_LIGNE]
    assert list(elements["id_anomalie"]) == list(anomalies["id_anomalie"])
    assert (elements["etat_decision"] == ETAT_SANS_DECISION).all()


# --- classeur d'arbitrage -----------------------------------------------------


def _remplir(chemin, decisions):
    """Remplit le classeur comme le métier : {id_anomalie: (décision, valeur, commentaire)}."""
    classeur = load_workbook(chemin)
    feuille = classeur[FEUILLE_ARBITRAGE]
    colonnes = {nom: indice for indice, nom in enumerate(COLONNES_ARBITRAGE, start=1)}
    for numero in range(2, feuille.max_row + 1):
        identifiant = feuille.cell(row=numero, column=colonnes["id_anomalie"]).value
        if identifiant in decisions:
            choix, valeur, commentaire = decisions[identifiant]
            feuille.cell(row=numero, column=colonnes["decision"]).value = choix
            if valeur is not ...:
                feuille.cell(row=numero, column=colonnes["valeur_retenue"]).value = valeur
            feuille.cell(row=numero, column=colonnes["commentaire"]).value = commentaire
    classeur.save(chemin)


def test_export_du_classeur(tmp_path):
    chemin = tmp_path / "decisions" / "arbitrage_2026-03.xlsx"
    elements = elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION)
    assert exporter_arbitrage(elements, "2026-03", chemin, nb_imports=0) == chemin
    classeur = load_workbook(chemin)
    assert classeur.sheetnames == ["mode_emploi", FEUILLE_ARBITRAGE]
    feuille = classeur[FEUILLE_ARBITRAGE]
    lignes = list(feuille.iter_rows(values_only=True))
    assert lignes[0] == COLONNES_ARBITRAGE
    # La ligne introuvable n'est pas exportée.
    assert len(lignes) == 1 + 8
    par_id = {ligne[COLONNES_ARBITRAGE.index("id_anomalie")]: dict(zip(COLONNES_ARBITRAGE, ligne))
              for ligne in lignes[1:]}
    assert par_id["f1"]["valeur"] == "31/02/2025"
    assert par_id["f1"]["num_ligne_source"] == 3
    # Valeur retenue préremplie par la proposition, pour une cellule seulement.
    assert par_id["g1"]["valeur_retenue"] == "04/01/2025"
    assert par_id["f1"]["valeur_retenue"] is None
    assert all(ligne["decision"] is None for ligne in par_id.values())
    # Colonnes techniques masquées, cellules écrites en texte, feuille protégée.
    lettre = {nom: feuille.cell(row=1, column=i).column_letter
              for i, nom in enumerate(COLONNES_ARBITRAGE, start=1)}
    assert feuille.column_dimensions[lettre["id_anomalie"]].hidden
    assert feuille.column_dimensions[lettre["empreinte"]].hidden
    assert not feuille.column_dimensions[lettre["sorte"]].hidden
    assert feuille.protection.sheet
    rang = COLONNES_ARBITRAGE.index("valeur_retenue") + 1
    assert feuille.cell(row=2, column=rang).number_format == "@"
    assert not feuille.cell(row=2, column=rang).protection.locked
    assert feuille.cell(row=2, column=1).protection.locked
    # Un menu par sorte.
    menus = {tuple(m.formula1.strip('"').split(",")): str(m.sqref)
             for m in feuille.data_validations.dataValidation}
    assert set(menus) == {("corriger", "garder"), ("retirer", "garder"), ("garder",)}
    # Repère des imports dans le mode d'emploi.
    repere = [ligne for ligne in classeur["mode_emploi"].iter_rows(values_only=True)
              if ligne[0] and ligne[0].startswith("Imports de décisions")]
    assert repere[0][1] == 0


def test_export_refuse_d_ecraser_et_repere_invalide(tmp_path):
    chemin = tmp_path / "a.xlsx"
    elements = elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION)
    exporter_arbitrage(elements, "2026-03", chemin, nb_imports=0)
    with pytest.raises(FileExistsError, match="existe déjà"):
        exporter_arbitrage(elements, "2026-03", chemin, nb_imports=0)
    exporter_arbitrage(elements, "2026-03", chemin, nb_imports=2, ecraser=True)
    for mauvais in (-1, True, "1"):
        with pytest.raises(ValueError, match="nb_imports"):
            exporter_arbitrage(elements, "2026-03", tmp_path / "b.xlsx", nb_imports=mauvais)


def test_proposition_masquee_non_preremplie(tmp_path):
    anomalies = _anomalies(
        _anomalie("gestion", "à arbitrer", 2, "d", id_anomalie="g1", proposition="masqué")
    )
    elements = elements_a_arbitrer(_tables(), anomalies, AUCUNE_DECISION)
    chemin = exporter_arbitrage(elements, "2026-03", tmp_path / "a.xlsx", nb_imports=0)
    ligne = list(load_workbook(chemin)[FEUILLE_ARBITRAGE].iter_rows(values_only=True))[1]
    assert dict(zip(COLONNES_ARBITRAGE, ligne))["valeur_retenue"] is None


def test_aller_retour_complet(tmp_path, monkeypatch):
    instants = iter([datetime(2026, 10, 20, 10, 0, s) for s in range(0, 60, 5)])
    monkeypatch.setattr(module_dictionnaire, "_maintenant", lambda: next(instants))
    fichier = tmp_path / "config" / "decisions.yaml"
    chemin = tmp_path / "a.xlsx"
    elements = elements_a_arbitrer(_tables(), _exemple(), charger_decisions(fichier))
    exporter_arbitrage(elements, "2026-03", chemin, nb_imports=0)
    _remplir(chemin, {
        "f1": ("corriger", "28.02.2025", "date saisie à la main"),
        "g1": ("corriger", ..., None),
        "d1": ("retirer", None, None),
        "d2": ("garder", None, None),
        "g2": ("garder", None, "vu"),
    })
    arbitrage = importer_arbitrage(chemin, DICTIONNAIRE, nb_imports=0)
    assert arbitrage.avertissements == []
    assert arbitrage.repere == 0
    decisions = arbitrage.decisions
    assert list(decisions.columns) == list(COLONNES_DECISIONS_LUES)
    assert list(decisions["id_anomalie"]) == ["f1", "d1", "d2", "g1", "g2"]
    par_id = decisions.set_index("id_anomalie")
    # Date lue dans une forme élargie, écrite dans la forme de la nature.
    assert par_id.at["f1", "valeur_retenue"] == "28/02/2025"
    assert par_id.at["f1", "nature"] == "date JJ/MM/AAAA"
    assert par_id.at["g1", "valeur_retenue"] == "04/01/2025"
    assert par_id.at["d1", "valeur_retenue"] is None
    assert par_id.at["d1", "nature"] is None
    assert par_id.at["f1", "num_ligne_source"] == 3
    assert par_id.at["g2", "commentaire"] == "vu"

    ecrit, entree = ecrire_decisions(arbitrage, "2026-03", fichier)
    assert ecrit == fichier
    assert entree == {
        "date": "2026-10-20 10:00:00", "classeur": "a.xlsx", "periode_fin": "2026-03",
        "decisions": 5, "remplacees": 0, "corriger": 2, "garder": 2, "retirer": 1,
    }
    enregistrees = charger_decisions(fichier)
    assert enregistrees["meta"]["imports"] == [entree]
    assert enregistrees["decisions"]["f1"] == {
        "sorte": "cellule", "table": "t", "fichier_source": "t.csv", "feuille_source": None,
        "num_ligne_source": 3, "champ": "d", "constat": "constat",
        "empreinte": elements.set_index("id_anomalie").at["f1", "empreinte"],
        "decision": "corriger", "valeur_retenue": "28/02/2025", "nature": "date JJ/MM/AAAA",
        "commentaire": "date saisie à la main", "decide_le": "2026-10-20 10:00:00",
        "classeur": "a.xlsx",
    }
    # Premier import : pas de copie de sauvegarde, rien à copier.
    assert not (fichier.parent / "sauvegardes").exists()
    # Le même classeur ne s'importe plus.
    assert repere_arbitrage(chemin) == 0
    with pytest.raises(ValueError, match="périmé"):
        importer_arbitrage(chemin, DICTIONNAIRE, nb_imports=1)
    with pytest.raises(ValueError, match="périmé"):
        ecrire_decisions(arbitrage, "2026-03", fichier)

    # Export suivant : seuls les éléments sans décision valable.
    elements = elements_a_arbitrer(_tables(), _exemple(), charger_decisions(fichier))
    a_arbitrer = elements[elements["etat_decision"] == ETAT_SANS_DECISION]
    assert list(a_arbitrer["id_anomalie"]) == ["d3", "g3", "g4"]
    # Une source modifiée rend la décision caduque.
    modifiees = _tables()
    modifiees["t"].at[1, "d"] = "30/02/2025"
    elements = elements_a_arbitrer(modifiees, _exemple(), charger_decisions(fichier))
    etats = elements.set_index("id_anomalie")["etat_decision"]
    assert etats["f1"] == ETAT_CADUQUE
    assert etats["g1"] == ETAT_VALABLE
    chemin2 = tmp_path / "b.xlsx"
    exporter_arbitrage(elements, "2026-03", chemin2, nb_imports=1)
    feuille = load_workbook(chemin2)[FEUILLE_ARBITRAGE]
    exportees = [
        ligne[COLONNES_ARBITRAGE.index("id_anomalie")]
        for ligne in feuille.iter_rows(min_row=2, values_only=True)
    ]
    # La ligne 3 a changé : la décision sur sa cellule (f1) et sur la ligne entière
    # (d1) sont caduques. Décisions valables et ligne introuvable : non exportées.
    assert etats["d1"] == ETAT_CADUQUE
    assert exportees == ["f1", "d1", "d3", "g3", "g4"]
    assert repere_arbitrage(chemin2) == 1
    _remplir(chemin2, {"f1": ("garder", None, None), "g3": ("garder", None, None)})
    arbitrage = importer_arbitrage(chemin2, DICTIONNAIRE, nb_imports=1)
    _, entree = ecrire_decisions(arbitrage, "2026-03", fichier)
    assert entree["remplacees"] == 1 and entree["decisions"] == 2
    enregistrees = charger_decisions(fichier)
    assert enregistrees["decisions"]["f1"]["decision"] == "garder"
    assert enregistrees["decisions"]["f1"]["valeur_retenue"] is None
    assert len(enregistrees["meta"]["imports"]) == 2
    copies = list((fichier.parent / "sauvegardes").iterdir())
    assert [c.name for c in copies] == ["decisions_2026-03_2026-10-20_100010.yaml"]
    assert not list(fichier.parent.glob("*.tmp"))


def test_import_sans_decision_rien_n_est_ecrit(tmp_path):
    fichier = tmp_path / "decisions.yaml"
    chemin = tmp_path / "a.xlsx"
    exporter_arbitrage(
        elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION), "2026-03", chemin,
        nb_imports=0,
    )
    _remplir(chemin, {"f1": (None, None, "à voir")})
    arbitrage = importer_arbitrage(chemin, DICTIONNAIRE, nb_imports=0)
    assert arbitrage.decisions.empty
    assert arbitrage.avertissements == ["ligne 2 : commentaire sans décision, ignoré"]
    assert ecrire_decisions(arbitrage, "2026-03", fichier) == (fichier, None)
    assert not fichier.exists()


def test_import_problemes_listes_en_une_fois(tmp_path):
    chemin = tmp_path / "a.xlsx"
    exporter_arbitrage(
        elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION), "2026-03", chemin,
        nb_imports=0,
    )
    # Lignes Excel : f1 2, d1 3, d2 4, d3 5, g1 6, g2 7, g3 8, g4 9.
    _remplir(chemin, {
        "f1": ("corriger", "31/02/2025", None),
        "d1": ("corriger", "x", None),
        "d2": ("supprimer", None, None),
        "d3": ("retirer", None, None),
        "g1": ("corriger", None, None),
        "g2": ("retirer", None, None),
        "g3": ("garder", None, None),
    })
    with pytest.raises(ValueError) as erreur:
        importer_arbitrage(chemin, DICTIONNAIRE, nb_imports=0)
    message = str(erreur.value)
    assert "5 problèmes" in message
    assert "ligne 2 : valeur retenue illisible dans la nature « date JJ/MM/AAAA » de t.d " \
        "(date impossible)" in message
    assert "ligne 3 : décision corriger impossible pour un élément de sorte ligne " \
        "(possibles : retirer, garder)" in message
    assert "ligne 4 : décision 'supprimer' hors du menu (corriger, garder, retirer)" in message
    assert "ligne 6 : décision corriger sans valeur retenue" in message
    assert "ligne 7 : décision retirer impossible pour un élément de sorte écart" in message


def test_import_cellule_contradictoire_et_colonne_inconnue(tmp_path):
    anomalies = _anomalies(
        _anomalie("format", "anomalie", 3, "d", id_anomalie="f1"),
        _anomalie("gestion", "à arbitrer", 3, "d", id_anomalie="g1"),
        _anomalie("format", "anomalie", 2, "nom", id_anomalie="f2"),
    )
    chemin = tmp_path / "a.xlsx"
    exporter_arbitrage(elements_a_arbitrer(_tables(), anomalies, AUCUNE_DECISION), "2026-03",
                       chemin, nb_imports=0)
    _remplir(chemin, {
        "f1": ("corriger", "01/03/2025", None),
        "g1": ("corriger", "02/03/2025", None),
        "f2": ("corriger", "Nom", None),
    })
    dictionnaire = {"t": {"d": {"nature": {"regle": "date JJ/MM/AAAA"}}}}
    with pytest.raises(ValueError) as erreur:
        importer_arbitrage(chemin, dictionnaire, nb_imports=0)
    message = str(erreur.value)
    assert "ligne 3 : la même cellule reçoit une autre valeur ligne 2" in message
    assert "ligne 4 : colonne t.nom inconnue du dictionnaire" in message
    # Deux fois la même valeur : accepté.
    _remplir(chemin, {"g1": ("corriger", "1.3.2025", None), "f1": ("corriger", "01/03/2025", None),
                      "f2": ("garder", None, None)})
    with pytest.raises(ValueError, match="1 problèmes"):
        importer_arbitrage(chemin, dictionnaire, nb_imports=0)


def test_import_avertissement_valeur_ignoree(tmp_path):
    chemin = tmp_path / "a.xlsx"
    exporter_arbitrage(
        elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION), "2026-03", chemin,
        nb_imports=0,
    )
    # g1 garde sa proposition préremplie : pas d'avertissement; d1 porte une valeur.
    _remplir(chemin, {"g1": ("garder", ..., None), "d1": ("retirer", "x", None)})
    arbitrage = importer_arbitrage(chemin, DICTIONNAIRE, nb_imports=0)
    assert arbitrage.avertissements == ["ligne 3 : valeur retenue ignorée avec la décision retirer"]
    assert list(arbitrage.decisions["valeur_retenue"]) == [None, None]


def test_import_classeur_modifie(tmp_path):
    chemin = tmp_path / "a.xlsx"
    exporter_arbitrage(
        elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION), "2026-03", chemin,
        nb_imports=0,
    )
    classeur = load_workbook(chemin)
    classeur[FEUILLE_ARBITRAGE].cell(row=1, column=1).value = "tableau"
    classeur.save(chemin)
    with pytest.raises(ValueError, match="en-tête modifié"):
        importer_arbitrage(chemin, DICTIONNAIRE, nb_imports=0)
    classeur = load_workbook(chemin)
    del classeur[FEUILLE_ARBITRAGE]
    classeur.save(chemin)
    with pytest.raises(ValueError, match="feuille « arbitrage » absente"):
        importer_arbitrage(chemin, DICTIONNAIRE, nb_imports=0)


# --- valeur retenue ---------------------------------------------------------


@pytest.mark.parametrize(
    "valeur, nature, attendu",
    [
        ("14.03.2025", "date JJ/MM/AAAA", ("14/03/2025", None)),
        (" 14/03/2025 ", "date JJ/MM/AAAA", ("14/03/2025", None)),
        ("2025-03-14", "date AAAAMMJJ", ("20250314", None)),
        ("14/03/2025", "date native", (datetime(2025, 3, 14), None)),
        (
            "14/03/25",
            "date JJ/MM/AAAA",
            (None, "année sur deux chiffres, date de coupure non fournie"),
        ),
        ("31/02/2025", "date native", (None, "date impossible")),
        ("07:30", "heure HH:MM", ("07:30", None)),
        ("7:30", "heure HH:MM", (None, "forme non reconnue")),
        ("12,50 €", "nombre natif", (12.5, None)),
        ("12", "nombre natif", (12.0, None)),
        (12, "nombre natif", (12.0, None)),
        ("076", "entier (texte)", ("076", None)),
        ("1 250,5", "décimal virgule (texte)", ("1250,5", None)),
        ("12.5", "décimal virgule (texte)", ("12,5", None)),
        ("1.250", "nombre natif", (None, "séparateur ambigu")),
        ("1h30", "entier (texte)", (None, "valeur avec une unité")),
        ("12,5", "entier (texte)", (None, "partie décimale non nulle")),
        ("  Nom corrigé ", "texte", ("Nom corrigé", None)),
        ("x", "à décider", ("x", None)),
        ("   ", "texte", (None, "valeur retenue vide")),
    ],
)
def test_lire_valeur_retenue(valeur, nature, attendu):
    assert lire_valeur_retenue(valeur, nature) == attendu


# --- fichier des décisions ------------------------------------------------------


def test_charger_decisions_absent_et_invalide(tmp_path):
    assert charger_decisions(tmp_path / "absent.yaml") == AUCUNE_DECISION
    fichier = tmp_path / "decisions.yaml"
    fichier.write_text(
        "meta:\n  structure: 1\n  imports: []\n"
        "decisions:\n  a1:\n    sorte: ligne\n    decision: corriger\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="décision a1 : champs attendus"):
        charger_decisions(fichier)
    fichier.write_text("meta: {}\ndecisions: []\n", encoding="utf-8")
    with pytest.raises(ValueError) as erreur:
        charger_decisions(fichier)
    assert "meta : structure 1 attendue" in str(erreur.value)
    assert "decisions : un dictionnaire par identifiant est attendu" in str(erreur.value)


def test_problemes_d_une_decision(tmp_path):
    from frigeo.decisions import problemes_decisions

    base = {
        "sorte": "ligne", "table": "t", "fichier_source": "t.csv", "feuille_source": None,
        "num_ligne_source": 2, "champ": "id", "constat": "c", "empreinte": "e",
        "decision": "retirer", "valeur_retenue": None, "nature": None, "commentaire": None,
        "decide_le": "2026-10-20 10:00:00", "classeur": "a.xlsx",
    }
    assert problemes_decisions({"meta": {"structure": 1, "imports": []},
                                "decisions": {"a": base}}) == []
    mauvaises = {
        "a": {**base, "decision": "corriger"},
        "b": {**base, "sorte": "colonne"},
        "c": {**base, "valeur_retenue": "x"},
        "d": {**base, "sorte": "cellule", "decision": "corriger"},
        "e": {**base, "empreinte": "", "decide_le": "20/10/2026"},
    }
    entree = {"date": "2026-10-20 10:00:00", "classeur": "a", "periode_fin": "2026-03",
              "decisions": 1, "remplacees": 0, "corriger": 0, "garder": 0, "retirer": 1}
    problemes = problemes_decisions(
        {"meta": {"structure": 1, "imports": [entree, entree]}, "decisions": mauvaises}
    )
    assert problemes == [
        "meta, import 2 : date antérieure ou égale à celle de l'import précédent",
        "décision a : décision 'corriger' impossible pour la sorte ligne",
        "décision a : valeur retenue absente ou d'un type non admis",
        "décision b : sorte 'colonne' inconnue",
        "décision c : valeur retenue donnée sans la décision corriger",
        "décision d : valeur retenue absente ou d'un type non admis",
        "décision e : empreinte absente",
        "décision e : decide_le 'AAAA-MM-JJ HH:MM:SS' attendu",
    ]


def test_ecriture_dans_la_meme_seconde_refusee(tmp_path, monkeypatch):
    monkeypatch.setattr(module_dictionnaire, "_maintenant", lambda: datetime(2026, 10, 20, 10))
    fichier = tmp_path / "decisions.yaml"
    chemins = [tmp_path / "a.xlsx", tmp_path / "b.xlsx"]
    exporter_arbitrage(elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION),
                       "2026-03", chemins[0], nb_imports=0)
    _remplir(chemins[0], {"d1": ("garder", None, None)})
    arbitrage = importer_arbitrage(chemins[0], DICTIONNAIRE, nb_imports=0)
    ecrire_decisions(arbitrage, "2026-03", fichier)
    exporter_arbitrage(elements_a_arbitrer(_tables(), _exemple(), charger_decisions(fichier)),
                       "2026-03", chemins[1], nb_imports=1)
    _remplir(chemins[1], {"d2": ("garder", None, None)})
    arbitrage = importer_arbitrage(chemins[1], DICTIONNAIRE, nb_imports=1)
    avant = fichier.read_bytes()
    with pytest.raises(ValueError, match="doit être postérieur"):
        ecrire_decisions(arbitrage, "2026-03", fichier)
    assert fichier.read_bytes() == avant


def test_date_et_nombre_natifs_enregistres(tmp_path):
    fichier = tmp_path / "decisions.yaml"
    chemin = tmp_path / "a.xlsx"
    exporter_arbitrage(
        elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION), "2026-03", chemin,
        nb_imports=0,
    )
    _remplir(chemin, {"f1": ("corriger", "28/02/2025", None)})
    dictionnaire = _dictionnaire(d="date native")
    arbitrage = importer_arbitrage(chemin, dictionnaire, nb_imports=0)
    ecrire_decisions(arbitrage, "2026-03", fichier)
    lue = charger_decisions(fichier)["decisions"]["f1"]
    assert lue["valeur_retenue"] == datetime(2025, 2, 28)
    assert type(lue["valeur_retenue"]) is datetime
    assert lue["nature"] == "date native"


def test_ajouter_decisions():
    elements = elements_a_arbitrer(_tables(), _exemple(), AUCUNE_DECISION).set_index("id_anomalie")

    def decision(choix, empreinte, sorte=SORTE_CELLULE):
        return {"sorte": sorte, "empreinte": empreinte, "decision": choix,
                "decide_le": "2026-10-20 10:00:00"}

    decisions = {
        "meta": {"structure": 1, "imports": []},
        "decisions": {
            "f1": decision("garder", elements.at["f1", "empreinte"]),
            "d1": decision("retirer", "autre", SORTE_LIGNE),
            "g5": decision("garder", "x"),
        },
    }
    anomalies = _exemple()
    elements = elements_a_arbitrer(_tables(), anomalies, decisions)
    marquees = ajouter_decisions(anomalies, elements, decisions)
    assert list(marquees.columns) == [*COLONNES_ANOMALIES, COLONNE_DECISION_METIER]
    par_id = marquees.set_index("id_anomalie")[COLONNE_DECISION_METIER]
    assert par_id["f1"] == "garder le 2026-10-20"
    assert par_id["d1"] == "retirer le 2026-10-20, caduque"
    # Une décision dont l'anomalie n'est pas un élément à arbitrer n'est pas reprise.
    assert par_id["g5"] is None
    assert par_id.drop(["f1", "d1"]).isna().all()
    assert list(anomalies.columns) == list(COLONNES_ANOMALIES)
