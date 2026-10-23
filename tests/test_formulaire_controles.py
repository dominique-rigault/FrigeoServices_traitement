"""Tests du formulaire des règles de contrôle (frigeo.formulaire_controles)."""

import pandas as pd
import pytest
import yaml
from openpyxl import load_workbook

from frigeo.controles import lire_fichier_controles
from frigeo.formulaire_controles import (
    COLONNES_BORNES,
    COLONNES_CHANGEMENTS,
    COLONNES_DECIMALES,
    COLONNES_DOUBLONS,
    COLONNES_LUHN,
    COLONNES_MOTIFS,
    COLONNES_TABLES,
    ENTETE_CONTROLES,
    LIGNES_AJOUT,
    REPERE_EMPREINTE,
    appliquer_formulaire,
    archiver_formulaire,
    dossier_archives,
    ecrire_controles,
    empreinte_controles,
    exporter_formulaire,
    importer_formulaire,
    lignes_bornes,
    lignes_decimales,
    lignes_doublons,
    lignes_luhn,
    lignes_motifs,
    lignes_tables,
    observer_groupe,
    observer_luhn,
    observer_nombres,
    texte_controles,
    texte_nombre,
)

SIRET_VALIDES = ["73282932000074", "35600000000048", "44306184100047"]


def _nature(nature):
    return {"nature": {"regle": nature}}


def _jeu():
    """Deux tables fictives, leur dictionnaire (natures en vigueur) et le contenu de
    config/controles.yaml."""
    tables = {
        "factures": pd.DataFrame(
            {
                "num": ["F-1", "F-2", "F-3", "F-4"],
                "montant": [12.5, 181.45000000000002, 3, "7,25"],
                "remise": [0, 5, 15, None],
                "code": ["A1", "A2", "A3", "A4"],
                "siret": SIRET_VALIDES + ["12345678901234"],
                "salaire": [1800.5, 2500, 3100, 2000],
            }
        ),
        "parametres": pd.DataFrame({"annee": [2025], "taux": [0.42], "libelle": ["unique"]}),
    }
    dictionnaire = {
        "factures": {
            "num": _nature("texte"),
            "montant": _nature("nombre natif"),
            "remise": _nature("nombre natif"),
            "code": _nature("texte"),
            "siret": _nature("entier (texte)"),
            "salaire": _nature("nombre natif"),
        },
        "parametres": {
            "annee": _nature("nombre natif"),
            "taux": _nature("nombre natif"),
            "libelle": _nature("texte"),
        },
    }
    contenu = {
        "regles": {
            "factures": {
                "montant": {"decimales": 2},
                "remise": {"bornes": {"min": 0, "max": 15, "max_exclu": True}},
                "code": {"motif": "[A-Z][0-9]+"},
            },
        },
        "doublons": {"factures": {"cle": ["num"], "identique_hors_cle": False}},
        "ecartees": {"decimales": {"factures": ["salaire"]}, "doublons": {"factures": [["siret"]]}},
        "commentaires": {
            "bornes": {"factures": {"remise": "remise commerciale plafonnée"}},
            "doublons": {"factures": {"num": "numéro de facture"}},
        },
    }
    sensibles = {"factures": frozenset({"salaire"})}
    return tables, dictionnaire, contenu, sensibles


def _par_colonne(lignes):
    return {(ligne["table"], ligne["colonne"]): ligne for ligne in lignes}


def test_observations_des_nombres():
    serie = pd.Series([12.5, 181.45000000000002, 3, "7,25", "1 200", "abc", None, ""])
    assert observer_nombres(serie) == {"valeurs": 4, "min": 3, "max": 181.45000000000002,
                                       "decimales": 2}
    assert observer_nombres(pd.Series(["abc", None])) == {
        "valeurs": 0, "min": None, "max": None, "decimales": None,
    }
    assert texte_nombre(181.45000000000002) == "181,45"
    assert (texte_nombre(3), texte_nombre(0.5), texte_nombre(2.0)) == ("3", "0,5", "2")


def test_observations_des_cles_de_luhn_et_des_groupes():
    observation = observer_luhn(pd.Series(SIRET_VALIDES + ["12345678901234", "12 34", None]))
    assert observation == {"valeurs": 4, "part_valide": 0.75}
    assert observer_luhn(pd.Series(["abc"]))["part_valide"] is None
    donnees = pd.DataFrame({"a": ["x", "y", "x", None], "b": [1, 2, 1, 3]})
    assert observer_groupe(donnees, ["a"]) == {"lignes": 4, "renseignees": 3, "repetees": 2}
    assert observer_groupe(donnees, ["b"]) == {"lignes": 4, "renseignees": 4, "repetees": 2}
    assert observer_groupe(donnees, ["a", "b"])["repetees"] == 2


def test_lignes_des_decimales():
    lignes = _par_colonne(lignes_decimales(*_jeu()[:3]))
    # Une ligne par colonne numérique, codes compris, dans l'ordre du dictionnaire,
    # sauf les entiers écrits en texte (siret), dont la nature interdit les décimales.
    assert list(lignes) == [
        ("factures", "montant"), ("factures", "remise"), ("factures", "salaire"),
        ("parametres", "annee"), ("parametres", "taux"),
    ]
    montant = lignes[("factures", "montant")]
    assert (montant["valeurs"], montant["observé"], montant["proposition"]) == (4, 2, 2)
    # Règle en vigueur : statut documenté, valeur dans la colonne à saisir.
    assert (montant["en_vigueur"], montant["statut"], montant["décimales"]) == (2, "documenté", 2)
    # Proposition écartée : statut invalide; sinon observé.
    assert lignes[("factures", "salaire")]["statut"] == "invalide"
    assert lignes[("factures", "remise")]["statut"] == "observé"
    assert lignes[("factures", "remise")]["décimales"] is None
    # Des entiers seuls n'ont aucune décimale.
    assert lignes[("parametres", "annee")]["proposition"] == 0


def test_lignes_des_bornes():
    tables, dictionnaire, contenu, sensibles = _jeu()
    lignes = _par_colonne(lignes_bornes(tables, dictionnaire, contenu, sensibles))
    remise = lignes[("factures", "remise")]
    assert (remise["observé"], remise["proposition"]) == ("de 0 à 15", "0 à 15")
    assert (remise["proposition_min"], remise["proposition_max"]) == (0, 15)
    assert (remise["en_vigueur"], remise["statut"]) == ("0 à 15 exclu", "documenté")
    # Une borne incluse laisse sa cellule vide, une borne exclue porte « non ».
    assert (remise["min"], remise["min_inclus"], remise["max"], remise["max_inclus"]) == (
        0, None, 15, "non",
    )
    assert remise["commentaire"] == "remise commerciale plafonnée"
    montant = lignes[("factures", "montant")]
    assert (montant["proposition"], montant["statut"], montant["min"]) == (
        "3 à 181,45", "observé", None,
    )
    # Colonne sensible : minimum et maximum masqués, aucune proposition.
    salaire = lignes[("factures", "salaire")]
    assert (salaire["observé"], salaire["proposition"]) == ("masqué", "aucune")
    assert (salaire["proposition_min"], salaire["proposition_max"]) == (None, None)
    # Règle en vigueur écrite avec une virgule décimale.
    contenu["regles"]["parametres"] = {
        "taux": {"bornes": {"min": 0.1, "max": 1, "max_exclu": True}}
    }
    taux = _par_colonne(lignes_bornes(tables, dictionnaire, contenu, sensibles))[
        ("parametres", "taux")
    ]
    assert (taux["en_vigueur"], taux["min"], taux["max_inclus"]) == ("0,1 à 1 exclu", 0.1, "non")
    assert "1800" not in repr(salaire) and "3100" not in repr(salaire)


def test_lignes_des_cles_de_luhn():
    lignes = lignes_luhn(*_jeu()[:3])
    assert len(lignes) == 1
    siret = lignes[0]
    assert (siret["colonne"], siret["valeurs"], siret["observé"], siret["proposition"]) == (
        "siret", 4, "75 % de clés valides", "aucune",
    )
    tables, dictionnaire, contenu, _ = _jeu()
    tables["factures"]["siret"] = SIRET_VALIDES + [SIRET_VALIDES[0]]
    contenu["regles"]["factures"]["siret"] = {"cle_luhn": True}
    siret = lignes_luhn(tables, dictionnaire, contenu)[0]
    assert (siret["proposition"], siret["en_vigueur"], siret["statut"]) == (
        "oui", "oui", "documenté",
    )


def test_lignes_des_doublons():
    lignes = lignes_doublons(*_jeu()[:3])
    resume = [(l["table"], l["type"], l["colonnes"], l["origine"], l["statut"]) for l in lignes]
    # En vigueur, puis clés candidates (texte ou entier, toujours renseignées, sans
    # répétition, pas déjà déclarées seules), puis groupes écartés. Une table d'une
    # seule ligne n'a pas de candidate; un montant ou un salaire non plus.
    assert resume == [
        ("factures", "clé", "num", "en vigueur", "documenté"),
        ("factures", "unique", "code", "proposée", "observé"),
        ("factures", "unique", "siret", "écartée", "invalide"),
    ]
    assert lignes[0]["observé"] == "4 lignes renseignées sur 4, aucune répétition"
    assert lignes[0]["commentaire"] == "numéro de facture"
    # Sans clé déclarée, la première candidate est proposée comme clé.
    tables, dictionnaire, contenu, _ = _jeu()
    del contenu["doublons"]
    resume = [(l["type"], l["colonnes"]) for l in lignes_doublons(tables, dictionnaire, contenu)]
    assert resume == [("clé", "num"), ("clé", "code"), ("unique", "siret")]


def test_lignes_des_tables_et_des_motifs():
    tables, dictionnaire, contenu, _ = _jeu()
    assert lignes_tables(tables, dictionnaire, contenu) == [
        {"table": "factures", "lignes": 4, "clé_en_vigueur": "num", "identiques_hors_clé": "non"},
        {"table": "parametres", "lignes": 1, "clé_en_vigueur": None,
         "identiques_hors_clé": "oui"},
    ]
    assert lignes_motifs(contenu) == [
        {"table": "factures", "colonne": "code", "motif": "[A-Z][0-9]+"}
    ]


def test_empreinte_du_fichier_des_controles(tmp_path):
    chemin = tmp_path / "controles.yaml"
    assert empreinte_controles(chemin) == "absent"
    chemin.write_bytes(b"regles:\n  t: {}\n")
    empreinte = empreinte_controles(chemin)
    chemin.write_bytes(b"regles:\r\n  t: {}\r\n")
    assert empreinte_controles(chemin) == empreinte
    chemin.write_bytes(b"regles:\n  u: {}\n")
    assert empreinte_controles(chemin) != empreinte


def _exporter(tmp_path, nom="formulaire.xlsx"):
    tables, dictionnaire, contenu, sensibles = _jeu()
    return exporter_formulaire(
        tables, dictionnaire, contenu, sensibles, "2026-03", tmp_path / "f" / nom,
        empreinte="abc123",
    )


def _lire(feuille):
    lignes = list(feuille.iter_rows(values_only=True))
    return [dict(zip(lignes[0], l)) for l in lignes[1:] if any(v is not None for v in l)]


def test_formulaire_exporte(tmp_path):
    classeur = load_workbook(_exporter(tmp_path))
    assert classeur.sheetnames == [
        "mode_emploi", "décimales", "bornes", "clé_luhn", "doublons", "tables", "motifs",
        "listes",
    ]
    for nom, colonnes in (
        ("décimales", COLONNES_DECIMALES), ("bornes", COLONNES_BORNES),
        ("clé_luhn", COLONNES_LUHN), ("doublons", COLONNES_DOUBLONS),
        ("tables", COLONNES_TABLES), ("motifs", COLONNES_MOTIFS),
    ):
        assert tuple(c.value for c in classeur[nom][1]) == colonnes
        assert classeur[nom].protection.sheet
    assert classeur["listes"].sheet_state == "hidden"
    assert [c.value for c in classeur["listes"]["A"]] == ["factures", "parametres"]
    assert len(_lire(classeur["décimales"])) == 5
    assert len(_lire(classeur["doublons"])) == 3
    # Les nombres restent des nombres, les colonnes techniques sont masquées.
    bornes = classeur["bornes"]
    remise = next(l for l in _lire(bornes) if l["colonne"] == "remise")
    assert (remise["min"], remise["max"], remise["proposition_max"]) == (0, 15, 15)
    masquees = {
        bornes.cell(row=1, column=i).value
        for i in range(1, len(COLONNES_BORNES) + 1)
        if bornes.column_dimensions[bornes.cell(row=1, column=i).column_letter].hidden
    }
    assert masquees == {"proposition_min", "proposition_max"}
    # Repère de l'empreinte et décomptes du mode d'emploi.
    notice = classeur["mode_emploi"]
    repere = next(r for r in notice.iter_rows(values_only=True) if r[0] == REPERE_EMPREINTE)
    assert repere[1] == "abc123"
    texte = " ".join(str(c.value) for c in notice["A"] if c.value)
    assert "Feuille « décimales » : 3 lignes à « observé » sur 5" in texte
    assert "Feuille « doublons » : 1 sur 3" in texte


def test_cellules_modifiables_et_menus(tmp_path):
    classeur = load_workbook(_exporter(tmp_path))

    def libres(feuille, numero):
        return {
            c.value for c, d in zip(feuille[1], feuille[numero]) if not d.protection.locked
        }

    assert libres(classeur["décimales"], 2) == {"statut", "décimales", "commentaire"}
    assert libres(classeur["bornes"], 2) == {
        "statut", "min", "min_inclus", "max", "max_inclus", "commentaire",
    }
    # Nombres saisis alignés à droite.
    entete = [c.value for c in classeur["bornes"][1]]
    for nom in ("min", "max"):
        cellule = classeur["bornes"].cell(row=2, column=entete.index(nom) + 1)
        assert cellule.alignment.horizontal == "right"
    entete = [c.value for c in classeur["décimales"][1]]
    cellule = classeur["décimales"].cell(row=2, column=entete.index("décimales") + 1)
    assert cellule.alignment.horizontal == "right"
    assert libres(classeur["clé_luhn"], 2) == {"statut", "commentaire"}
    assert libres(classeur["doublons"], 2) == {"type", "statut", "commentaire"}
    # Lignes d'ajout en bas de la feuille des doublons.
    doublons = classeur["doublons"]
    assert libres(doublons, 5) == {"table", "type", "colonnes", "statut", "commentaire"}
    assert libres(doublons, 4 + LIGNES_AJOUT) == libres(doublons, 5)
    assert libres(doublons, 5 + LIGNES_AJOUT) == set()
    assert libres(classeur["tables"], 2) == {"identiques_hors_clé"}
    assert libres(classeur["motifs"], 2) == set()
    menus = {m.formula1 for m in doublons.data_validations.dataValidation}
    assert menus == {
        '"observé,valide,invalide,documenté"', '"clé,unique,présumé unique"',
        "'listes'!$A$1:$A$2",
    }
    menus = [m.formula1 for m in classeur["bornes"].data_validations.dataValidation]
    assert '"oui,non"' in menus
    for feuille in classeur.worksheets:
        assert all(m.showErrorMessage for m in feuille.data_validations.dataValidation)


def test_export_refuse_d_ecraser_un_formulaire(tmp_path):
    chemin = _exporter(tmp_path)
    with pytest.raises(FileExistsError, match="existe déjà"):
        _exporter(tmp_path)
    tables, dictionnaire, contenu, sensibles = _jeu()
    exporter_formulaire(
        tables, dictionnaire, contenu, sensibles, "2026-03", chemin, empreinte="x",
        ecraser=True,
    )


def test_aucune_valeur_d_une_colonne_sensible_dans_le_formulaire(tmp_path):
    classeur = load_workbook(_exporter(tmp_path))
    valeurs = {
        str(v) for feuille in classeur.worksheets for ligne in feuille.iter_rows(values_only=True)
        for v in ligne if v is not None
    }
    assert not valeurs & {"1800.5", "2500", "3100", "2000", "1800,5"}


def test_cles_du_formulaire_admises_dans_le_fichier_des_controles(tmp_path):
    from frigeo.controles import charger_regles_complementaires, lire_fichier_controles
    from frigeo.doublons import charger_declarations_doublons

    chemin = tmp_path / "controles.yaml"
    chemin.write_text(
        "regles:\n  t:\n    a: {decimales: 2}\n"
        "doublons:\n  t:\n    cle: [a]\n"
        "ecartees:\n  bornes:\n    t: [a]\n"
        "commentaires:\n  decimales:\n    t:\n      a: montant\n",
        encoding="utf-8",
    )
    assert set(lire_fichier_controles(chemin)) == {"regles", "doublons", "ecartees", "commentaires"}
    # Sans effet sur les contrôles.
    assert charger_regles_complementaires(chemin) == {"t": {"a": {"decimales": 2}}}
    assert charger_declarations_doublons(chemin) == {"t": {"cle": ["a"]}}
    chemin.write_text("regle:\n  t: {}\n", encoding="utf-8")
    with pytest.raises(ValueError) as erreur:
        lire_fichier_controles(chemin)
    assert (
        "les seules clés attendues sont « regles », « doublons », « ecartees » et "
        "« commentaires » (clé inconnue\u00a0: regle)"
    ) in str(erreur.value)


# --- Import du formulaire et écriture de config/controles.yaml (morceau c) ---------


def _modifier(chemin, feuille, cle, **valeurs):
    """Modifie une ligne d'une feuille repérée par ses premières cellules (`cle`),
    comme le ferait le métier dans Excel; avec `cle=None`, écrit la première ligne
    vide (ligne d'ajout)."""
    classeur = load_workbook(chemin)
    page = classeur[feuille]
    entete = [c.value for c in page[1]]
    for numero in range(2, page.max_row + 2):
        cellules = [page.cell(row=numero, column=i + 1).value for i in range(len(entete))]
        if cle is None and all(v is None for v in cellules) or (
            cle is not None and tuple(cellules[: len(cle)]) == tuple(cle)
        ):
            break
    else:
        raise AssertionError(f"ligne {cle} introuvable dans {feuille}")
    for nom, valeur in valeurs.items():
        page.cell(row=numero, column=entete.index(nom) + 1).value = valeur
    classeur.save(chemin)
    return numero


def _importer(chemin, contenu=None, empreinte="abc123"):
    tables, dictionnaire, contenu_jeu, _ = _jeu()
    return importer_formulaire(
        chemin, dictionnaire, contenu if contenu is not None else contenu_jeu,
        empreinte=empreinte,
    )


def _changements(formulaire):
    return [tuple(None if pd.isna(v) else v for v in ligne)
            for ligne in formulaire.changements.itertuples(index=False)]


def test_import_sans_modification_redonne_le_contenu(tmp_path):
    formulaire = _importer(_exporter(tmp_path))
    assert formulaire.contenu == _jeu()[2]
    assert formulaire.changements.empty
    assert list(formulaire.changements.columns) == list(COLONNES_CHANGEMENTS)
    assert formulaire.avertissements == []


def test_import_des_regles_selon_le_statut(tmp_path):
    chemin = _exporter(tmp_path)
    # Décimales : valide adopte la proposition, documenté prend la saisie (texte
    # d'une cellule au format texte), invalide retire la règle et mémorise le
    # refus, observé efface un refus.
    _modifier(chemin, "décimales", ("factures", "remise"), statut="valide")
    _modifier(chemin, "décimales", ("parametres", "taux"), statut="documenté", décimales=" 3 ")
    _modifier(chemin, "décimales", ("factures", "montant"), statut="invalide")
    _modifier(chemin, "décimales", ("factures", "salaire"), statut="observé")
    # Une saisie sur une ligne qui n'est pas documentée est ignorée, avec un
    # avertissement.
    ligne_annee = _modifier(chemin, "décimales", ("parametres", "annee"), décimales=4)
    # Bornes : proposition observée, saisie avec une borne exclue, règle en vigueur
    # remise à observé (retirée sans refus).
    _modifier(chemin, "bornes", ("factures", "montant"), statut="valide")
    _modifier(
        chemin, "bornes", ("parametres", "taux"), statut="documenté", min="0",
        min_inclus="non", max=1, max_inclus="oui", commentaire="taux de charges",
    )
    _modifier(chemin, "bornes", ("factures", "remise"), statut="observé", commentaire=None)
    _modifier(chemin, "clé_luhn", ("factures", "siret"), statut="documenté")
    formulaire = _importer(chemin)
    regles = formulaire.contenu["regles"]
    assert regles["factures"] == {
        "montant": {"bornes": {"min": 3, "max": 181.45}},
        "remise": {"decimales": 0},
        "code": {"motif": "[A-Z][0-9]+"},
        "siret": {"cle_luhn": True},
    }
    assert regles["parametres"] == {
        "taux": {"decimales": 3, "bornes": {"min": 0, "max": 1, "min_exclu": True}},
    }
    # Ordre du dictionnaire : colonnes, puis types de règle.
    assert list(regles["factures"]) == ["montant", "remise", "code", "siret"]
    assert formulaire.contenu["ecartees"] == {
        "decimales": {"factures": ["montant"]}, "doublons": {"factures": [["siret"]]},
    }
    assert formulaire.contenu["commentaires"] == {
        "bornes": {"parametres": {"taux": "taux de charges"}},
        "doublons": {"factures": {"num": "numéro de facture"}},
    }
    assert formulaire.avertissements == [
        f"feuille « décimales », ligne {ligne_annee} (parametres.annee)\u00a0: valeurs saisies "
        "ignorées (statut observé, seul le statut documenté prend les valeurs saisies)"
    ]
    assert _changements(formulaire) == [
        ("factures", "montant", "decimales", "2", "écartée"),
        ("factures", "remise", "decimales", None, "0"),
        ("factures", "salaire", "decimales", "écartée", None),
        ("factures", "montant", "bornes", None, "3 à 181,45"),
        ("factures", "remise", "bornes", "0 à 15 exclu", None),
        ("factures", "siret", "cle_luhn", None, "oui"),
        ("parametres", "taux", "decimales", None, "3"),
        ("parametres", "taux", "bornes", None, "0 exclu à 1"),
    ]


def test_valeurs_saisies_egales_a_la_regle_en_vigueur_sans_avertissement(tmp_path):
    chemin = _exporter(tmp_path)
    # Les valeurs exportées d'une règle en vigueur restent dans les cellules quand
    # le métier la retire : ce n'est pas une saisie.
    _modifier(chemin, "bornes", ("factures", "remise"), statut="invalide")
    _modifier(chemin, "décimales", ("factures", "montant"), statut="observé")
    formulaire = _importer(chemin)
    assert formulaire.avertissements == []
    assert "remise" not in formulaire.contenu["regles"]["factures"]
    assert formulaire.contenu["ecartees"]["bornes"] == {"factures": ["remise"]}
    # Observé sur une règle en vigueur : retirée sans refus mémorisé.
    assert formulaire.contenu["ecartees"]["decimales"] == {"factures": ["salaire"]}
    assert "montant" not in formulaire.contenu["regles"]["factures"]


def test_problemes_des_regles_listes_en_une_fois(tmp_path):
    chemin = _exporter(tmp_path)
    lignes = [
        _modifier(chemin, "clé_luhn", ("factures", "siret"), statut="valide"),
        _modifier(chemin, "décimales", ("factures", "remise"), statut="documenté"),
        _modifier(chemin, "décimales", ("parametres", "taux"), statut="documenté",
                  décimales="1,5"),
        _modifier(chemin, "bornes", ("factures", "montant"), statut="documenté", min=10,
                  max=2),
        _modifier(chemin, "bornes", ("parametres", "taux"), statut="documenté",
                  max_inclus="non"),
        _modifier(chemin, "bornes", ("factures", "salaire"), statut="valide"),
        _modifier(chemin, "décimales", ("parametres", "annee"), statut="à voir"),
    ]
    with pytest.raises(ValueError) as erreur:
        _importer(chemin)
    message = str(erreur.value)
    assert "7 problèmes" in message
    attendus = [
        "feuille « clé_luhn », ligne {} (factures.siret)\u00a0: statut valide impossible",
        "feuille « décimales », ligne {} (factures.remise)\u00a0: statut documenté sans "
        "règle saisie",
        "feuille « décimales », ligne {} (parametres.taux)\u00a0: décimales\u00a0: un nombre "
        "entier positif ou nul est attendu",
        "feuille « bornes », ligne {} (factures.montant)\u00a0: le minimum dépasse le maximum",
        "feuille « bornes », ligne {} (parametres.taux)\u00a0: max_inclus à non sans max",
        "feuille « bornes », ligne {} (factures.salaire)\u00a0: statut valide impossible",
        "feuille « décimales », ligne {} (parametres.annee)\u00a0: statut 'à voir' inconnu",
    ]
    for attendu, numero in zip(attendus, lignes):
        assert attendu.format(numero) in message


def test_import_des_doublons(tmp_path):
    chemin = _exporter(tmp_path)
    # La proposition retenue avec son type, le groupe écarté repris comme présumé
    # unique, un groupe de deux colonnes ajouté en bas, la recherche des lignes
    # identiques hors clé rétablie.
    _modifier(chemin, "doublons", ("factures", "unique", "code"), statut="valide")
    _modifier(chemin, "doublons", ("factures", "unique", "siret"), statut="valide",
              type="présumé unique", commentaire="SIRET du client")
    _modifier(chemin, "doublons", None, table="factures", type="unique",
              colonnes=" montant ,remise", statut="documenté", commentaire="à voir")
    _modifier(chemin, "tables", ("factures",), identiques_hors_clé="oui")
    formulaire = _importer(chemin)
    assert formulaire.contenu["doublons"] == {
        "factures": {
            "cle": ["num"],
            "unique": [["code"], ["montant", "remise"]],
            "presume_unique": [["siret"]],
        },
    }
    assert "ecartees" not in formulaire.contenu or "doublons" not in formulaire.contenu["ecartees"]
    assert formulaire.contenu["commentaires"]["doublons"] == {
        "factures": {
            "num": "numéro de facture", "siret": "SIRET du client", "montant, remise": "à voir",
        },
    }
    assert _changements(formulaire) == [
        # Groupes d'une colonne dans l'ordre du dictionnaire, puis les autres.
        ("factures", "code", "doublons", None, "unique"),
        ("factures", "siret", "doublons", "écartée", "présumé unique"),
        ("factures", "montant, remise", "doublons", None, "unique"),
        ("factures", None, "identique_hors_cle", "non", "oui"),
    ]


def test_retrait_et_refus_d_un_groupe_en_vigueur(tmp_path):
    chemin = _exporter(tmp_path)
    _modifier(chemin, "doublons", ("factures", "clé", "num"), statut="invalide")
    _modifier(chemin, "tables", ("factures",), identiques_hors_clé="oui")
    formulaire = _importer(chemin)
    assert "doublons" not in formulaire.contenu
    assert formulaire.contenu["ecartees"]["doublons"] == {"factures": [["num"], ["siret"]]}


def test_problemes_des_doublons(tmp_path):
    chemin = _exporter(tmp_path)
    deux_cles = _modifier(chemin, "doublons", ("factures", "unique", "code"), type="clé",
                          statut="valide")
    inconnue = _modifier(chemin, "doublons", None, table="factures", type="unique",
                         colonnes="montant, prix", statut="documenté")
    sans_statut = _modifier(chemin, "doublons", None, table="parametres", type="unique",
                            colonnes="libelle", statut="valide")
    _modifier(chemin, "tables", ("parametres",), identiques_hors_clé="non")
    with pytest.raises(ValueError) as erreur:
        _importer(chemin)
    message = str(erreur.value)
    assert "4 problèmes" in message
    assert (
        f"feuille « doublons »\u00a0: factures, une seule clé par table (lignes 2, {deux_cles})"
    ) in message
    assert (
        f"feuille « doublons », ligne {inconnue} (factures, montant, prix)\u00a0: colonnes "
        "inconnues de la table (prix)"
    ) in message
    assert (
        f"feuille « doublons », ligne {sans_statut} (parametres, libelle)\u00a0: une ligne "
        "ajoutée prend le statut documenté (statut valide)"
    ) in message
    assert "feuille « tables » (parametres)\u00a0: non sans clé déclarée" in message


def test_ligne_en_vigueur_supprimee_et_groupe_en_double(tmp_path):
    chemin = _exporter(tmp_path)
    classeur = load_workbook(chemin)
    classeur["doublons"].delete_rows(2)
    classeur.save(chemin)
    double = _modifier(chemin, "doublons", None, table="factures", type="unique",
                       colonnes="code", statut="documenté")
    with pytest.raises(ValueError) as erreur:
        _importer(chemin)
    message = str(erreur.value)
    assert "la ligne du groupe en vigueur factures, num manque" in message
    assert f"factures, colonnes code présentes sur plusieurs lignes ({double})" not in message
    _modifier(chemin, "doublons", ("factures", "unique", "code"), statut="valide")
    with pytest.raises(ValueError, match="colonnes code présentes sur plusieurs lignes \\(2, "):
        _importer(chemin)


def test_formulaire_perime_ou_modifie(tmp_path):
    chemin = _exporter(tmp_path)
    with pytest.raises(ValueError, match="Formulaire périmé .*empreinte abc123 à l'export, "
                       "def456 aujourd'hui"):
        _importer(chemin, empreinte="def456")
    classeur = load_workbook(chemin)
    classeur["bornes"].title = "limites"
    classeur.save(chemin)
    with pytest.raises(ValueError, match="feuille « bornes » absente"):
        _importer(chemin)
    chemin = _exporter(tmp_path, "autre.xlsx")
    classeur = load_workbook(chemin)
    classeur["doublons"].cell(row=1, column=2).value = "genre"
    classeur.save(chemin)
    with pytest.raises(ValueError, match="feuille « doublons »\u00a0: en-tête modifié"):
        _importer(chemin)


def test_regles_hors_formulaire_reprises_telles_quelles(tmp_path):
    tables, dictionnaire, contenu, sensibles = _jeu()
    contenu["regles"]["archives"] = {"montant": {"decimales": 2}}
    contenu["regles"]["factures"]["num"] = {"motif": "F-[0-9]+", "decimales": 0}
    contenu["doublons"]["archives"] = {"cle": ["id"]}
    contenu["commentaires"]["decimales"] = {"archives": {"montant": "ancienne table"}}
    chemin = exporter_formulaire(
        tables, dictionnaire, contenu, sensibles, "2026-03", tmp_path / "f.xlsx",
        empreinte="e",
    )
    _modifier(chemin, "décimales", ("factures", "remise"), statut="valide")
    formulaire = importer_formulaire(chemin, dictionnaire, contenu, empreinte="e")
    assert formulaire.contenu["regles"]["archives"] == {"montant": {"decimales": 2}}
    # Une règle de décimales sur une colonne de texte n'est pas dans le formulaire.
    assert formulaire.contenu["regles"]["factures"]["num"] == {"decimales": 0, "motif": "F-[0-9]+"}
    assert formulaire.contenu["doublons"]["archives"] == {"cle": ["id"]}
    assert formulaire.contenu["commentaires"]["decimales"] == {
        "archives": {"montant": "ancienne table"}
    }
    # Les tables hors du dictionnaire viennent après les autres.
    assert list(formulaire.contenu["regles"]) == ["factures", "archives"]


def test_texte_du_fichier_des_controles():
    contenu = {
        "regles": {
            "t": {"a": {"decimales": 2, "bornes": {"min": 0, "max": 1, "min_exclu": True}},
                  "mail": {"motif": "[^@\\s]+@[^@\\s]+"}},
        },
        "doublons": {"t": {"cle": ["a"], "unique": [["b", "c"]], "identique_hors_cle": False}},
        "ecartees": {"bornes": {"t": ["b"]}, "doublons": {"t": [["d"]]}},
        "commentaires": {
            "decimales": {"t": {"a": "deux lignes\npuis\u00a0: fin", "b": "oui"}},
            "doublons": {"t": {"b, c": "true"}},
        },
    }
    texte = texte_controles(contenu)
    assert yaml.safe_load(texte) == contenu
    assert texte.startswith(ENTETE_CONTROLES + "\nregles:\n  t:\n")
    assert "    a: {decimales: 2, bornes: {min: 0, max: 1, min_exclu: true}}\n" in texte
    assert "    cle: [a]\n    unique:\n      - [b, c]\n    identique_hors_cle: false\n" in texte
    assert "  bornes:\n    t: [b]\n  doublons:\n    t:\n      - [d]\n" in texte
    assert '      a: "deux lignes\\npuis\u00a0: fin"\n' in texte
    assert "      b, c: 'true'\n" in texte
    assert texte.endswith("\n") and "\r" not in texte
    # Sans section vide.
    assert texte_controles({"regles": {"t": {"a": {"decimales": 1}}}}).count(":\n") == 2


def test_ecriture_avec_copie_datee(tmp_path, monkeypatch):
    import frigeo.dictionnaire as module_dictionnaire
    from datetime import datetime

    monkeypatch.setattr(
        module_dictionnaire, "_maintenant", lambda: datetime(2026, 10, 22, 9, 5, 7)
    )
    chemin, sauvegardes = tmp_path / "config" / "controles.yaml", tmp_path / "sauvegardes"
    # Fichier absent : écrit, sans copie.
    contenu = {"regles": {"t": {"a": {"decimales": 2}}}}
    resultat = ecrire_controles(contenu, chemin, empreinte="absent", dossier=sauvegardes)
    assert resultat == (chemin, None, True)
    assert yaml.safe_load(chemin.read_text(encoding="utf-8")) == contenu
    assert b"\r" not in chemin.read_bytes()
    ancien = chemin.read_bytes()
    # Contenu inchangé : rien n'est écrit.
    empreinte = empreinte_controles(chemin)
    assert ecrire_controles(contenu, chemin, empreinte=empreinte, dossier=sauvegardes) == (
        chemin, None, False,
    )
    assert not sauvegardes.exists()
    # Contenu changé : copie de l'ancien fichier, puis écriture.
    nouveau = {"regles": {"t": {"a": {"decimales": 3}}}}
    resultat = ecrire_controles(nouveau, chemin, empreinte=empreinte, dossier=sauvegardes)
    copie = sauvegardes / "controles_2026-10-22_090507.yaml"
    assert resultat == (chemin, copie, True)
    assert copie.read_bytes() == ancien
    assert yaml.safe_load(chemin.read_text(encoding="utf-8")) == nouveau
    # Fichier modifié depuis la lecture du formulaire : refus.
    with pytest.raises(ValueError, match="a changé depuis la lecture du formulaire"):
        ecrire_controles(contenu, chemin, empreinte=empreinte, dossier=sauvegardes)
    # Copie du même nom déjà présente : rien n'est modifié.
    actuel = chemin.read_bytes()
    with pytest.raises(FileExistsError, match="rien n'a été modifié"):
        ecrire_controles(contenu, chemin, empreinte=empreinte_controles(chemin),
                         dossier=sauvegardes)
    assert chemin.read_bytes() == actuel and copie.read_bytes() == ancien
    # Contenu invalide : refus, tous les problèmes listés.
    with pytest.raises(ValueError, match="règle bornes"):
        ecrire_controles(
            {"regles": {"t": {"a": {"bornes": {"min": 2, "max": 1}}}}, "autre": 1}, chemin,
            empreinte=empreinte_controles(chemin), dossier=sauvegardes,
        )


def test_aller_retour_complet_stable(tmp_path, monkeypatch):
    """Exporter, revoir, écrire, puis réexporter et réimporter sans modification
    redonne exactement le même fichier."""
    tables, dictionnaire, contenu, sensibles = _jeu()
    chemin_yaml = tmp_path / "controles.yaml"
    chemin_yaml.write_text(texte_controles(contenu), encoding="utf-8")
    empreinte = empreinte_controles(chemin_yaml)
    chemin = exporter_formulaire(
        tables, dictionnaire, lire_fichier_controles(chemin_yaml), sensibles, "2026-03",
        tmp_path / "f1.xlsx", empreinte=empreinte,
    )
    _modifier(chemin, "bornes", ("parametres", "taux"), statut="documenté", min=0,
              min_inclus="non", max=1)
    _modifier(chemin, "doublons", ("factures", "unique", "code"), statut="invalide")
    formulaire = importer_formulaire(
        chemin, dictionnaire, lire_fichier_controles(chemin_yaml), empreinte=empreinte
    )
    ecrire_controles(formulaire.contenu, chemin_yaml, empreinte=empreinte,
                     dossier=tmp_path / "s")
    ecrit = chemin_yaml.read_bytes()
    empreinte = empreinte_controles(chemin_yaml)
    # L'ancien formulaire est désormais périmé.
    with pytest.raises(ValueError, match="périmé"):
        importer_formulaire(chemin, dictionnaire, lire_fichier_controles(chemin_yaml),
                            empreinte=empreinte)
    chemin = exporter_formulaire(
        tables, dictionnaire, lire_fichier_controles(chemin_yaml), sensibles, "2026-03",
        tmp_path / "f2.xlsx", empreinte=empreinte,
    )
    doublons = _lire(load_workbook(chemin)["doublons"])
    lignes = {(l["table"], l["colonnes"]): l["statut"] for l in doublons}
    assert lignes[("factures", "code")] == "invalide"
    formulaire = importer_formulaire(
        chemin, dictionnaire, lire_fichier_controles(chemin_yaml), empreinte=empreinte
    )
    assert formulaire.changements.empty
    assert ecrire_controles(formulaire.contenu, chemin_yaml, empreinte=empreinte,
                            dossier=tmp_path / "s").ecrit is False
    assert chemin_yaml.read_bytes() == ecrit


def test_contenu_remis_dans_l_ordre_du_dictionnaire(tmp_path):
    tables, dictionnaire, contenu, sensibles = _jeu()
    contenu["regles"] = {
        "parametres": {"taux": {"bornes": {"min": 0, "max": 1}, "decimales": 2}},
        "factures": {"remise": {"decimales": 0}, "montant": {"decimales": 2}},
    }
    contenu["ecartees"]["decimales"] = {"factures": ["siret", "salaire"]}
    chemin = exporter_formulaire(
        tables, dictionnaire, contenu, sensibles, "2026-03", tmp_path / "f.xlsx",
        empreinte="e",
    )
    formulaire = importer_formulaire(chemin, dictionnaire, contenu, empreinte="e")
    regles = formulaire.contenu["regles"]
    assert list(regles) == ["factures", "parametres"]
    assert list(regles["factures"]) == ["montant", "remise"]
    assert list(regles["parametres"]["taux"]) == ["decimales", "bornes"]
    assert formulaire.contenu["ecartees"]["decimales"] == {"factures": ["siret", "salaire"]}
    assert list(formulaire.contenu) == ["regles", "doublons", "ecartees", "commentaires"]
    assert formulaire.changements.empty


def test_nombre_de_decimales_negatif_refuse(tmp_path):
    chemin = _exporter(tmp_path)
    numero = _modifier(chemin, "décimales", ("factures", "remise"), statut="documenté",
                       décimales="-1")
    with pytest.raises(ValueError) as erreur:
        _importer(chemin)
    assert (
        f"ligne {numero} (factures.remise) : décimales : un nombre entier positif "
        "ou nul est attendu"
    ) in str(erreur.value)


# --- Application d'un formulaire (morceau d) -------------------------------------


def _horloge(monkeypatch, *instant):
    import frigeo.dictionnaire as module_dictionnaire
    from datetime import datetime

    monkeypatch.setattr(module_dictionnaire, "_maintenant", lambda: datetime(*instant))


def test_archivage_d_un_formulaire(tmp_path, monkeypatch):
    _horloge(monkeypatch, 2026, 10, 22, 18, 4, 5)
    chemin = tmp_path / "formulaires" / "controles_2026-03.xlsx"
    chemin.parent.mkdir()
    chemin.write_bytes(b"classeur")
    archive = archiver_formulaire(chemin)
    assert archive == tmp_path / "formulaires" / "appliques" / (
        "controles_2026-03_2026-10-22_180405.xlsx"
    )
    assert archive.read_bytes() == b"classeur" and not chemin.exists()
    assert dossier_archives(chemin) == archive.parent
    # Une archive du même nom n'est jamais remplacée.
    chemin.write_bytes(b"autre")
    with pytest.raises(FileExistsError, match="rien n'a été modifié"):
        archiver_formulaire(chemin)
    assert chemin.read_bytes() == b"autre" and archive.read_bytes() == b"classeur"


def test_classeur_ouvert_non_archive(tmp_path, monkeypatch):
    from pathlib import Path

    chemin = tmp_path / "controles_2026-03.xlsx"
    chemin.write_bytes(b"classeur")

    def refuser(self, cible):
        raise PermissionError("ouvert")

    monkeypatch.setattr(Path, "rename", refuser)
    with pytest.raises(PermissionError, match="fermer le classeur dans Excel"):
        archiver_formulaire(chemin, tmp_path / "a")
    assert chemin.exists()


def _application(tmp_path, monkeypatch, modifier=True):
    """Fichier des contrôles du jeu d'essai, formulaire exporté (et revu si
    `modifier`), puis importé."""
    _horloge(monkeypatch, 2026, 10, 22, 18, 30, 0)
    tables, dictionnaire, contenu, sensibles = _jeu()
    chemin_yaml = tmp_path / "config" / "controles.yaml"
    chemin_yaml.parent.mkdir()
    chemin_yaml.write_text(texte_controles(contenu), encoding="utf-8")
    contenu = lire_fichier_controles(chemin_yaml)
    empreinte = empreinte_controles(chemin_yaml)
    chemin = exporter_formulaire(
        tables, dictionnaire, contenu, sensibles, "2026-03",
        tmp_path / "formulaires" / "controles_2026-03.xlsx", empreinte=empreinte,
    )
    if modifier:
        _modifier(chemin, "décimales", ("factures", "remise"), statut="valide")
    formulaire = importer_formulaire(chemin, dictionnaire, contenu, empreinte=empreinte)
    options = {"empreinte": empreinte, "chemin": chemin_yaml, "dossier": tmp_path / "s"}
    return formulaire, chemin, contenu, chemin_yaml, options


def test_formulaire_non_rempli_rien_n_est_fait(tmp_path, monkeypatch):
    formulaire, chemin, contenu, chemin_yaml, options = _application(
        tmp_path, monkeypatch, modifier=False
    )
    avant = chemin_yaml.read_bytes()
    assert appliquer_formulaire(formulaire, chemin, contenu, **options) == (False, None, None)
    assert chemin.exists() and chemin_yaml.read_bytes() == avant
    assert not (tmp_path / "s").exists()


def test_formulaire_applique_puis_archive(tmp_path, monkeypatch):
    formulaire, chemin, contenu, chemin_yaml, options = _application(tmp_path, monkeypatch)
    avant = chemin_yaml.read_bytes()
    resultat = appliquer_formulaire(formulaire, chemin, contenu, **options)
    archive = chemin.parent / "appliques" / "controles_2026-03_2026-10-22_183000.xlsx"
    assert resultat == (True, tmp_path / "s" / "controles_2026-10-22_183000.yaml", archive)
    assert not chemin.exists() and archive.exists()
    assert resultat.sauvegarde.read_bytes() == avant
    assert lire_fichier_controles(chemin_yaml)["regles"]["factures"]["remise"] == {
        "decimales": 0, "bornes": {"min": 0, "max": 15, "max_exclu": True},
    }


def test_ecriture_refusee_formulaire_remis_en_place(tmp_path, monkeypatch):
    formulaire, chemin, contenu, chemin_yaml, options = _application(tmp_path, monkeypatch)
    chemin_yaml.write_text(chemin_yaml.read_text(encoding="utf-8") + "\n", encoding="utf-8")
    modifie = chemin_yaml.read_bytes()
    with pytest.raises(ValueError, match="a changé depuis la lecture du formulaire"):
        appliquer_formulaire(formulaire, chemin, contenu, **options)
    assert chemin.exists() and not list((chemin.parent / "appliques").iterdir())
    assert chemin_yaml.read_bytes() == modifie


def test_bornes_incluses_par_defaut_ou_par_oui(tmp_path):
    chemin = _exporter(tmp_path)
    # Vide ou « oui » inclut la borne, « non » l'exclut.
    _modifier(chemin, "bornes", ("parametres", "taux"), statut="documenté", min=0,
              min_inclus="oui", max="0,9")
    _modifier(chemin, "bornes", ("factures", "remise"), max_inclus=None)
    formulaire = _importer(chemin)
    assert formulaire.contenu["regles"]["parametres"]["taux"] == {
        "bornes": {"min": 0, "max": 0.9}
    }
    assert formulaire.contenu["regles"]["factures"]["remise"] == {
        "bornes": {"min": 0, "max": 15}
    }
    assert ("factures", "remise", "bornes", "0 à 15 exclu", "0 à 15") in _changements(formulaire)
    assert ("parametres", "taux", "bornes", None, "0 à 0,9") in _changements(formulaire)
    numero = _modifier(chemin, "bornes", ("factures", "montant"), statut="documenté", min=1,
                       min_inclus="peut-être")
    with pytest.raises(ValueError) as erreur:
        _importer(chemin)
    assert (
        f"ligne {numero} (factures.montant) : min_inclus : oui, non ou vide est attendu"
    ) in str(erreur.value)
