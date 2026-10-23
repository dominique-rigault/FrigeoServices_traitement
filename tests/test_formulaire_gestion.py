"""Tests du formulaire des règles de gestion (frigeo.formulaire_gestion)."""

import datetime as dt

import pandas as pd
import pytest
from openpyxl import load_workbook

from frigeo.formulaire_gestion import (
    COLONNES_AUTRES,
    COLONNES_FERIES,
    COLONNES_NOUVELLES,
    COLONNES_PERIODES,
    FEUILLE_AUTRES,
    FEUILLE_FERIES,
    FEUILLE_MODE_EMPLOI,
    FEUILLE_NOUVELLES,
    FEUILLE_PERIODES,
    REPERE_EMPREINTE,
    annees_des_exports,
    empreinte_gestion,
    exporter_formulaire_gestion,
    feries_france,
    lignes_autres,
    lignes_feries,
    lignes_nouvelles,
    lignes_periodes,
    problemes_contenu_formulaire,
    texte_parametres,
)
from frigeo.gestion import lire_fichier_gestion

J = dt.date


def _export(periodes, **colonnes):
    return pd.DataFrame({"periode_export": periodes, **colonnes}, dtype=object)


def _tables():
    """Un export « t » (six dates par colonne) et un référentiel « r »."""
    return {
        "t": _export(
            ["2025"] * 5 + ["2026-01"],
            d=["01/02/2025"] * 5 + ["05/01/2026"],
            e=["01/02/2025"] * 5 + ["05/01/2026"],
            f=["01/02/2025"] * 4 + ["01/02/2052", "05/01/2026"],
            g=["01/02/2025"] * 6,
            h=["01/02/2025", "", "x", None, "01/03/2025", "05/01/2026"],
            n=["1"] * 6,
        ),
        "r": pd.DataFrame({"code": ["A"], "debut": ["01/01/2020"]}, dtype=object),
    }


def _dictionnaire():
    date = {"nature": {"regle": "date JJ/MM/AAAA"}}
    return {
        "r": {"code": {"nature": {"regle": "texte"}}, "debut": date},
        "t": {c: date for c in "defgh"} | {"n": {"nature": {"regle": "entier (texte)"}}},
    }


def _periode(colonne, ref="decoupage", **autres):
    return {"ref": ref, "libelle": f"{colonne} dans la période", "type":
            "date_dans_periode_export", "table": "t", "colonne": colonne, **autres}


def _contenu():
    return {
        "jours_feries": {2025: ["2025-01-01", "2025-12-26"]},
        "regles": [
            _periode("d", issue="anomalie", commentaire="Vu avec le métier"),
            {"ref": "R22", "libelle": "Distance", "type": "plage", "table": "t",
             "colonne": "n", "min": 2, "max": 120},
            _periode("d", ref="autre", filtre=[{"colonne": "n", "valeurs": ["1", "2"]}]),
            _periode("e", filtre=[{"colonne": "n", "motif": "1.*"}]),
            _periode("d", filtre=[{"colonne": "n", "valeurs": ["1"]},
                                  {"colonne": "e", "valeurs": ["x"]}]),
        ],
        "retirees": [_periode("h", libelle="Date h retirée")],
        "ecartees": {
            "periodes_export": [{"table": "t", "colonne": "g"}],
            "jours_feries": {2025: ["2025-05-08"]},
        },
        "demandes": [{"regle": "Un avoir suit sa facture", "tables": "factures",
                      "issue": "à arbitrer"}],
    }


def test_feries_france():
    assert list(feries_france(2025)) == [
        J(2025, 1, 1), J(2025, 4, 21), J(2025, 5, 1), J(2025, 5, 8), J(2025, 5, 29),
        J(2025, 6, 9), J(2025, 7, 14), J(2025, 8, 15), J(2025, 11, 1), J(2025, 11, 11),
        J(2025, 12, 25),
    ]
    assert feries_france(2027)[J(2027, 5, 6)] == "Ascension"
    assert list(feries_france(2027))[3:5] == [J(2027, 5, 6), J(2027, 5, 8)]
    # Pâques aux dates extrêmes du calendrier grégorien (22 mars et 25 avril).
    assert feries_france(2285)[J(2285, 3, 23)] == "Lundi de Pâques"
    assert feries_france(2038)[J(2038, 4, 26)] == "Lundi de Pâques"
    assert feries_france(2024)[J(2024, 5, 20)] == "Lundi de Pentecôte"
    # Années où la règle corrige le 26 ou le 25 avril (Pâques le 19 ou le 18 avril).
    assert feries_france(2049)[J(2049, 4, 19)] == "Lundi de Pâques"
    assert feries_france(2076)[J(2076, 4, 20)] == "Lundi de Pâques"
    assert len(feries_france(2026)) == 11


def _reduites(lignes, *colonnes):
    return [tuple(ligne.get(c) for c in colonnes) for ligne in lignes]


def test_lignes_periodes():
    lignes = lignes_periodes(_tables(), _dictionnaire(), _contenu())
    assert _reduites(lignes, "colonne", "origine", "statut", "ref", "issue",
                     "filtre_colonne", "filtre_valeurs", "rang") == [
        ("d", "en vigueur", "documenté", "decoupage", "anomalie", None, None, 0),
        ("d", "en vigueur", "documenté", "autre", "à arbitrer", "n", "1, 2", 2),
        # e a une règle à filtre par motif (feuille des autres règles) : pas de
        # proposition.
        ("f", "proposée", "observé", "decoupage", "à arbitrer", None, None, None),
        ("h", "retirée", "invalide", "decoupage", "à arbitrer", None, None, 0),
        ("g", "écartée", "invalide", "decoupage", "à arbitrer", None, None, None),
    ]
    assert _reduites(lignes, "colonne", "nature", "dates", "observé", "proposition") == [
        ("d", "date JJ/MM/AAAA", 6, "6 sur 6 dans la période de l'export (100 %)", "oui"),
        ("d", "date JJ/MM/AAAA", 6, "6 sur 6 dans la période de l'export (100 %)", "oui"),
        ("f", "date JJ/MM/AAAA", 6, "5 sur 6 dans la période de l'export (83,3 %)",
         "aucune"),
        # Cellules vides ou illisibles non comptées.
        ("h", "date JJ/MM/AAAA", 3, "3 sur 3 dans la période de l'export (100 %)", "oui"),
        ("g", "date JJ/MM/AAAA", 6, "5 sur 6 dans la période de l'export (83,3 %)",
         "aucune"),
    ]
    assert lignes[0]["libellé"] == "d dans la période"
    assert lignes[0]["commentaire"] == "Vu avec le métier"
    assert lignes[2]["libellé"] == "f dans la période de l'export"
    assert all(set(ligne) <= set(COLONNES_PERIODES) for ligne in lignes)
    # Seuil de proposition : 95 % des dates, sans compter les lignes sans période.
    # La fin de la période est exclue : le 1er janvier 2026 est hors de l'export 2025.
    tables = {"t": _export(["2025", "2025"], d=["31/12/2025", "01/01/2026"])}
    assert lignes_periodes(tables, _dictionnaire(), {})[0]["observé"] == (
        "1 sur 2 dans la période de l'export (50 %)"
    )
    tables = _tables()
    tables["t"]["periode_export"] = ["2025"] * 5 + ["inconnue"]
    lignes = lignes_periodes(tables, _dictionnaire(), {})
    assert _reduites(lignes, "colonne", "dates", "proposition") == [
        ("d", 5, "oui"), ("e", 5, "oui"), ("f", 5, "aucune"), ("g", 5, "oui"),
        ("h", 2, "oui"),
    ]


def test_regle_sur_une_colonne_hors_des_exports():
    contenu = {"regles": [_periode("debut") | {"table": "r"}, _periode("zz")]}
    lignes = lignes_periodes({"r": _tables()["r"]}, _dictionnaire(), contenu)
    assert _reduites(lignes, "table", "colonne", "dates", "observé", "statut") == [
        ("r", "debut", None, None, "documenté"),
        ("t", "zz", None, None, "documenté"),
    ]


def test_annees_et_lignes_feries():
    tables = _tables()
    assert annees_des_exports(tables) == [2025, 2026]
    lignes = lignes_feries(tables, _contenu())
    de_2025 = [ligne for ligne in lignes if ligne["année"] == 2025]
    assert len(de_2025) == 12
    assert _reduites(de_2025[:4], "date", "jour", "fête", "origine", "statut") == [
        (J(2025, 1, 1), "mercredi", "Jour de l'an", "en vigueur", "documenté"),
        (J(2025, 4, 21), "lundi", "Lundi de Pâques", "proposée", "observé"),
        (J(2025, 5, 1), "jeudi", "Fête du travail", "proposée", "observé"),
        (J(2025, 5, 8), "jeudi", "Victoire 1945", "écartée", "invalide"),
    ]
    assert _reduites(de_2025[-2:], "date", "fête", "origine", "statut") == [
        (J(2025, 12, 25), "Noël", "proposée", "observé"),
        (J(2025, 12, 26), None, "en vigueur", "documenté"),
    ]
    de_2026 = [ligne for ligne in lignes if ligne["année"] == 2026]
    assert {ligne["statut"] for ligne in de_2026} == {"observé"}
    assert len(de_2026) == 11
    assert all(set(ligne) == set(COLONNES_FERIES) for ligne in lignes)
    # Une année déclarée ou écartée hors des exports est montrée.
    lignes = lignes_feries(
        {}, {"jours_feries": {2030: []}, "ecartees": {"jours_feries": {2031: ["2031-01-01"]}}}
    )
    assert {ligne["année"] for ligne in lignes} == {2030, 2031}


def test_lignes_autres_et_parametres_en_clair():
    lignes = lignes_autres(_contenu())
    assert _reduites(lignes, "ref", "type", "paramètres", "filtre", "issue", "origine",
                     "statut", "rang") == [
        ("R22", "plage", "colonne = n; min = 2; max = 120", None, "anomalie", "en vigueur",
         "documenté", 1),
        ("decoupage", "date_dans_periode_export", "colonne = e", "n suit le motif 1.*",
         "à arbitrer", "en vigueur", "documenté", 3),
        ("decoupage", "date_dans_periode_export", "colonne = d", "n = 1 et e = x",
         "à arbitrer", "en vigueur", "documenté", 4),
    ]
    assert all(set(ligne) <= set(COLONNES_AUTRES) for ligne in lignes)
    regle = {
        "ref": "x", "libelle": "x", "type": "sans_ligne_liee", "table": "t", "cle": ["a", "b"],
        "table_lignes": "l",
        "filtre_lignes": [{"colonne": "c", "valeurs": ["V"]}, {"colonne": "d", "motif": "M.*"}],
        "exceptions": [{"libelle": "cas normal",
                        "conditions": [{"colonne": "s", "valeurs": ["A", "B"]}]}],
        "issue": "information", "commentaire": "c",
    }
    assert texte_parametres(regle) == (
        "cle = a, b; table_lignes = l; filtre_lignes = c = V et d suit le motif M.*; "
        "exceptions = « cas normal » si s = A ou B"
    )
    regle = {
        "ref": "x", "libelle": "x", "type": "produit_calcule", "table": "t", "cible": "m",
        "facteurs": ["q", {"colonne": "r", "constante": 1, "coefficient": -0.01, "vide": 0}],
        "tolerance": 0.005,
    }
    assert texte_parametres(regle) == (
        "cible = m; facteurs = [q, {colonne: r, constante: 1, coefficient: -0.01, vide: 0}]; "
        "tolerance = 0.005"
    )
    regle = {
        "ref": "x", "libelle": "x", "type": "solde_jamais_negatif", "table": "t",
        "date_initial": dt.date(2025, 1, 1), "hors": True,
    }
    assert texte_parametres(regle) == "date_initial = 2025-01-01; hors = true"


def test_lignes_nouvelles():
    assert lignes_nouvelles(_contenu()) == [
        {"règle": "Un avoir suit sa facture", "tables": "factures", "colonnes": None,
         "tolérance_ou_délai": None, "issue_souhaitée": "à arbitrer", "commentaire": None}
    ]
    assert set(lignes_nouvelles(_contenu())[0]) == set(COLONNES_NOUVELLES)
    assert lignes_nouvelles({}) == []


def test_contenu_lu_et_problemes(tmp_path):
    chemin = tmp_path / "gestion.yaml"
    chemin.write_text(
        "jours_feries: {2025: ['2025-01-01']}\nregles: []\nretirees: []\n"
        "ecartees: {periodes_export: [{table: t, colonne: d}]}\n"
        "demandes: [{regle: Une règle}]\n",
        encoding="utf-8",
    )
    contenu = lire_fichier_gestion(chemin)
    assert problemes_contenu_formulaire(contenu) == []
    contenu = {
        "jours_feries": {2025: ["2026-01-01"]},
        "regles": [_periode("d"), {"ref": "x"}],
        "retirees": [_periode("d", issue="grave")],
        "ecartees": {
            "periodes_export": [{"table": "t"}],
            "jours_feries": {"an": []},
            "autre": [],
        },
        "demandes": [{"regle": " "}, "texte", {"regle": "r", "delai": "3 j", "tables": 2}],
    }
    horaires = {"ref": "R04", "libelle": "h", "type": "hors_horaires_selon_valeur",
                "table": "t", "date": "d", "heure": "h", "colonne": "c", "valeurs": [],
                "horaires": {"lundi": [{"debut": "08:00", "fin": "18:00"}]}}
    contenu["regles"].append(horaires)
    assert problemes_contenu_formulaire(contenu) == [
        "jours_feries : année 2025 : 2026-01-01 n'est pas de cette année",
        "regles, règle n° 2 (x) : clé manquante (libelle, type, table)",
        "regles, règle n° 3 (R04) : paramètre manquant (feries) (les jours fériés peuvent "
        "aussi se déclarer une fois pour le fichier, sous la clé jours_feries)",
        "retirees, règle n° 1 (decoupage), issue : issue inconnue 'grave' "
        "(issues possibles : anomalie, à arbitrer, information)",
        "ecartees, periodes_export : une liste de {table, colonne} (textes) est attendue",
        "ecartees, jours_feries : année 'an' : une année AAAA est attendue",
        "ecartees : famille inconnue 'autre' (familles possibles : periodes_export, "
        "jours_feries)",
        "demandes, demande n° 1, regle : un texte non vide est attendu",
        "demandes, demande n° 2 : un dictionnaire de clés est attendu",
        "demandes, demande n° 3 : clé inconnue delai (clés possibles : regle, tables, "
        "colonnes, tolerance, issue, commentaire)",
        "demandes, demande n° 3 : des textes sont attendus",
    ]
    assert problemes_contenu_formulaire({"ecartees": [], "demandes": {}}) == [
        "ecartees : un dictionnaire par famille est attendu",
        "demandes : une liste est attendue",
    ]


def test_empreinte_gestion(tmp_path):
    chemin = tmp_path / "gestion.yaml"
    assert empreinte_gestion(chemin) == "absent"
    chemin.write_bytes(b"regles:\n  - ref: x\n")
    lf = empreinte_gestion(chemin)
    chemin.write_bytes(b"regles:\r\n  - ref: x\r\n")
    assert empreinte_gestion(chemin) == lf
    chemin.write_bytes(b"regles:\n  - ref: y\n")
    assert empreinte_gestion(chemin) != lf


def _verrouillee(feuille, ligne, nom, colonnes):
    return feuille.cell(row=ligne, column=colonnes.index(nom) + 1).protection.locked


def test_export_du_classeur(tmp_path):
    chemin = tmp_path / "formulaires" / "gestion_2026-03.xlsx"
    assert exporter_formulaire_gestion(
        _tables(), _dictionnaire(), _contenu(), "2026-03", chemin, empreinte="abc123"
    ) == chemin
    classeur = load_workbook(chemin)
    assert classeur.sheetnames == [
        FEUILLE_MODE_EMPLOI, FEUILLE_PERIODES, FEUILLE_FERIES, FEUILLE_AUTRES,
        FEUILLE_NOUVELLES, "listes",
    ]
    assert classeur["listes"].sheet_state == "hidden"
    notice = classeur[FEUILLE_MODE_EMPLOI]
    textes = [ligne[0] for ligne in notice.iter_rows(values_only=True)]
    rang = textes.index(REPERE_EMPREINTE) + 1
    assert notice.cell(row=rang, column=2).value == "abc123"
    assert any(
        texte and texte.startswith("Feuille « périodes_export » : 1 lignes à « observé » sur 5")
        for texte in textes
    )

    feuille = classeur[FEUILLE_PERIODES]
    assert [c.value for c in feuille[1]] == list(COLONNES_PERIODES)
    assert feuille.column_dimensions["O"].hidden
    assert feuille.cell(row=2, column=15).value == 0
    assert _verrouillee(feuille, 2, "table", COLONNES_PERIODES)
    assert _verrouillee(feuille, 2, "filtre_colonne", COLONNES_PERIODES)
    assert not _verrouillee(feuille, 2, "statut", COLONNES_PERIODES)
    assert not _verrouillee(feuille, 2, "libellé", COLONNES_PERIODES)
    # Lignes d'ajout : table, colonne et filtre se saisissent.
    assert not _verrouillee(feuille, 8, "table", COLONNES_PERIODES)
    assert not _verrouillee(feuille, 8, "filtre_colonne", COLONNES_PERIODES)
    assert _verrouillee(feuille, 8, "observé", COLONNES_PERIODES)
    plages = {str(menu.sqref): menu.formula1 for menu in feuille.data_validations.dataValidation}
    assert plages["H2:H56"] == '"observé,valide,invalide,documenté"'
    assert plages["K2:K56"] == '"anomalie,à arbitrer,information"'
    assert "A7:A56" in plages and "B7:B56" in plages

    feuille = classeur[FEUILLE_FERIES]
    assert feuille.cell(row=2, column=2).value == dt.datetime(2025, 1, 1)
    assert feuille.cell(row=2, column=2).number_format == "DD/MM/YYYY"
    # 12 dates en 2025, 11 en 2026, puis les lignes d'ajout.
    assert feuille.max_row == 1 + 23 + 50
    assert not _verrouillee(feuille, 25, "date", COLONNES_FERIES)
    assert feuille.cell(row=25, column=2).number_format == "DD/MM/YYYY"
    assert _verrouillee(feuille, 2, "date", COLONNES_FERIES)

    feuille = classeur[FEUILLE_AUTRES]
    assert feuille.max_row == 4
    assert not _verrouillee(feuille, 2, "issue", COLONNES_AUTRES)
    assert _verrouillee(feuille, 2, "paramètres", COLONNES_AUTRES)
    feuille = classeur[FEUILLE_NOUVELLES]
    assert feuille.cell(row=2, column=1).value == "Un avoir suit sa facture"
    assert not _verrouillee(feuille, 30, "règle", COLONNES_NOUVELLES)

    # Un formulaire existant n'est pas remplacé; un contenu mal formé est refusé.
    with pytest.raises(FileExistsError, match="export refusé"):
        exporter_formulaire_gestion(
            _tables(), _dictionnaire(), _contenu(), "2026-03", chemin, empreinte="x"
        )
    with pytest.raises(ValueError, match="demandes : une liste est attendue"):
        exporter_formulaire_gestion(
            _tables(), _dictionnaire(), {"demandes": {}}, "2026-03", chemin,
            empreinte="x", ecraser=True,
        )
    exporter_formulaire_gestion(
        _tables(), _dictionnaire(), {}, "2026-03", chemin, empreinte="absent", ecraser=True
    )
    assert load_workbook(chemin)[FEUILLE_AUTRES].max_row == 1
